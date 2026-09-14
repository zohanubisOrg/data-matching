from __future__ import annotations

import json
import logging
from pathlib import Path

from openpyxl import load_workbook

from ..models import POSProduct
from ..vision.barcode_utils import barcode_lookup_keys, normalize_barcode

logger = logging.getLogger("product_mapper")


def load_knowledge_base(path: Path) -> dict[str, dict[str, str]]:
    """Load human-approved mappings used as highest-priority deterministic evidence."""
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("knowledge base root must be an object")
        logger.info("Loaded %d approved mappings from knowledge base: %s", len(data), path)
        return data
    except Exception as exc:
        logger.error("Failed to load knowledge base from %s: %s", path, exc)
        return {}


def save_knowledge_base(mappings: dict[str, dict[str, str]], path: Path) -> None:
    """Save approved mappings without synthesizing any POS identity fields."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(mappings, indent=2, ensure_ascii=False), encoding="utf-8")
    logger.info("Saved %d mappings to knowledge base: %s", len(mappings), path)


def resolve_approved_mapping(entry: dict[str, str], products: list[POSProduct]) -> list[POSProduct]:
    """Resolve a human-approved tuple against current POS rows exactly.

    New review workbooks persist Product ID + SKU + Barcode, which can disambiguate two
    POS rows sharing one barcode. Legacy barcode-only knowledge remains supported but
    intentionally stays ambiguous when the current POS contains duplicate barcodes.
    """
    product_id = str(entry.get("product_id") or "").strip()
    sku = str(entry.get("sku") or "").strip()
    barcode = normalize_barcode(entry.get("barcode"))
    barcode_keys = set(barcode_lookup_keys(barcode))

    hits: list[POSProduct] = []
    for product in products:
        if product_id and product.product_id != product_id:
            continue
        if sku and product.sku != sku:
            continue
        if barcode_keys and not barcode_keys.intersection(barcode_lookup_keys(product.barcode)):
            continue
        hits.append(product)
    return hits


def _cell_text(row: tuple, col_map: dict[str, int], header: str) -> str:
    index = col_map.get(header)
    if index is None or index >= len(row):
        return ""
    return str(row[index] or "").strip()


def _approved_entry(
    selection: str,
    barcode: str = "",
    product_id: str = "",
    sku: str = "",
    source: str = "customer_review",
) -> dict[str, str] | None:
    selection = str(selection or "").strip()
    if selection in {"1", "2", "3"}:
        if not any(str(value or "").strip() for value in (barcode, product_id, sku)):
            return None
        return {
            "barcode": str(barcode or "").strip(),
            "product_id": str(product_id or "").strip(),
            "sku": str(sku or "").strip(),
            "selection": selection,
            "source": str(source or "customer_review").strip() or "customer_review",
        }
    if selection.casefold() in {"no match", "nomatch"}:
        return {
            "barcode": "",
            "product_id": "",
            "sku": "",
            "selection": "No Match",
            "source": str(source or "customer_review").strip() or "customer_review",
        }
    return None


def import_review_decisions(review_file: Path, kb_file: Path) -> int:
    """Persist customer Excel review selections as deterministic POS identity tuples."""
    if not review_file.exists():
        raise FileNotFoundError(f"Review workbook not found: {review_file}")

    wb = load_workbook(review_file, data_only=True)
    if "Review" not in wb.sheetnames:
        raise ValueError("Workbook missing 'Review' sheet")

    ws = wb["Review"]
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    col_map = {header: index for index, header in enumerate(headers)}
    for required in ("Image", "Customer Selection"):
        if required not in col_map:
            raise ValueError(f"Missing required column in Review sheet: {required}")

    kb = load_knowledge_base(kb_file)
    added_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        image_name = _cell_text(row, col_map, "Image")
        selection = _cell_text(row, col_map, "Customer Selection")
        if not image_name or not selection:
            continue

        if selection in {"1", "2", "3"}:
            candidate_index = int(selection)
            entry = _approved_entry(
                selection=selection,
                barcode=_cell_text(row, col_map, f"Candidate {candidate_index} Barcode"),
                product_id=_cell_text(row, col_map, f"Candidate {candidate_index} Product ID"),
                sku=_cell_text(row, col_map, f"Candidate {candidate_index} SKU"),
                source="customer_review",
            )
        else:
            entry = _approved_entry(selection=selection, source="customer_review")

        if entry is not None:
            kb[image_name] = entry
            added_count += 1

    save_knowledge_base(kb, kb_file)
    return added_count


def import_review_json_decisions(review_file: Path, kb_file: Path) -> int:
    """Merge decisions exported by review_report.html into the deterministic KB."""
    if not review_file.exists():
        raise FileNotFoundError(f"Review JSON not found: {review_file}")
    data = json.loads(review_file.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Review JSON root must be an object keyed by image filename")

    kb = load_knowledge_base(kb_file)
    added_count = 0
    for raw_image_name, raw_entry in data.items():
        image_name = str(raw_image_name or "").strip()
        if not image_name or not isinstance(raw_entry, dict):
            continue
        entry = _approved_entry(
            selection=str(raw_entry.get("selection") or ""),
            barcode=str(raw_entry.get("barcode") or ""),
            product_id=str(raw_entry.get("product_id") or ""),
            sku=str(raw_entry.get("sku") or ""),
            source=str(raw_entry.get("source") or "review_report_html"),
        )
        if entry is not None:
            kb[image_name] = entry
            added_count += 1

    save_knowledge_base(kb, kb_file)
    return added_count

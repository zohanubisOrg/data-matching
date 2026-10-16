"""Atomic dual-image embedder inserting Webshop Front and Back photos into Excel."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Dict, Optional, Tuple

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image


class ExcelDualEmbedder:
    """Inserts two dedicated image columns into the master catalog spreadsheet."""

    def __init__(self, master_excel_path: Path):
        self.master_excel_path = master_excel_path

    @staticmethod
    def _create_thumbnail_image(img_path: Path, max_dim: int = 150) -> Optional[OpenpyxlImage]:
        """Creates a memory-optimized thumbnail image for embedding into an Excel cell."""
        try:
            pil_img = Image.open(img_path)
            if pil_img.mode != "RGB":
                pil_img = pil_img.convert("RGB")
            # Calculate thumbnail dimensions fitting row height ~120pt (160px)
            pil_img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)

            buf = io.BytesIO()
            pil_img.save(buf, format="JPEG", quality=85, optimize=True)
            buf.seek(0)

            xl_img = OpenpyxlImage(buf)
            xl_img.width = pil_img.width
            xl_img.height = pil_img.height
            return xl_img
        except Exception:
            return None

    def embed_dual_images(
        self,
        output_excel_path: Path,
        product_front_map: Dict[str, Path],  # record_id -> front_hd_path
        product_back_map: Dict[str, Path],   # record_id -> back_hd_path
        target_record_ids: Optional[set] = None,
    ) -> Path:
        """Embeds front webshop and back photos adjacent to Column W ('Ảnh mặt trước')."""
        output_excel_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.load_workbook(self.master_excel_path)
        ws = wb.active

        # Header is at row 3
        # Column W is column 23 ('Ảnh mặt trước')
        front_demo_col_idx = 23
        webshop_col_idx = 24  # Column X
        back_col_idx = 25     # Column Y

        # Insert 2 new columns right after Column W
        ws.insert_cols(webshop_col_idx, amount=2)

        # Style headers
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        cell_webshop = ws.cell(3, webshop_col_idx, value="Ảnh webshop")
        cell_webshop.fill = header_fill
        cell_webshop.font = header_font
        cell_webshop.alignment = Alignment(horizontal="center", vertical="center")

        cell_back = ws.cell(3, back_col_idx, value="Ảnh mặt sau")
        cell_back.fill = header_fill
        cell_back.font = header_font
        cell_back.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions[openpyxl.utils.get_column_letter(webshop_col_idx)].width = 24
        ws.column_dimensions[openpyxl.utils.get_column_letter(back_col_idx)].width = 24

        # Iterate rows
        for r in range(4, ws.max_row + 1):
            record_id_val = ws.cell(r, 5).value  # Record ID is in Column E (col 5)
            if not record_id_val:
                continue
            record_id = str(record_id_val).strip()

            if target_record_ids is not None and record_id not in target_record_ids:
                continue

            # Lock row height for clean cell presentation
            ws.row_dimensions[r].height = 120

            # Embed Front Webshop Image
            front_path = product_front_map.get(record_id)
            if front_path and front_path.exists():
                xl_front = self._create_thumbnail_image(front_path, max_dim=150)
                if xl_front is not None:
                    col_letter = openpyxl.utils.get_column_letter(webshop_col_idx)
                    cell_ref = f"{col_letter}{r}"
                    ws.add_image(xl_front, cell_ref)

            # Embed Back Image
            back_path = product_back_map.get(record_id)
            if back_path and back_path.exists():
                xl_back = self._create_thumbnail_image(back_path, max_dim=150)
                if xl_back is not None:
                    col_letter = openpyxl.utils.get_column_letter(back_col_idx)
                    cell_ref = f"{col_letter}{r}"
                    ws.add_image(xl_back, cell_ref)

        wb.save(output_excel_path)
        return output_excel_path

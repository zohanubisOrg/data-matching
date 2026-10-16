"""Reconciler module for 3-way barcode reconciliation, gap analysis, and missing SKU reporting."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Set

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from product_mapper.zohanubis.barcode_scanner import BarcodeScanResult
from product_mapper.zohanubis.sanitizer import RawPhotoRecord


@dataclass
class CatalogProduct:
    row_idx: int
    barcode: str
    record_id: str
    name_en: str
    brand: str
    category: str
    partition: str  # G1, G2, G3, G4, Nyt
    demo_image_cell: str
    verified_back_photos: List[Path] = field(default_factory=list)
    selected_front_photo: Optional[Path] = None
    barcode_center: Optional[Tuple[float, float]] = None
    status: str = "PENDING"
    strategy: str = "STANDARD"  # STANDARD, SEQUENTIAL_PAIR, CARTON_BOX, SINGLE_SURFACE


@dataclass
class ReconciliationSummary:
    total_catalog_rows: int
    matched_count: int
    missing_count: int
    orphan_count: int
    missing_report_path: Path
    catalog_products: Dict[str, CatalogProduct]  # keyed by record_id
    orphan_photos: List[BarcodeScanResult]


class BarcodeReconciler:
    """Performs 3-way reconciliation against Master Excel catalog."""

    def __init__(self, excel_path: Path, output_dir: Path):
        self.excel_path = excel_path
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def derive_partition_from_record_id(record_id: str) -> str:
        """Map Record ID (e.g. G1-0001, G4-0012, N25-0001) to folder partition."""
        rid = record_id.strip().upper()
        if rid.startswith("G1"):
            return "G1"
        elif rid.startswith("G2"):
            return "G2"
        elif rid.startswith("G3"):
            return "G3"
        elif rid.startswith("G4"):
            return "G4"
        elif rid.startswith("N"):
            return "Nyt"
        return "OTHER"

    def load_catalog(self) -> Dict[str, CatalogProduct]:
        """Load catalog products from Excel, keeping string barcodes intact."""
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb.active
        catalog: Dict[str, CatalogProduct] = {}

        # Rows start at 4
        for r in range(4, ws.max_row + 1):
            raw_bc = ws.cell(r, 1).value
            if not raw_bc:
                continue
            barcode = str(raw_bc).strip()
            # If float e.g. 8850043126456.0, strip decimal
            if barcode.endswith(".0"):
                barcode = barcode[:-2]

            name_en = str(ws.cell(r, 2).value or "").strip()
            brand = str(ws.cell(r, 4).value or "").strip()
            record_id = str(ws.cell(r, 5).value or f"ROW-{r}").strip()
            category = str(ws.cell(r, 8).value or "").strip()
            partition = self.derive_partition_from_record_id(record_id)

            prod = CatalogProduct(
                row_idx=r,
                barcode=barcode,
                record_id=record_id,
                name_en=name_en,
                brand=brand,
                category=category,
                partition=partition,
                demo_image_cell=f"W{r}",
            )
            catalog[record_id] = prod

        return catalog

    def reconcile(
        self,
        scan_results: List[BarcodeScanResult],
        catalog: Optional[Dict[str, CatalogProduct]] = None,
        customer_front_map: Optional[Dict[str, Path]] = None,
    ) -> ReconciliationSummary:
        """Reconcile scanned barcodes with catalog, resolve sequential pairs and gap analysis."""
        if catalog is None:
            catalog = self.load_catalog()

        # Build reverse lookup by barcode
        barcode_to_prods: Dict[str, List[CatalogProduct]] = {}
        for prod in catalog.values():
            barcode_to_prods.setdefault(prod.barcode, []).append(prod)

        # Index scans by parent directory to resolve sequential shooting pairs
        folder_scans: Dict[Path, List[BarcodeScanResult]] = {}
        for scan in scan_results:
            folder_scans.setdefault(scan.path.parent, []).append(scan)

        # Sort each directory by file name to reconstruct chronological timeline
        for f_path, s_list in folder_scans.items():
            s_list.sort(key=lambda s: s.path.name.lower())

        orphan_photos: List[BarcodeScanResult] = []

        # Process barcode scan results across each folder sequence
        for f_path, s_list in folder_scans.items():
            for idx, scan in enumerate(s_list):
                if not scan.barcodes:
                    continue

                matched_any = False
                for bc in scan.barcodes:
                    matched_prods = barcode_to_prods.get(bc)
                    if not matched_prods and bc.startswith("0"):
                        matched_prods = barcode_to_prods.get(bc.lstrip("0"))
                    if not matched_prods:
                        matched_prods = barcode_to_prods.get("0" + bc)

                    if matched_prods:
                        matched_any = True
                        for p in matched_prods:
                            # Register verified back photo
                            if scan.path not in p.verified_back_photos:
                                p.verified_back_photos.append(scan.path)
                                p.status = "MATCHED_BACK"

                            # Detect adjacent non-barcode candidate for sequential pairs
                            front_candidate: Optional[Path] = None
                            if idx > 0 and len(s_list[idx - 1].barcodes) == 0:
                                front_candidate = s_list[idx - 1].path
                            elif idx < len(s_list) - 1 and len(s_list[idx + 1].barcodes) == 0:
                                front_candidate = s_list[idx + 1].path

                            # Apply Strategy based on partition and shooting paradigm
                            folder_name = f_path.name.lower()
                            is_carton = "carton" in p.name_en.lower() or "thùng" in p.name_en.lower()

                            if "nyt" in folder_name:
                                if front_candidate and scan.path.name.upper().startswith("IMG_"):
                                    p.strategy = "SEQUENTIAL_PAIR"
                                    p.selected_front_photo = front_candidate
                                    p.barcode_center = None
                                elif scan.path.name.lower() in ("tofu.jpg", "tofu2.jpg", "mychinh.jpg", "super.jpg"):
                                    p.strategy = "STANDARD"
                                    if p.selected_front_photo is None:
                                        p.selected_front_photo = scan.path
                                    if p.barcode_center is None and bc in scan.barcode_centers:
                                        p.barcode_center = scan.barcode_centers[bc]
                                else:
                                    p.strategy = "SINGLE_SURFACE"
                                    p.selected_front_photo = front_candidate or scan.path
                                    p.barcode_center = None

                            elif "g4" in folder_name:
                                if is_carton:
                                    p.strategy = "CARTON_BOX"
                                    p.selected_front_photo = front_candidate or scan.path
                                    p.barcode_center = None
                                elif front_candidate and p.record_id in ("G4-0005", "G4-0006"):
                                    p.strategy = "SEQUENTIAL_PAIR"
                                    p.selected_front_photo = front_candidate
                                    p.barcode_center = None
                                elif p.partition == "G4":
                                    # Fresh produce / bagged items (single-surface with barcode on bag)
                                    p.strategy = "SINGLE_SURFACE"
                                    p.selected_front_photo = scan.path
                                    p.barcode_center = None
                                else:
                                    # Cross-partition supplement photo for G1/G2/G3
                                    if front_candidate:
                                        p.strategy = "SEQUENTIAL_PAIR"
                                        p.selected_front_photo = front_candidate
                                        p.barcode_center = None
                                    else:
                                        p.strategy = "SINGLE_SURFACE"
                                        p.selected_front_photo = scan.path
                                        p.barcode_center = None

                            else:
                                # G1, G2, G3 standard dual-hero
                                p.strategy = "STANDARD"
                                if p.selected_front_photo is None:
                                    p.selected_front_photo = scan.path
                                if p.barcode_center is None and bc in scan.barcode_centers:
                                    p.barcode_center = scan.barcode_centers[bc]

                if not matched_any:
                    orphan_photos.append(scan)

        # Separate matched vs missing
        matched_count = sum(1 for p in catalog.values() if p.verified_back_photos)
        missing_count = sum(1 for p in catalog.values() if not p.verified_back_photos)

        # Write missing report
        report_path = self.output_dir / "Zohanubis_Missing_SKU_Report.xlsx"
        self._write_missing_report(
            report_path,
            [p for p in catalog.values() if not p.verified_back_photos],
            customer_front_map=customer_front_map,
        )

        return ReconciliationSummary(
            total_catalog_rows=len(catalog),
            matched_count=matched_count,
            missing_count=missing_count,
            orphan_count=len(orphan_photos),
            missing_report_path=report_path,
            catalog_products=catalog,
            orphan_photos=orphan_photos,
        )

    def _write_missing_report(
        self,
        report_path: Path,
        missing_products: List[CatalogProduct],
        customer_front_map: Optional[Dict[str, Path]] = None,
    ) -> None:
        """Create a styled Excel spreadsheet detailing all unphotographed or missing SKUs."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Missing SKUs"

        headers = [
            "Record ID",
            "Barcode",
            "Product Name (EN)",
            "Brand",
            "Category",
            "Source Partition",
            "Issue Reason",
            "Recommended Action",
        ]
        ws.append(headers)

        # Header styling
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(1, col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 28

        for idx, p in enumerate(missing_products, start=2):
            is_cust = False
            if customer_front_map:
                bc = p.barcode.strip()
                if bc in customer_front_map or bc.lstrip("0") in customer_front_map or ("0" + bc) in customer_front_map:
                    is_cust = True

            if is_cust:
                reason = "No raw studio photo found; Customer provided front packshot"
                action = "Front packshot ready for webshop. Customer to provide back photo if needed."
            else:
                reason = "No raw photo found with readable barcode"
                action = "Customer to provide photography (both Front and Back) for this SKU"

            ws.append([
                p.record_id,
                p.barcode,
                p.name_en,
                p.brand,
                p.category,
                p.partition,
                reason,
                action,
            ])
            ws.row_dimensions[idx].height = 20
            # format barcode cell as text
            ws.cell(idx, 2).number_format = "@"

            # alternating fill
            fill_color = "F2F5F9" if idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(idx, col_idx)
                c.fill = row_fill
                c.border = thin_border
                c.font = Font(name="Calibri", size=10)

        # Set column widths
        col_widths = {1: 14, 2: 18, 3: 45, 4: 20, 5: 18, 6: 18, 7: 35, 8: 45}
        for col_idx, width in col_widths.items():
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        wb.save(report_path)

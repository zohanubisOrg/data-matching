from pathlib import Path
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import DataQualityIssue, MatchResult
from .serializers import flat_result

BLUE = "1F4E78"
_TEXT_HINTS = ("barcode", "sku", "product id", "image")


def _sheet(wb: Workbook, name: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for index, cell in enumerate(row):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            header = str(headers[index]).casefold() if index < len(headers) else ""
            if any(hint in header for hint in _TEXT_HINTS):
                cell.number_format = "@"


def _candidate_columns(result: MatchResult) -> list[object]:
    values: list[object] = []
    for candidate in result.top_candidates[:3]:
        values.extend([
            candidate.product.product_name,
            candidate.product.product_id,
            candidate.product.sku,
            candidate.product.barcode or "",
            "; ".join(candidate.evidence),
        ])
    while len(values) < 15:
        values.append("")
    return values


def _review_candidate(candidate):
    if candidate is None:
        return ["", "", "", ""]
    return [
        candidate.product.product_name,
        candidate.product.barcode or "",
        candidate.product.product_id,
        candidate.product.sku,
    ]


def write_reports(
    results: list[MatchResult],
    issues: list[DataQualityIssue],
    profile: dict,
    output: Path,
) -> None:
    output.mkdir(parents=True, exist_ok=True)
    statuses = ["AUTO_MATCHED", "NEED_REVIEW", "UNMATCHED", "ERROR"]
    counts = {status: sum(result.status == status for result in results) for status in statuses}
    total = len(results)
    pct = lambda count: count / total if total else 0

    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_rows = [
        ["Total POS products", profile.get("pos", {}).get("rows", 0)],
        ["Total imported products", profile.get("imports", {}).get("rows", 0)],
        ["Total valid images", total],
        ["Invalid image inputs", profile.get("images", {}).get("parse_failures", 0)],
        ["Auto matched", counts["AUTO_MATCHED"]],
        ["Need review", counts["NEED_REVIEW"]],
        ["Unmatched", counts["UNMATCHED"]],
        ["Errors", counts["ERROR"]],
        ["Auto match rate", pct(counts["AUTO_MATCHED"])],
        ["Review rate", pct(counts["NEED_REVIEW"])],
        ["Unmatched rate", pct(counts["UNMATCHED"])],
        ["Measured accuracy", "NOT MEASURED - requires gold dataset/human verification"],
    ]
    _sheet(workbook, "Summary", ["Metric", "Value"], summary_rows)
    for row in range(10, 13):
        workbook["Summary"].cell(row=row, column=2).number_format = "0.00%"

    auto_headers = [
        "Image", "Source", "Source Product ID", "Parsed Name", "POS Product Name",
        "Official Name", "Product ID", "SKU", "Barcode", "Evidence Strength",
        "Match Method", "Evidence",
    ]
    auto_rows = [
        [
            result.image.filename,
            result.image.source,
            result.image.source_product_id,
            result.image.normalized_name,
            result.candidate.product.product_name if result.candidate else "",
            result.candidate.product.official_name if result.candidate else "",
            result.candidate.product.product_id if result.candidate else "",
            result.candidate.product.sku if result.candidate else "",
            result.candidate.product.barcode if result.candidate else "",
            "STRONG_DETERMINISTIC" if result.status == "AUTO_MATCHED" else "",
            result.match_method,
            "; ".join(result.evidence),
        ]
        for result in results
        if result.status == "AUTO_MATCHED"
    ]
    _sheet(workbook, "Auto Matched", auto_headers, auto_rows)

    review_headers = ["Image", "Source", "Parsed Name"] + sum(([
        f"Candidate {index} Name",
        f"Candidate {index} Product ID",
        f"Candidate {index} SKU",
        f"Candidate {index} Barcode",
        f"Candidate {index} Evidence",
    ] for index in range(1, 4)), []) + ["Reason", "Evidence"]
    review_rows = [
        [
            result.image.filename,
            result.image.source,
            result.image.normalized_name,
            *_candidate_columns(result),
            result.reason,
            "; ".join(result.evidence),
        ]
        for result in results
        if result.status == "NEED_REVIEW"
    ]
    _sheet(workbook, "Need Review", review_headers, review_rows)

    unmatched_rows = [
        [result.image.filename, result.image.source, result.image.normalized_name, result.reason]
        for result in results
        if result.status == "UNMATCHED"
    ]
    _sheet(workbook, "Unmatched", ["Image", "Source", "Parsed Name", "Reason"], unmatched_rows)

    error_rows = [
        [result.image.filename, "matching", "ERROR", result.error or result.reason]
        for result in results
        if result.status == "ERROR"
    ]
    _sheet(workbook, "Errors", ["Image", "Stage", "Error Type", "Error Message"], error_rows)
    quality_rows = [[issue.issue_type, issue.field, issue.value, issue.count, issue.details] for issue in issues]
    _sheet(workbook, "Data Quality", ["Issue", "Field", "Value", "Count", "Details"], quality_rows)
    workbook.save(output / "mapping_result.xlsx")

    flat = [flat_result(result) for result in results]
    with (output / "mapping_result.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(flat[0]) if flat else ["image_path"])
        writer.writeheader()
        writer.writerows(flat)
    (output / "mapping_result.json").write_text(
        json.dumps(flat, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    customer_headers = [
        "Image", "Detected Product", "Matched POS Product", "Barcode", "SKU",
        "Match Method", "Status", "Need Customer Confirmation", "Notes",
    ]
    customer_rows = [
        [
            result.image.filename,
            result.image.normalized_name,
            result.candidate.product.product_name if result.candidate else "",
            result.candidate.product.barcode if result.candidate else "",
            result.candidate.product.sku if result.candidate else "",
            result.match_method,
            result.status,
            "YES" if result.status in {"NEED_REVIEW", "UNMATCHED"} else "NO",
            result.reason,
        ]
        for result in results
    ]
    customer_workbook = Workbook()
    customer_workbook.remove(customer_workbook.active)
    _sheet(customer_workbook, "Customer Result", customer_headers, customer_rows)
    customer_workbook.save(output / "customer_result.xlsx")

    # Keep the first 11 columns backward-compatible so Customer Selection remains J.
    review_workbook = Workbook()
    review_workbook.remove(review_workbook.active)
    customer_review_headers = [
        "Image", "Detected Name",
        "Candidate 1", "Candidate 1 Barcode",
        "Candidate 2", "Candidate 2 Barcode",
        "Candidate 3", "Candidate 3 Barcode",
        "Suggested Candidate", "Customer Selection", "Customer Note",
        "Candidate 1 Product ID", "Candidate 1 SKU",
        "Candidate 2 Product ID", "Candidate 2 SKU",
        "Candidate 3 Product ID", "Candidate 3 SKU",
    ]
    customer_review_rows: list[list[object]] = []
    for result in results:
        if result.status != "NEED_REVIEW":
            continue
        candidates = list(result.top_candidates[:3]) + [None] * 3
        first = _review_candidate(candidates[0])
        second = _review_candidate(candidates[1])
        third = _review_candidate(candidates[2])
        customer_review_rows.append([
            result.image.filename,
            result.image.normalized_name,
            first[0], first[1],
            second[0], second[1],
            third[0], third[1],
            "",  # No preselected suggestion: deterministic candidates are not probabilities.
            "",
            "",
            first[2], first[3],
            second[2], second[3],
            third[2], third[3],
        ])
    _sheet(review_workbook, "Review", customer_review_headers, customer_review_rows)
    if customer_review_rows:
        validation = DataValidation(type="list", formula1='"1,2,3,No Match"', allow_blank=True)
        review_workbook["Review"].add_data_validation(validation)
        validation.add(f"J2:J{len(customer_review_rows) + 1}")
    review_workbook.save(output / "customer_review.xlsx")

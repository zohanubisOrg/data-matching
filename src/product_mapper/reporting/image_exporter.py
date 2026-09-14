from __future__ import annotations

import csv
import logging
from pathlib import Path
import re
import shutil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps

from ..models import MatchResult

logger = logging.getLogger("product_mapper")

HEADER_FILL_BLUE = "1F4E78"
_SAFE_FILENAME_BARCODE = re.compile(r"^[A-Za-z0-9._-]+$")


def _format_worksheet(ws, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL_BLUE)
        cell.alignment = Alignment(wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(60, max(12, max(len(str(x.value or "")) for x in column) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "@"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _copy_image(source: Path, targets: list[Path], copy_failures: set[str]) -> bool:
    try:
        for target in targets:
            shutil.copy2(source, target)
        return True
    except Exception as exc:
        copy_failures.add(str(source))
        logger.error("Failed to copy image %s: %s", source, exc)
        return False


def _to_rgb_for_jpeg(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    if "A" in image.getbands():
        rgba = image.convert("RGBA")
        background = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        background.alpha_composite(rgba)
        return background.convert("RGB")
    return image.convert("RGB")


def _export_matched_jpeg(source: Path, targets: list[Path], copy_failures: set[str]) -> bool:
    """Write a canonical JPEG for website/import delivery regardless of source format."""
    try:
        if source.suffix.casefold() in {".jpg", ".jpeg"}:
            for target in targets:
                shutil.copy2(source, target)
            return True

        with Image.open(source) as image:
            rgb = _to_rgb_for_jpeg(image)
            first = targets[0]
            rgb.save(first, format="JPEG", quality=95, subsampling=0, optimize=True)
            for target in targets[1:]:
                shutil.copy2(first, target)
        return True
    except Exception as exc:
        copy_failures.add(str(source))
        logger.error("Failed to export canonical JPEG %s: %s", source, exc)
        return False


def _unique_unresolved_name(filename: str, used: set[str]) -> str:
    """Return a flat-directory-safe unique name without dropping any input image."""
    original = Path(filename).name or "image"
    candidate = original
    stem = Path(original).stem or "image"
    suffix = Path(original).suffix
    counter = 1
    while candidate.casefold() in used:
        counter += 1
        candidate = f"{stem}_{counter}{suffix}"
    used.add(candidate.casefold())
    return candidate


def _append_error_rows(
    result: MatchResult,
    deliverable_rows: list[dict[str, str]],
    mapping_rows: list[dict[str, str]],
) -> None:
    deliverable_rows.append({
        "Barcode": "",
        "SKU": "",
        "Original Name": result.image.raw_name,
        "Final Name": result.image.normalized_name,
        "Original Image": result.image.filename,
        "Final Image": "",
        "Status": "ERROR",
    })
    mapping_rows.append({
        "barcode": "",
        "pos_sku": "",
        "product_id": "",
        "source_name": result.image.raw_name,
        "normalized_name": result.image.normalized_name,
        "supplier": result.image.source or "",
        "supplier_sku": result.image.source_product_id or "",
        "original_image": result.image.filename,
        "final_image": "",
        "match_method": result.match_method or "error",
        "match_score": "0.0000",
        "status": "INVALID",
    })


def export_standard_deliverables(
    results: list[MatchResult],
    output_dir: Path,
    total_pos_count: int = 0,
    invalid_input_count: int = 0,
) -> dict:
    """Export client deliverables and account for both valid and invalid image inputs."""
    matched_dir = output_dir / "matched-images"
    images_dir = output_dir / "images"
    unmatched_dir = output_dir / "unmatched-images"
    review_dir = output_dir / "review"
    unmatched_sub_dir = output_dir / "unmatched"

    for directory in (matched_dir, images_dir, unmatched_dir, review_dir, unmatched_sub_dir):
        directory.mkdir(parents=True, exist_ok=True)
        for old_file in directory.glob("*.*"):
            try:
                old_file.unlink()
            except Exception as exc:
                logger.warning("Could not remove stale output file %s: %s", old_file, exc)

    invalid_input_count = max(0, int(invalid_input_count))
    barcode_counts: dict[str, int] = {}
    image_to_barcode: dict[Path, str] = {}
    used_unresolved_names: set[str] = set()
    conflict_count = 0
    copy_failures: set[str] = set()

    deliverable_rows: list[dict[str, str]] = []
    mapping_rows: list[dict[str, str]] = []

    for result in results:
        if result.status == "AUTO_MATCHED" and result.candidate and result.candidate.product.barcode:
            barcode = result.candidate.product.barcode.strip()
            original_image = result.image.path.resolve()
            if original_image in image_to_barcode and image_to_barcode[original_image] != barcode:
                logger.warning(
                    "Conflict detected: Image %s claimed by barcode %s and %s. Downgrading to REVIEW.",
                    result.image.filename,
                    image_to_barcode[original_image],
                    barcode,
                )
                result.status = "NEED_REVIEW"
                result.reason = (
                    f"Conflict: image assigned to multiple barcodes "
                    f"({image_to_barcode[original_image]} and {barcode})"
                )
                conflict_count += 1
            else:
                image_to_barcode[original_image] = barcode

    matched_count = 0
    review_count = 0
    unmatched_count = 0
    error_count = 0

    for result in results:
        status = result.status
        original_image_name = result.image.filename
        source_path = result.image.path
        supplier_name = result.image.source or ""
        supplier_sku = result.image.source_product_id or ""

        if status == "AUTO_MATCHED" and result.candidate and result.candidate.product.barcode:
            barcode = result.candidate.product.barcode.strip()
            sku = result.candidate.product.sku.strip()
            product_id = result.candidate.product.product_id.strip()
            original_name = result.candidate.product.product_name.strip()
            final_name = (result.candidate.product.official_name or result.candidate.product.product_name).strip()

            if not _SAFE_FILENAME_BARCODE.fullmatch(barcode):
                result.status = "ERROR"
                result.reason = "matched POS barcode is unsafe as canonical image filename"
                unresolved_name = _unique_unresolved_name(original_image_name, used_unresolved_names)
                _copy_image(source_path, [unmatched_dir / unresolved_name], copy_failures)
                _append_error_rows(result, deliverable_rows, mapping_rows)
                error_count += 1
                continue

            index = barcode_counts.get(barcode, 0) + 1
            final_image_name = f"{barcode}.jpg" if index == 1 else f"{barcode}_{index}.jpg"
            exported = _export_matched_jpeg(
                source_path,
                [matched_dir / final_image_name, images_dir / final_image_name],
                copy_failures,
            )
            if not exported:
                result.status = "ERROR"
                result.reason = "canonical matched image export failed"
                unresolved_name = _unique_unresolved_name(original_image_name, used_unresolved_names)
                _copy_image(source_path, [unmatched_dir / unresolved_name], copy_failures)
                _append_error_rows(result, deliverable_rows, mapping_rows)
                error_count += 1
                continue

            barcode_counts[barcode] = index
            deliverable_rows.append({
                "Barcode": barcode,
                "SKU": sku,
                "Original Name": original_name,
                "Final Name": final_name,
                "Original Image": original_image_name,
                "Final Image": final_image_name,
                "Status": "MATCHED",
            })
            mapping_rows.append({
                "barcode": barcode,
                "pos_sku": sku,
                "product_id": product_id,
                "source_name": original_name,
                "normalized_name": final_name,
                "supplier": supplier_name,
                "supplier_sku": supplier_sku,
                "original_image": original_image_name,
                "final_image": final_image_name,
                "match_method": result.match_method,
                "match_score": f"{result.confidence:.4f}",
                "status": "MATCHED",
            })
            matched_count += 1

        elif status == "NEED_REVIEW":
            barcode = result.candidate.product.barcode.strip() if (result.candidate and result.candidate.product.barcode) else ""
            sku = result.candidate.product.sku.strip() if (result.candidate and result.candidate.product.sku) else ""
            product_id = result.candidate.product.product_id.strip() if result.candidate else ""
            original_name = result.candidate.product.product_name.strip() if result.candidate else result.image.raw_name
            final_name = (
                (result.candidate.product.official_name or result.candidate.product.product_name).strip()
                if result.candidate
                else result.image.normalized_name
            )
            unresolved_name = _unique_unresolved_name(original_image_name, used_unresolved_names)
            _copy_image(
                source_path,
                [unmatched_dir / unresolved_name, review_dir / unresolved_name],
                copy_failures,
            )

            deliverable_rows.append({
                "Barcode": barcode,
                "SKU": sku,
                "Original Name": original_name,
                "Final Name": final_name,
                "Original Image": original_image_name,
                "Final Image": "",
                "Status": "REVIEW",
            })
            mapping_rows.append({
                "barcode": barcode,
                "pos_sku": sku,
                "product_id": product_id,
                "source_name": original_name,
                "normalized_name": final_name,
                "supplier": supplier_name,
                "supplier_sku": supplier_sku,
                "original_image": original_image_name,
                "final_image": "",
                "match_method": result.match_method,
                "match_score": f"{result.confidence:.4f}",
                "status": "REVIEW_REQUIRED",
            })
            review_count += 1

        elif status == "UNMATCHED":
            unresolved_name = _unique_unresolved_name(original_image_name, used_unresolved_names)
            _copy_image(
                source_path,
                [unmatched_dir / unresolved_name, unmatched_sub_dir / unresolved_name],
                copy_failures,
            )

            deliverable_rows.append({
                "Barcode": "",
                "SKU": "",
                "Original Name": result.image.raw_name or result.image.normalized_name,
                "Final Name": result.image.normalized_name,
                "Original Image": original_image_name,
                "Final Image": "",
                "Status": "UNMATCHED",
            })
            mapping_rows.append({
                "barcode": "",
                "pos_sku": "",
                "product_id": "",
                "source_name": result.image.raw_name or result.image.normalized_name,
                "normalized_name": result.image.normalized_name,
                "supplier": supplier_name,
                "supplier_sku": supplier_sku,
                "original_image": original_image_name,
                "final_image": "",
                "match_method": "none",
                "match_score": f"{result.confidence:.4f}",
                "status": "UNMATCHED",
            })
            unmatched_count += 1

        else:
            unresolved_name = _unique_unresolved_name(original_image_name, used_unresolved_names)
            _copy_image(source_path, [unmatched_dir / unresolved_name], copy_failures)
            _append_error_rows(result, deliverable_rows, mapping_rows)
            error_count += 1

    headers_7 = ["Barcode", "SKU", "Original Name", "Final Name", "Original Image", "Final Image", "Status"]
    table_rows_7 = [[row[header] for header in headers_7] for row in deliverable_rows]
    result_workbook = Workbook()
    result_sheet = result_workbook.active
    result_sheet.title = "Result"
    _format_worksheet(result_sheet, headers_7, table_rows_7)
    try:
        result_workbook.save(output_dir / "result.xlsx")
    except PermissionError:
        logger.warning("Permission denied writing result.xlsx (file open in Excel). Saved to result_updated.xlsx")
        result_workbook.save(output_dir / "result_updated.xlsx")

    with (output_dir / "result.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers_7)
        writer.writeheader()
        writer.writerows(deliverable_rows)

    headers_12 = [
        "barcode", "pos_sku", "product_id", "source_name", "normalized_name",
        "supplier", "supplier_sku", "original_image", "final_image",
        "match_method", "match_score", "status",
    ]
    table_rows_12 = [[row[header] for header in headers_12] for row in mapping_rows]
    mapping_workbook = Workbook()
    mapping_sheet = mapping_workbook.active
    mapping_sheet.title = "Mapping"
    _format_worksheet(mapping_sheet, headers_12, table_rows_12)
    try:
        mapping_workbook.save(output_dir / "mapping.xlsx")
    except PermissionError:
        logger.warning("Permission denied writing mapping.xlsx (file open in Excel). Saved to mapping_updated.xlsx")
        mapping_workbook.save(output_dir / "mapping_updated.xlsx")

    with (output_dir / "mapping.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers_12)
        writer.writeheader()
        writer.writerows(mapping_rows)

    valid_images = len(results)
    total_input_images = valid_images + invalid_input_count
    matched_pos_identities = {
        (row["product_id"], row["pos_sku"], row["barcode"])
        for row in mapping_rows
        if row["status"] == "MATCHED"
    }
    percent = lambda count: (count / total_input_images * 100.0) if total_input_images else 0.0
    reconciled = matched_count + review_count + unmatched_count + error_count + invalid_input_count
    reconciliation_state = "PASS" if reconciled == total_input_images else "FAIL"
    conflict_state = "PASS" if conflict_count == 0 else "REVIEW REQUIRED"
    copy_state = "PASS" if not copy_failures else "FAIL"

    recon_report = f"""================================================================================
FINAL DATA MATCHING & RECONCILIATION REPORT
================================================================================

1. IMAGES ACCOUNTING
--------------------------------------------------------------------------------
Total Input Images Discovered:      {total_input_images}
Valid Images Processed:             {valid_images}
Invalid input files:                {invalid_input_count}
  - MATCHED (in /matched-images/):  {matched_count} ({percent(matched_count):.1f}%)
  - REVIEW (in /unmatched-images/): {review_count} ({percent(review_count):.1f}%)
  - UNMATCHED:                      {unmatched_count} ({percent(unmatched_count):.1f}%)
  - PROCESSING/EXPORT ERROR:        {error_count} ({percent(error_count):.1f}%)
  ------------------------------------------------------------------------------
  Reconciled Check:                 {reconciled} / {total_input_images} [{reconciliation_state}]

2. POS PRODUCTS ACCOUNTING
--------------------------------------------------------------------------------
Total POS Products in System:       {total_pos_count}
  - Unique matched POS identities:  {len(matched_pos_identities)}
  - POS rows without matched image: {total_pos_count - len(matched_pos_identities) if total_pos_count else 'N/A'}
  ------------------------------------------------------------------------------
  Barcode & SKU handling:           Preserved from POS records; never synthesized.

3. IMAGE & QUALITY SAFEGUARDS
--------------------------------------------------------------------------------
Duplicate image conflicts:          {conflict_count} [{conflict_state}]
Output copy/conversion failures:    {len(copy_failures)} [{copy_state}]
Output naming policy:               <barcode>.jpg or <barcode>_N.jpg
Gold-dataset accuracy:              NOT MEASURED
Final Matched Accuracy:             NOT MEASURED (requires gold dataset or human verification)

The strict matcher is designed to maximize precision by sending uncertain evidence to
REVIEW/UNMATCHED. A row is only delivered as MATCHED when its canonical image export
succeeds. Invalid inputs remain accounted for through Data Quality reporting.
================================================================================
"""
    (output_dir / "reconciliation_report.txt").write_text(recon_report, encoding="utf-8")

    logger.info("Deliverables exported:\n%s", recon_report)
    return {
        "total_images": valid_images,
        "total_input_images": total_input_images,
        "invalid_input_count": invalid_input_count,
        "matched_count": matched_count,
        "review_count": review_count,
        "unmatched_count": unmatched_count,
        "error_count": error_count,
        "conflict_count": conflict_count,
        "copy_failure_count": len(copy_failures),
        "matched_pos_count": len(matched_pos_identities),
    }

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


def profile_workbooks(pos_path: Path, imports_path: Path, images: list[Path], pos_meta: dict, import_meta: dict, image_meta: dict, issues: list) -> dict:
    profile = {
        "inputs": {"pos": str(pos_path), "imports": str(imports_path), "images": [str(x) for x in images]},
        "pos": pos_meta, "imports": import_meta, "images": image_meta,
        "data_quality": [x.model_dump() for x in issues],
    }
    return profile


def write_profile(profile: dict, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    (output / "data_profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Profile"
    ws.append(["Section", "Key", "Value"])
    for section, values in profile.items():
        if isinstance(values, dict):
            for key, value in values.items(): ws.append([section, key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        else: ws.append([section, "", json.dumps(values, ensure_ascii=False)])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 100
    wb.save(output / "data_profile.xlsx")


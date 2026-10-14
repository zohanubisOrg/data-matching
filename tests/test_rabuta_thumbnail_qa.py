import sys
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts import qa_row_by_row as qa


class Embedded:
    def _data(self):
        im = Image.new('RGB', (115,115), 'white')
        ImageDraw.Draw(im).rectangle((25,10,90,105), fill='green')
        buf = BytesIO()
        im.save(buf, 'PNG')
        return buf.getvalue()


def test_clean_thumbnail_never_grants_full_resolution_pass():
    passed, flags = qa.inspect_col24_image(Embedded())
    assert passed is False
    assert any('NOT_VERIFIED' in x for x in flags)


def test_barcode_on_authentic_front_is_an_observation_not_automatic_defect(monkeypatch):
    class Hit:
        text = '12345678'
    monkeypatch.setattr(qa.zxingcpp, 'read_barcodes', lambda *a, **kw: [Hit()])
    passed, flags = qa.inspect_col24_image(Embedded())
    assert not passed
    assert any('Barcode observed' in x for x in flags)
    assert not any('Barcode detected on webshop front' in x for x in flags)


def test_product_with_record_id_but_no_barcode_or_images_is_not_a_blank_row(tmp_path,monkeypatch):
    import openpyxl
    path=tmp_path/'catalog.xlsx'
    wb=openpyxl.Workbook()
    wb.active.cell(4,2,'Product needing source')
    wb.active.cell(4,5,'N240924-003')
    wb.save(path)
    monkeypatch.setattr(qa,'EXCEL_PATH',path)
    monkeypatch.setattr(qa,'CHUNKS_DIR',tmp_path/'screening')
    result=qa.run_row_qa(4,4)
    assert result['results'][0]['status']=='BLOCKED_NEEDS_SOURCE'
    assert result['passed_rows']==0

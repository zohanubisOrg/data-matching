from product_mapper.models import POSProduct, SupplierProduct
from product_mapper.matching.no_fuzzy import (
    match_import_names_no_fuzzy,
    summarize_no_fuzzy_matches,
)


def p(pid, name, official="", printer=""):
    return POSProduct(
        product_id=str(pid), sku=f"SKU-{pid}", barcode=f"BC-{pid}",
        product_name=name, official_name=official, printer_name=printer,
    )


def s(name, supplier="EUR", sheet="EUR"):
    return SupplierProduct(name=name, supplier=supplier, sheet=sheet)


def test_exact_name_is_auto_matched_after_basic_normalization():
    rows = match_import_names_no_fuzzy(
        [p(1, "Aroy-D Coconut Milk 500 ml")],
        [s("AROY D Coconut Milk 500ml")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["match_method"] == "canonical_exact"
    assert rows[0]["pos_product_id"] == "1"


def test_logistics_noise_and_pack_count_can_match_deterministically():
    rows = match_import_names_no_fuzzy(
        [p(1, "Basmati Rice 5 kg")],
        [s("Basmati Rice 1 X 5 KG AKASH")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["match_method"] in {"core_exact", "safe_subset"}


def test_multilingual_product_family_alias_is_deterministic_not_fuzzy():
    rows = match_import_names_no_fuzzy(
        [p(1, "Glass Noodles 250 g")],
        [s("THANH LOC Glass Vermicelli 30X250 G")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["match_method"] in {"core_exact", "safe_subset"}


def test_unique_generic_pos_name_can_match_longer_import_name_when_family_is_preserved():
    rows = match_import_names_no_fuzzy(
        [p(1, "Aloe Vera Drink")],
        [s("Aloe Vera Drink Lychee 20 X 500 ML OKF")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["match_method"] == "safe_subset"


def test_duplicate_same_canonical_pos_name_can_auto_match_but_identity_stays_ambiguous():
    rows = match_import_names_no_fuzzy(
        [p(1, "Sesame Oil 500ml"), p(2, "Sesame Oil 500 ml")],
        [s("Sesame Oil 500 ml")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["matched_pos_rows"] == 2
    assert rows[0]["pos_identity_ambiguous"] is True
    assert rows[0]["pos_product_id"] == ""


def test_two_distinct_deterministic_candidates_go_to_review_not_auto():
    rows = match_import_names_no_fuzzy(
        [p(1, "Rice Noodles"), p(2, "Rice Vermicelli")],
        [s("Rice Noodles Vermicelli 400 g")],
    )
    assert rows[0]["status"] == "REVIEW"
    assert rows[0]["reason"] == "multiple_deterministic_candidates"


def test_size_conflict_blocks_auto_match():
    rows = match_import_names_no_fuzzy(
        [p(1, "Basmati Rice 1 kg")],
        [s("Basmati Rice 10 kg")],
    )
    assert rows[0]["status"] != "AUTO_MATCHED"
    assert rows[0]["size_conflict"] is True


def test_short_query_does_not_expand_to_broader_product():
    rows = match_import_names_no_fuzzy(
        [p(1, "Chili Mango")],
        [s("MANGO")],
    )
    assert rows[0]["status"] == "UNMATCHED"


def test_processed_product_cannot_collapse_to_ingredient():
    rows = match_import_names_no_fuzzy(
        [p(1, "Nata de Coco")],
        [s("TIKTAK Watermelon Juice With Nata De Coco 24x320ml")],
    )
    assert rows[0]["status"] != "AUTO_MATCHED"


def test_composite_food_cannot_collapse_to_ingredient():
    rows = match_import_names_no_fuzzy(
        [p(1, "Black Bean")],
        [s("Fried Dace with Salted Black Bean 24 X 184 G EAGLE COIN")],
    )
    assert rows[0]["status"] != "AUTO_MATCHED"


def test_unrelated_family_never_auto_matches():
    rows = match_import_names_no_fuzzy(
        [p(1, "Rice 5 kg")],
        [s("Korean Pear 6 PCS 1 X 5 KG - Origin:CN")],
    )
    assert rows[0]["status"] == "UNMATCHED"


def test_summary_reports_auto_review_unmatched_and_target_flags():
    summary = summarize_no_fuzzy_matches([
        {"status": "AUTO_MATCHED"},
        {"status": "AUTO_MATCHED"},
        {"status": "REVIEW"},
        {"status": "UNMATCHED"},
    ])
    assert summary["total_import_rows"] == 4
    assert summary["auto_matched"] == 2
    assert summary["review"] == 1
    assert summary["unmatched"] == 1
    assert summary["auto_match_rate_pct"] == 50.0
    assert summary["review_rate_pct"] == 25.0
    assert summary["candidate_coverage_pct"] == 75.0
    assert summary["target_auto_ge_80"] is False
    assert summary["target_review_lt_20"] is False


def test_equivalent_units_match_without_a_score():
    rows = match_import_names_no_fuzzy(
        [p(1, "Coconut Water 1 L")],
        [s("Coconut Water 12 X 1000 ML BAMBOO TREE")],
    )
    assert rows[0]["status"] == "AUTO_MATCHED"
    assert rows[0]["item_size"] == "1000 ml"
    assert "score" not in rows[0]


def test_observed_typo_and_compound_aliases_are_exact_not_fuzzy():
    rows = match_import_names_no_fuzzy(
        [p(1, "Turmeric Powder"), p(2, "Lemongrass")],
        [s("Tumeric powder 10X400 G"), s("Lemon Grass 1 X 2 KG")],
    )
    assert [row["status"] for row in rows] == ["AUTO_MATCHED", "AUTO_MATCHED"]


def test_no_fuzzy_dependency_is_used_by_deterministic_matcher():
    import inspect
    import product_mapper.matching.no_fuzzy as module

    assert "rapidfuzz" not in inspect.getsource(module).casefold()


def test_write_outputs_keeps_summary_and_matches_auditable(tmp_path):
    from openpyxl import load_workbook
    from product_mapper.matching.no_fuzzy import write_no_fuzzy_outputs

    rows = match_import_names_no_fuzzy([p(1, "Sesame Oil 500 ml")], [s("Sesame Oil 500ml")])
    paths = write_no_fuzzy_outputs(rows, summarize_no_fuzzy_matches(rows), tmp_path)
    workbook = load_workbook(paths["xlsx"], read_only=True)
    assert workbook.sheetnames == ["summary", "matches"]
    assert paths["csv"].exists()
    assert paths["summary_json"].exists()

from pathlib import Path

from openpyxl import load_workbook

from product_mapper.models import POSProduct, ProductImage, MatchCandidate, MatchResult
from product_mapper.reporting.excel_report import write_reports
from product_mapper.reporting.html_report import write_html


def test_excel_and_html_export(tmp_path: Path):
    image = ProductImage(path=tmp_path / "image.jpg", filename="image.jpg", source="x", source_product_id="1", normalized_name="beans")
    p = POSProduct(product_id="1", sku="S1", barcode="0201000000179", product_name="Beans", official_name="Beans", printer_name="Beans")
    candidate = MatchCandidate(product=p, score=.95)
    result = MatchResult(image=image, status="NEED_REVIEW", candidate=candidate, top_candidates=[candidate], confidence=.95)
    write_reports([result], [], {"pos": {"rows": 1}, "imports": {"rows": 1}, "images": {"directories": {}}}, tmp_path)
    write_html([result], tmp_path)
    wb = load_workbook(tmp_path / "mapping_result.xlsx", read_only=True)
    assert wb.sheetnames == ["Summary", "Auto Matched", "Need Review", "Unmatched", "Errors", "Data Quality"]
    review = load_workbook(tmp_path / "customer_review.xlsx")
    assert review["Review"]["J2"].value is None
    assert (tmp_path / "review_report.html").exists()


def test_export_standard_deliverables(tmp_path: Path):
    from product_mapper.reporting.image_exporter import export_standard_deliverables

    img1_path = tmp_path / "test1.jpg"
    img1_path.write_bytes(b"fake image data 1")
    img2_path = tmp_path / "test2.jpg"
    img2_path.write_bytes(b"fake image data 2")

    img1 = ProductImage(path=img1_path, filename="test1.jpg", source="heuschen", source_product_id="101", normalized_name="soy sauce")
    img2 = ProductImage(path=img2_path, filename="test2.jpg", source="asropa", source_product_id="202", normalized_name="unknown item")

    p1 = POSProduct(product_id="1001", sku="SKU-1001", barcode="8935001234567", product_name="Soy Sauce 500ml", official_name="Soy Sauce 500ml")
    cand1 = MatchCandidate(product=p1, score=0.98)

    r1 = MatchResult(image=img1, status="AUTO_MATCHED", candidate=cand1, top_candidates=[cand1], confidence=0.98)
    r2 = MatchResult(image=img2, status="UNMATCHED", confidence=0.2)

    out_dir = tmp_path / "output"
    summary = export_standard_deliverables([r1, r2], out_dir, total_pos_count=50)

    assert (out_dir / "matched-images" / "8935001234567.jpg").exists()
    assert (out_dir / "unmatched-images" / "test2.jpg").exists()
    assert (out_dir / "result.xlsx").exists()
    assert (out_dir / "result.csv").exists()
    assert (out_dir / "reconciliation_report.txt").exists()
    assert summary["matched_count"] == 1
    assert summary["unmatched_count"] == 1



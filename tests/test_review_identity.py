import json

from openpyxl import load_workbook

from product_mapper.matching.knowledge_base import (
    import_review_decisions,
    import_review_json_decisions,
    resolve_approved_mapping,
)
from product_mapper.models import MatchCandidate, MatchResult, POSProduct, ProductImage
from product_mapper.reporting.excel_report import write_reports


def product(pid: str, sku: str, barcode: str, name: str) -> POSProduct:
    return POSProduct(
        product_id=pid,
        sku=sku,
        barcode=barcode,
        product_name=name,
        official_name=name,
    )


def test_customer_review_persists_product_id_sku_and_barcode(tmp_path):
    image_path = tmp_path / "product.jpg"
    image_path.write_bytes(b"image")
    image = ProductImage(path=image_path, filename="product.jpg", normalized_name="sesame oil")
    first = product("1", "SKU-A", "4006381333931", "Sesame Oil A")
    second = product("2", "SKU-B", "4006381333931", "Sesame Oil B")
    result = MatchResult(
        image=image,
        status="NEED_REVIEW",
        top_candidates=[
            MatchCandidate(product=first, score=1.0, evidence=["duplicate barcode candidate"]),
            MatchCandidate(product=second, score=1.0, evidence=["duplicate barcode candidate"]),
        ],
        reason="DUPLICATE_OR_EQUIVALENT_POS_BARCODE",
    )

    write_reports(
        [result],
        [],
        {"pos": {"rows": 2}, "imports": {"rows": 0}, "images": {"directories": {}}},
        tmp_path,
    )

    review_path = tmp_path / "customer_review.xlsx"
    workbook = load_workbook(review_path)
    sheet = workbook["Review"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=headers["Customer Selection"], value="2")
    workbook.save(review_path)

    kb_path = tmp_path / "knowledge_base.json"
    assert import_review_decisions(review_path, kb_path) == 1
    data = json.loads(kb_path.read_text(encoding="utf-8"))

    assert data["product.jpg"]["product_id"] == "2"
    assert data["product.jpg"]["sku"] == "SKU-B"
    assert data["product.jpg"]["barcode"] == "4006381333931"


def test_customer_review_has_one_decision_row_per_front_back_capture(tmp_path):
    candidate_product = product("1", "SKU-A", "4006381333931", "Sesame Oil")
    candidate = MatchCandidate(product=candidate_product, score=1.0)
    front = ProductImage(
        path=tmp_path / "000123_front.jpg",
        filename="000123_front.jpg",
        capture_id="000123",
        image_side="front",
    )
    back = ProductImage(
        path=tmp_path / "000123_back.jpg",
        filename="000123_back.jpg",
        capture_id="000123",
        image_side="back",
    )
    results = [
        MatchResult(image=front, status="NEED_REVIEW", top_candidates=[candidate], reason="review"),
        MatchResult(image=back, status="NEED_REVIEW", top_candidates=[candidate], reason="review"),
    ]

    write_reports(
        results,
        [],
        {"pos": {"rows": 1}, "imports": {"rows": 0}, "images": {"directories": {}}},
        tmp_path,
    )

    workbook = load_workbook(tmp_path / "customer_review.xlsx", read_only=True)
    sheet = workbook["Review"]
    assert sheet.max_row == 2  # header + one physical-product decision


def test_html_review_json_can_merge_into_knowledge_base(tmp_path):
    review_json = tmp_path / "reviewed_knowledge.json"
    review_json.write_text(json.dumps({
        "product.jpg": {
            "selection": "1",
            "product_id": "2",
            "sku": "SKU-B",
            "barcode": "4006381333931",
            "source": "review_report_html",
        }
    }), encoding="utf-8")
    kb_path = tmp_path / "knowledge_base.json"

    assert import_review_json_decisions(review_json, kb_path) == 1
    saved = json.loads(kb_path.read_text(encoding="utf-8"))
    assert saved["product.jpg"]["product_id"] == "2"
    assert saved["product.jpg"]["source"] == "review_report_html"


def test_human_approved_tuple_resolves_one_row_even_when_barcode_is_duplicated():
    first = product("1", "SKU-A", "4006381333931", "Sesame Oil A")
    second = product("2", "SKU-B", "4006381333931", "Sesame Oil B")
    entry = {"product_id": "2", "sku": "SKU-B", "barcode": "4006381333931"}

    hits = resolve_approved_mapping(entry, [first, second])

    assert hits == [second]


def test_legacy_barcode_only_human_mapping_remains_ambiguous_when_pos_barcode_is_duplicated():
    first = product("1", "SKU-A", "4006381333931", "Sesame Oil A")
    second = product("2", "SKU-B", "4006381333931", "Sesame Oil B")

    hits = resolve_approved_mapping({"barcode": "4006381333931"}, [first, second])

    assert hits == [first, second]

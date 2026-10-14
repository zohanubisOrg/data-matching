import sys
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.qa_audit_delivery import audit_delivery


def books(tmp_path, alter_w=False, alter_data=False):
    source_image = tmp_path/'source.png'
    Image.new('RGB',(100,100),'green').save(source_image)
    altered = tmp_path/'altered.png'
    Image.new('RGB',(100,100),'red').save(altered)
    source, delivery = tmp_path/'source.xlsx', tmp_path/'delivery.xlsx'
    for path in (source,delivery):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(3,1,'Barcode')
        ws.cell(4,1,'0123456789')
        ws.cell(4,5,'R-1')
        ws.cell(4,24,'ingredients')
        if path == delivery:
            ws.insert_cols(24,2)
            if alter_data: ws.cell(4,26,'wrong ingredients')
        ws.add_image(XLImage(altered if path == delivery and alter_w else source_image),'W4')
        if path == delivery:
            ws.add_image(XLImage(source_image),'X4')
            ws.add_image(XLImage(source_image),'Y4')
        wb.save(path)
    front, back = tmp_path/'front',tmp_path/'back'
    front.mkdir(); back.mkdir()
    Image.new('RGB',(2400,2400),'white').save(front/'0123456789-front.png')
    return source,delivery,front,back


def test_structural_success_cannot_certify_images_from_clean_borders(tmp_path):
    result = audit_delivery(*books(tmp_path))
    assert result['status'] == 'NOT_VERIFIED'
    assert result['qa_pass'] is False
    assert result['w_preserved'] is True


def test_any_w_image_replacement_blocks_delivery(tmp_path):
    result = audit_delivery(*books(tmp_path, alter_w=True))
    assert result['status'] == 'BLOCKED'
    assert result['w_preserved'] is False


def test_data_shift_mismatch_outside_old_sample_rows_is_detected(tmp_path):
    result = audit_delivery(*books(tmp_path, alter_data=True))
    assert result['status'] == 'BLOCKED'
    assert result['cell_mismatch_count'] == 1

"""Structural delivery audit, never an automatic visual acceptance certificate."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SRC_EXCEL = ROOT/'source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx'
DELIVERY_EXCEL = ROOT/'output/catalog_pairs/Rabuta_master_final_Delivery_Embedded.xlsx'
FRONT_V3_DIR = ROOT/'output/catalog_pairs/tagged/front-v3'
BACK_V2_DIR = ROOT/'output/catalog_pairs/tagged/back-v2'


def image_index(ws):
    result = {}
    for image in ws._images:
        cell = getattr(getattr(image, 'anchor', None), '_from', None)
        if cell is not None:
            result.setdefault((cell.row+1,cell.col+1), []).append(hashlib.sha256(image._data()).hexdigest())
    return {k: sorted(v) for k,v in result.items()}


def audit_delivery(source, delivery, front_dir, back_dir):
    src = openpyxl.load_workbook(source, data_only=True)
    dest = openpyxl.load_workbook(delivery, data_only=True)
    a,b = src.active,dest.active
    shape_ok = a.max_row == b.max_row and b.max_column == a.max_column+2
    mismatches = []
    mismatch_count = 0
    for row in range(1,a.max_row+1):
        for col in range(1,a.max_column+1):
            dest_col = col if col <= 23 else col+2
            if a.cell(row,col).value != b.cell(row,dest_col).value:
                mismatch_count += 1
                if len(mismatches) < 100:
                    mismatches.append({'row':row,'source_col':col,'delivery_col':dest_col})
    source_images, delivery_images = image_index(a),image_index(b)
    original_w = {k:v for k,v in source_images.items() if k[1] == 23}
    delivered_w = {k:v for k,v in delivery_images.items() if k[1] == 23}
    w_preserved = original_w == delivered_w
    missing_pair_rows = [r for r in range(4,a.max_row+1)
                         if (r,24) not in delivery_images or (r,25) not in delivery_images]
    collisions = {}
    for row in range(4,b.max_row+1):
        bc,rid = str(b.cell(row,1).value or '').strip(),str(b.cell(row,5).value or '').strip()
        if bc: collisions.setdefault(bc,set()).add(rid)
    collisions = {bc:sorted(ids) for bc,ids in collisions.items() if len(ids)>1}
    file_checks = []
    for path in sorted(Path(front_dir).glob('*')):
        if path.suffix.lower() not in ('.jpg','.jpeg','.png'): continue
        try:
            with Image.open(path) as image:
                arr = np.asarray(image.convert('RGB'))
                white_border = bool((arr[0]==255).all() and (arr[-1]==255).all()
                                    and (arr[:,0]==255).all() and (arr[:,-1]==255).all())
                file_checks.append({'path':str(path.resolve()),'size':list(image.size),
                                    'white_outer_border':white_border,
                                    'nonwhite_pixels':int(np.count_nonzero(np.any(arr!=255,axis=2))),
                                    'qa_status':'NOT_VERIFIED'})
        except (OSError,ValueError) as exc:
            file_checks.append({'path':str(path.resolve()),'error':str(exc),'qa_status':'NEEDS_MANUAL_REVIEW'})
    status = 'BLOCKED' if (not shape_ok or mismatch_count or not w_preserved or missing_pair_rows or collisions) else 'NOT_VERIFIED'
    src.close(); dest.close()
    return {'status':status,'qa_pass':False,'shape_ok':shape_ok,'cell_mismatch_count':mismatch_count,
            'cell_mismatch_examples':mismatches,'w_preserved':w_preserved,'w_image_count':len(original_w),
            'missing_pair_rows':missing_pair_rows,'barcode_identity_conflicts':collisions,
            'front_files':file_checks,'back_file_count':sum(1 for p in Path(back_dir).glob('*') if p.is_file()),
            'limitations':['Outer white borders do not prove background removal, full packaging, correct face or identity.',
                           'Native source and output must be visually compared; no automatic catalog PASS.']}


def run_qa():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report',type=Path,default=ROOT/'output/catalog_pairs/qa_delivery_structural_v2.json')
    args = parser.parse_args()
    result = audit_delivery(SRC_EXCEL,DELIVERY_EXCEL,FRONT_V3_DIR,BACK_V2_DIR)
    args.report.parent.mkdir(parents=True,exist_ok=True)
    args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:result[k] for k in ('status','qa_pass','shape_ok','cell_mismatch_count','w_preserved')},ensure_ascii=True))
    return 2


if __name__ == '__main__':
    raise SystemExit(run_qa())

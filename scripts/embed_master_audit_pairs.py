#!/usr/bin/env python3
"""
scripts/embed_master_audit_pairs.py
------------------------------------
Nhúng 2 cột ảnh webshop (mặt trước v3 và mặt sau v2) ngay cạnh cột ảnh mặt trước gốc của khách
trong file: Rabuta_master_SKU_categories_ingredients_final.xlsx

Cấu trúc 3 cột đối chiếu song song:
- Cột 23 (W): 'Ảnh mặt trước' (Ảnh gốc khách hàng - giữ nguyên 100%)
- Cột 24 (X): 'Ảnh webshop mặt trước' (Nhúng từ output/catalog_pairs/tagged/front-v3)
- Cột 25 (Y): 'Ảnh mặt sau' (Nhúng từ output/catalog_pairs/tagged/back-v2)
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path
from typing import Optional

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx"
BACKUP_SOURCE = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.original_backup.xlsx"

FRONT_DIR = ROOT / "output/catalog_pairs/tagged/front-v3"
BACK_DIR = ROOT / "output/catalog_pairs/tagged/back-v2"

OUTPUT_DELIVERY = ROOT / "output/catalog_pairs/Rabuta_master_final_Delivery_Embedded.xlsx"
OUTPUT_STANDALONE = ROOT / "output/catalog_pairs/Rabuta_master_SKU_categories_ingredients_final_audit.xlsx"
THUMBS_DIR = ROOT / "output/catalog_pairs/thumbnails_cache"
THUMBS_DIR.mkdir(parents=True, exist_ok=True)


def find_image(folder: Path, barcode: str, suffix: str) -> Optional[Path]:
    """Tìm file ảnh theo barcode với nhiều biến thể định dạng (dấu gạch nối, bỏ số 0 đầu, thêm số 0 đầu)."""
    for sep in ("-", "_"):
        cand = folder / f"{barcode}{sep}{suffix}.jpg"
        if cand.exists():
            return cand

    # Nếu có số 0 ở đầu (UPC/EAN)
    if barcode.startswith("0"):
        stripped = barcode.lstrip("0")
        for sep in ("-", "_"):
            cand = folder / f"{stripped}{sep}{suffix}.jpg"
            if cand.exists():
                return cand

    # Nếu mã 12 số thiếu số 0 đầu thành 13 số EAN
    if len(barcode) == 12:
        with_zero = "0" + barcode
        for sep in ("-", "_"):
            cand = folder / f"{with_zero}{sep}{suffix}.jpg"
            if cand.exists():
                return cand

    return None


def create_excel_thumbnail(img_path: Path, max_size: int = 115) -> Path:
    """Tạo thumbnail tỷ lệ chuẩn để nhúng vào ô Excel mà không làm phình dung lượng file."""
    thumb_path = THUMBS_DIR / f"thumb_{img_path.name}"
    
    # Kiểm tra cache và so sánh thời gian sửa đổi (mtime)
    if thumb_path.exists() and thumb_path.stat().st_size > 0:
        if thumb_path.stat().st_mtime >= img_path.stat().st_mtime:
            return thumb_path

    try:
        with PILImage.open(img_path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_size, max_size), PILImage.Resampling.LANCZOS)
            im.save(thumb_path, "JPEG", quality=90, optimize=True)
            return thumb_path
    except Exception:
        return img_path


def embed_audit_columns():
    print("=" * 80)
    print("  NHÚNG 2 CỘT ĐỐI CHIẾU: [ẢNH WEBSHOP MẶT TRƯỚC] VÀ [ẢNH MẶT SAU] VÀO EXCEL")
    print("=" * 80)

    if not EXCEL_SOURCE.exists():
        print(f"[-] Lỗi: Không tìm thấy file nguồn: {EXCEL_SOURCE}")
        return

    # 1. Tạo backup an toàn nếu chưa có
    if not BACKUP_SOURCE.exists():
        print(f"--> Đang tạo bản sao lưu an toàn: {BACKUP_SOURCE.name}...")
        shutil.copy2(EXCEL_SOURCE, BACKUP_SOURCE)
        print("    [+] Đã sao lưu bản gốc thành công.")
    else:
        print(f"    [i] Bản sao lưu đã tồn tại: {BACKUP_SOURCE.name}")

    print(f"--> Đang tải file Excel nguồn: {EXCEL_SOURCE}...")
    wb = openpyxl.load_workbook(EXCEL_SOURCE)
    ws = wb.active

    print(f"    * Kích thước ban đầu: {ws.max_row} dòng, {ws.max_column} cột.")
    print(f"    * Số ảnh gốc của khách hàng trong file: {len(ws._images)} ảnh.")

    # 2. Chèn 2 cột mới liền kề bên phải Cột 23 ('Ảnh mặt trước')
    print("--> Đang chèn 2 cột mới tại vị trí 24 (Col X) và vị trí 25 (Col Y)...")
    ws.insert_cols(24, 2)

    # 3. Định dạng tiêu đề cho 2 cột mới tại Row 3
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c24 = ws.cell(3, 24)
    c24.value = "Ảnh webshop mặt trước"
    c24.font = header_font
    c24.fill = header_fill
    c24.alignment = align_center

    c25 = ws.cell(3, 25)
    c25.value = "Ảnh mặt sau"
    c25.font = header_font
    c25.fill = header_fill
    c25.alignment = align_center

    # Đặt độ rộng chuẩn cho 3 cột ảnh đối chiếu
    ws.column_dimensions["W"].width = 18  # Ảnh mặt trước (khách)
    ws.column_dimensions["X"].width = 18  # Ảnh webshop mặt trước (front-v3)
    ws.column_dimensions["Y"].width = 18  # Ảnh mặt sau (back-v2)

    front_count = 0
    back_count = 0
    total_products = 0

    print("--> Bắt đầu quét barcode và nhúng ảnh thumbnail vào từng hàng...")

    for r in range(4, ws.max_row + 1):
        bc_val = ws.cell(r, 1).value
        if not bc_val:
            continue

        barcode = str(bc_val).strip()
        total_products += 1

        # Chuẩn hóa chiều cao hàng để ảnh hiển thị rõ ràng, không bị chèn ép
        current_h = ws.row_dimensions[r].height
        if current_h is None or current_h < 95:
            ws.row_dimensions[r].height = 95

        # Nhúng Ảnh webshop mặt trước (Cột 24 / Col X)
        front_path = find_image(FRONT_DIR, barcode, "front")
        if front_path:
            try:
                t_front = create_excel_thumbnail(front_path, max_size=115)
                img_front = OpenpyxlImage(str(t_front))
                ws.add_image(img_front, f"X{r}")
                front_count += 1
            except Exception as e:
                print(f"    [!] Lỗi nhúng front row {r} ({barcode}): {e}")

        # Nhúng Ảnh mặt sau (Cột 25 / Col Y)
        back_path = find_image(BACK_DIR, barcode, "back")
        if back_path:
            try:
                t_back = create_excel_thumbnail(back_path, max_size=115)
                img_back = OpenpyxlImage(str(t_back))
                ws.add_image(img_back, f"Y{r}")
                back_count += 1
            except Exception as e:
                print(f"    [!] Lỗi nhúng back row {r} ({barcode}): {e}")

        if total_products % 200 == 0 or r == ws.max_row:
            print(f"    [Tiến độ] Đã duyệt {total_products}/{ws.max_row - 3} SKU | Đã nhúng: Front={front_count}, Back={back_count}")

    # 4. Lưu ra file đầu ra chính và copy sang các đích khác
    print(f"\n--> Đang lưu workbook vào: {OUTPUT_DELIVERY}...")
    ws.title = "Master SKU Delivery"
    wb.save(OUTPUT_DELIVERY)
    print("    [+] Đã lưu hoàn tất OUTPUT_DELIVERY.")

    print(f"--> Đang sao chép sang bản độc lập: {OUTPUT_STANDALONE}...")
    shutil.copy2(OUTPUT_DELIVERY, OUTPUT_STANDALONE)
    print("    [+] Đã tạo bản OUTPUT_STANDALONE thành công.")

    # Cập nhật trực tiếp vào file nguồn của khách (do đã có bản backup an toàn)
    try:
        print(f"--> Đang cập nhật trực tiếp vào file gốc: {EXCEL_SOURCE}...")
        shutil.copy2(OUTPUT_DELIVERY, EXCEL_SOURCE)
        print("    [+] Đã cập nhật trực tiếp vào file gốc thành công.")
    except Exception as e:
        print(f"    [!] Không thể ghi đè file gốc (có thể đang mở trong Excel): {e}")

    print("\n" + "=" * 80)
    print("                 HOÀN TẤT NHÚNG 3 CỘT ĐỐI CHIẾU CHUẨN XÁC")
    print("=" * 80)
    print(f"  * Tổng số sản phẩm trong bảng : {total_products}")
    print(f"  * Cột 23 (W) 'Ảnh mặt trước'          : 100% giữ nguyên ảnh gốc của khách")
    print(f"  * Cột 24 (X) 'Ảnh webshop mặt trước'  : Đã nhúng {front_count} ảnh từ front-v3")
    print(f"  * Cột 25 (Y) 'Ảnh mặt sau'            : Đã nhúng {back_count} ảnh từ back-v2")
    print(f"  * File sẵn sàng mở kiểm tra   : {OUTPUT_STANDALONE}")
    print(f"  * File bàn giao hoàn thiện    : {OUTPUT_DELIVERY}")
    print("=" * 80 + "\n")


if __name__ == "__main__":
    embed_audit_columns()

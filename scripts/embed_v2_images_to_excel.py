#!/usr/bin/env python3
"""
scripts/embed_v2_images_to_excel.py
------------------------------------
Nhúng trực tiếp ảnh Mặt trước (front-v2) và Mặt sau (back-v2)
vào file Excel chuẩn Rabuta.12.10.2026_original.xlsx theo yêu cầu của khách hàng.
"""

from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/12-10-new-requirment/Rabuta.12.10.2026_original.xlsx"
FRONT_DIR = ROOT / "output/catalog_pairs/tagged/front-v2"
BACK_DIR = ROOT / "output/catalog_pairs/tagged/back-v2"
OUTPUT_EXCEL = ROOT / "output/catalog_pairs/Rabuta_12_10_Delivery_Embedded_v2.xlsx"
THUMBS_DIR = ROOT / "output/catalog_pairs/thumbnails_cache"
THUMBS_DIR.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)


def create_excel_thumbnail(img_path: Path, max_size: int = 120) -> Path:
    """Tạo thumbnail tỷ lệ chuẩn để nhúng vào ô Excel mà không làm phình dung lượng file."""
    thumb_path = THUMBS_DIR / f"thumb_{img_path.name}"
    if thumb_path.exists():
        return thumb_path

    try:
        with PILImage.open(img_path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_size, max_size), PILImage.Resampling.LANCZOS)
            im.save(thumb_path, "JPEG", quality=90, optimize=True)
            return thumb_path
    except Exception:
        return img_path


def embed_catalog_images():
    print("===============================================================================")
    print("        NHÚNG ẢNH MẶT TRƯỚC (V2) VÀ MẶT SAU (V2) VÀO EXCEL RABUTA 12.10        ")
    print("===============================================================================\n")

    if not EXCEL_SOURCE.exists():
        print(f"[-] Lỗi: Không tìm thấy file nguồn: {EXCEL_SOURCE}")
        return

    # 1. Đọc dữ liệu từ file gốc
    print("--> Đang đọc file gốc...")
    src_wb = openpyxl.load_workbook(EXCEL_SOURCE, read_only=True, data_only=True)
    src_sheet = src_wb.active
    rows = list(src_sheet.iter_rows(values_only=True))
    src_wb.close()

    print(f"    Tổng cộng: {len(rows)} dòng dữ liệu.")

    # 2. Tạo workbook mới
    dest_wb = openpyxl.Workbook()
    dest_sheet = dest_wb.active
    dest_sheet.title = "Rabuta Catalog v2"
    dest_sheet.views.sheetView[0].showGridLines = True

    # Sao chép 2 dòng tiêu đề đầu
    if len(rows) > 0 and rows[0]:
        dest_sheet.append(list(rows[0]))
    if len(rows) > 1 and rows[1]:
        dest_sheet.append(list(rows[1]))

    # Hàng tiêu đề cột chính thức (Header ở Row 3)
    orig_headers = list(rows[2]) if len(rows) > 2 else []
    new_headers = list(orig_headers)
    new_headers.insert(1, "Ảnh Mặt Trước (Front Web v2)")
    new_headers.insert(2, "Ảnh Mặt Sau (Back Label v2)")
    dest_sheet.append(new_headers)

    header_row_idx = 3
    dest_sheet.row_dimensions[header_row_idx].height = 28
    for col_idx in range(1, len(new_headers) + 1):
        cell = dest_sheet.cell(header_row_idx, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    dest_sheet.column_dimensions["B"].width = 18  # Front
    dest_sheet.column_dimensions["C"].width = 18  # Back
    dest_sheet.column_dimensions["A"].width = 16  # Barcode
    dest_sheet.column_dimensions["D"].width = 35  # Full name (EN)

    front_count = 0
    back_count = 0

    print("--> Đang nhúng ảnh vào từng dòng sản phẩm...")
    current_out_row = 4

    for r in rows[3:]:
        if not r or not r[0]:
            continue

        barcode = str(r[0]).strip()
        row_vals = list(r)
        row_vals.insert(1, "")
        row_vals.insert(2, "")

        dest_sheet.append(row_vals)
        dest_sheet.row_dimensions[current_out_row].height = 95

        for c_idx in range(1, len(row_vals) + 1):
            c_cell = dest_sheet.cell(current_out_row, c_idx)
            c_cell.font = REGULAR_FONT
            c_cell.border = CELL_BORDER
            c_cell.alignment = Alignment(vertical="center", wrap_text=True)

        # Nhúng Mặt Trước v2
        front_p = FRONT_DIR / f"{barcode}-front.jpg"
        if not front_p.exists():
            front_p = FRONT_DIR / f"{barcode}_front.jpg"

        if front_p.exists():
            try:
                t_p = create_excel_thumbnail(front_p, max_size=115)
                img_front = OpenpyxlImage(str(t_p))
                dest_sheet.add_image(img_front, f"B{current_out_row}")
                front_count += 1
            except Exception:
                pass

        # Nhúng Mặt Sau v2
        back_p = BACK_DIR / f"{barcode}-back.jpg"
        if not back_p.exists():
            back_p = BACK_DIR / f"{barcode}_back.jpg"

        if back_p.exists():
            try:
                t_p = create_excel_thumbnail(back_p, max_size=115)
                img_back = OpenpyxlImage(str(t_p))
                dest_sheet.add_image(img_back, f"C{current_out_row}")
                back_count += 1
            except Exception:
                pass

        current_out_row += 1

    print("--> Đang lưu file Excel...")
    dest_wb.save(OUTPUT_EXCEL)
    print(f"\n=== HOÀN TẤT NHÚNG ẢNH V2 ===")
    print(f"  * File kết quả đã lưu tại : {OUTPUT_EXCEL}")
    print(f"  * Số ảnh Mặt Trước đã nhúng: {front_count}")
    print(f"  * Số ảnh Mặt Sau đã nhúng  : {back_count}")


if __name__ == "__main__":
    embed_catalog_images()

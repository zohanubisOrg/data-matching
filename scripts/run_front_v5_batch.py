#!/usr/bin/env python3
"""Index original photos, render conservative front-v5 packshots, write the ledger."""
from __future__ import annotations

import hashlib
import argparse
import html
import json
import os
import re
import sys
from urllib.parse import quote
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import openpyxl


if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import pillow_heif
from PIL import Image

pillow_heif.register_heif_opener()

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from product_mapper.vision.barcode_validator import get_barcode_lookup_variants
from product_mapper.vision.front_v5_ledger import (
    build_run_fingerprint,
    can_reuse_result,
    accepted_set_digest,
    first_batch_digest,
    iter_source_files,
    load_apptance_records,
    load_catalog_rows,
    load_column_media,
    normalize_barcode_variants,
)
from product_mapper.vision.front_v5_workers import index_one, render_one
from product_mapper.vision.front_v5_pipeline import (
    _read_pilot_approval,
    promote_candidate,
    review_with_repairs,
)

WORKBOOK = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx"
SOURCE_DIRS = [
    ROOT / "source-import/G1_original",
    ROOT / "source-import/G2_original",
    ROOT / "source-import/G3_Original",
    ROOT / "source-import/G4-original",
    ROOT / "source-import/Nyt25.09.26",
]
RUN_DIR = ROOT / "output/catalog_pairs/front-v5-run"
FRONT_DIR = ROOT / "output/catalog_pairs/tagged/front-v5"
APPTANCE_DIR = ROOT / "output/catalog_pairs/tagged/apptance"
WORKERS = 8
RUN_STAGE = "render"
ALLOW_RESUME = False
QA_MODE = "shadow"
PILOT_APPROVAL_PATH: Path | None = None
PILOT_LABELS_PATH: Path | None = None
SHARD_INDEX: int | None = None
RUN_SCOPE = "pilot"
BUDGET_ID: str | None = None


def _load_confirmed_pilot(path: Path | None) -> dict[str, dict]:
    if path is None or not path.is_file():
        raise SystemExit("render requires --labels with the 48 human-confirmed pilot cases")
    labels = _read_jsonl(path)
    by_barcode = {str(item.get("barcode", "")): item for item in labels}
    if len(by_barcode) != 48 or "" in by_barcode or any(
        item.get("confirmed") is not True
        or not isinstance(item.get("barcode"), str)
        or not item["barcode"].isdigit()
        or not 8 <= len(item["barcode"]) <= 14
        for item in labels
    ):
        raise SystemExit("pilot labels must contain exactly 48 unique, human-confirmed barcodes")
    groups = {"box", "bottle_jar", "sealed_bag", "transparent_bag", "produce_tray", "reflective_frozen"}
    if any(item.get("group") not in groups for item in labels):
        raise SystemExit("every pilot label needs one of the six package groups")
    positives = [item for item in labels if item.get("case_type") == "positive" and item.get("expected_verdict") == "PASS"]
    regressions = [item for item in labels if item.get("case_type") == "regression" and item.get("expected_verdict") in {"FAIL", "REVIEW"}]
    difficult = [item for item in labels if item.get("case_type") == "difficult" and item.get("group") in groups]
    if (len(positives), len(regressions), len(difficult)) != (12, 12, 24):
        raise SystemExit("pilot labels need 12 positives, 12 regression cases, and 24 difficult cases")
    if any(
        item.get("expected_verdict") not in {"PASS", "REVIEW", "FAIL"}
        or ("expected_reason_codes" in item and (
            not isinstance(item["expected_reason_codes"], list)
            or any(not isinstance(code, str) or not code for code in item["expected_reason_codes"])
        ))
        for item in labels
    ):
        raise SystemExit("pilot labels contain an invalid expected verdict or reason-code list")
    if {item["group"] for item in difficult} != groups or any(
        sum(item.get("group") == group for item in difficult) < 4 for group in groups
    ):
        raise SystemExit("the 24 difficult labels must cover all six package groups, at least four each")
    return by_barcode


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def _write_jsonl_atomic(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + f".{os.getpid()}.tmp")
    _write_jsonl(temporary, rows)
    os.replace(temporary, path)


def _read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _shard_name(index: int, total: int) -> str:
    if index < 150:
        return "a"
    span = max(1, (total - 150 + 2) // 3)
    offset = index - 150
    if offset < span:
        return "b"
    if offset < span * 2:
        return "c"
    return "d"


def run_render() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(f"missing workbook: {WORKBOOK}")
    workbook_hash = _sha256(WORKBOOK)
    pilot_labels = _load_confirmed_pilot(PILOT_LABELS_PATH) if RUN_STAGE == "render" else {}
    labels_sha256 = _sha256(PILOT_LABELS_PATH) if PILOT_LABELS_PATH is not None else None
    if RUN_DIR.exists() and not ALLOW_RESUME:
        raise SystemExit(f"run already exists; pass --resume to validate and reuse current evidence: {RUN_DIR}")
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    (RUN_DIR / "workbook.sha256").write_text(workbook_hash, encoding="utf-8")
    implementation = {
        str(path.relative_to(ROOT)): _sha256(path)
        for path in (
            ROOT / "src/product_mapper/vision/conservative_front.py",
            ROOT / "src/product_mapper/vision/front_v5_quality.py",
            ROOT / "src/product_mapper/vision/front_v5_workers.py",
            ROOT / "src/product_mapper/vision/front_v5_ledger.py",
            ROOT / "src/product_mapper/vision/front_v5_gemini.py",
            ROOT / "src/product_mapper/vision/front_v5_pipeline.py",
            ROOT / "scripts/run_front_v5_batch.py",
            ROOT / "scripts/run_front_shard_0.py",
            ROOT / "scripts/run_front_shard_1.py",
            ROOT / "scripts/run_front_shard_2.py",
        )
    }
    run_config = {
        "quality": "Q1-Q6-gemini-v1", "canvas": 2400,
        "product_long_side": 2064, "gemini": "gemini-3.1-pro-preview",
    }
    index_path = RUN_DIR / "index.jsonl"
    indexed = []
    legacy_index = ROOT / "output/catalog_pairs/front-v5-run/index.jsonl"
    if index_path.exists():
        indexed = [json.loads(line) for line in index_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    elif legacy_index.exists():
        indexed = [json.loads(line) for line in legacy_index.read_text(encoding="utf-8").splitlines() if line.strip()]
    files = [str(Path(path).resolve()) for path in iter_source_files(SOURCE_DIRS)]
    actual_inventory = {str(Path(path).resolve()): _sha256(Path(path)) for path in files}
    inventory_digest = hashlib.sha256(
        json.dumps(actual_inventory, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    run_fingerprint = build_run_fingerprint(
        workbook_hash, implementation,
        run_config | {"source_inventory_sha256": inventory_digest, "pilot_labels_sha256": labels_sha256},
    )
    if QA_MODE == "enforce" and not _read_pilot_approval(
        PILOT_APPROVAL_PATH, run_fingerprint, labels_sha256
    ):
        raise SystemExit("enforce mode requires a passing, hash-bound 48-case pilot benchmark approval")
    (RUN_DIR / "run_fingerprint.txt").write_text(run_fingerprint, encoding="utf-8")
    indexed_inventory = {str(Path(item.get("path")).resolve()): item.get("sha256") for item in indexed if item.get("path")}
    if indexed_inventory != actual_inventory:
        print(f"indexing {len(files)} originals with {WORKERS} workers", flush=True)
        indexed = []
        with ProcessPoolExecutor(max_workers=WORKERS) as pool:
            futures = [pool.submit(index_one, path) for path in files]
            for done, future in enumerate(as_completed(futures), start=1):
                indexed.append(future.result())
                if done % 100 == 0 or done == len(files):
                    print(f"  indexed {done}/{len(files)}", flush=True)
        _write_jsonl(index_path, indexed)
    else:
        print(f"reusing verified index {len(indexed)}", flush=True)
        if not index_path.is_file():
            _write_jsonl(index_path, indexed)
    if RUN_STAGE == "index":
        print(f"indexed {len(indexed)} source files -> {index_path}", flush=True)
        return
    rows = load_catalog_rows(WORKBOOK)
    rows_by_barcode = {row["barcode"]: row for row in rows}
    missing = sorted(set(pilot_labels) - rows_by_barcode.keys())
    if missing:
        raise SystemExit(f"pilot labels contain barcodes absent from the workbook: {missing[:8]}")
    rows = [rows_by_barcode[barcode] for barcode in pilot_labels] if RUN_SCOPE == "pilot" else list(rows_by_barcode.values())
    if SHARD_INDEX is not None:
        rows = [row for row in rows if row["barcode"].isdigit() and int(row["barcode"]) % 3 == SHARD_INDEX]
    if len(rows) == 0:
        raise SystemExit("this shard has no barcodes in the confirmed pilot")
    shared_qa_state = ROOT / "output/catalog_pairs/front-v5-pilot-state" / labels_sha256
    budget_id = BUDGET_ID or labels_sha256
    (RUN_DIR / "pilot-labels.sha256").write_text(labels_sha256, encoding="utf-8")
    print(f"{RUN_SCOPE} rows {len(rows)} conflicts {sum(row['identity_conflict'] for row in rows)}", flush=True)
    by_variant: dict[str, list[dict]] = {}
    for record in indexed:
        if not record["barcode"]:
            continue
        for variant in get_barcode_lookup_variants(record["barcode"]):
            by_variant.setdefault(variant, []).append(record)
    path_to_records = {str(Path(r["path"]).resolve()).lower(): r for r in indexed}
    manifest_path = ROOT / "output/catalog_pairs/manifest.json"
    plan_path = ROOT / "output/catalog_pairs/front_v3_master_plan.json"
    if manifest_path.is_file():
        manifest_data = json.loads(manifest_path.read_text(encoding="utf-8"))
        for bc, entry in manifest_data.items():
            src = entry.get("source")
            if src:
                norm_src = str(Path(src).resolve()).lower()
                if norm_src in path_to_records:
                    rec = dict(path_to_records[norm_src])
                    rec["barcode"] = bc
                    for variant in get_barcode_lookup_variants(bc):
                        existing = by_variant.setdefault(variant, [])
                        if not any(x["path"] == rec["path"] for x in existing):
                            existing.append(rec)
    if plan_path.is_file():
        plan_data = json.loads(plan_path.read_text(encoding="utf-8"))
        for bc, entry in plan_data.items():
            for src in entry.get("candidates", []):
                norm_src = str(Path(src).resolve()).lower()
                if norm_src in path_to_records:
                    rec = dict(path_to_records[norm_src])
                    rec["barcode"] = bc
                    for variant in get_barcode_lookup_variants(bc):
                        existing = by_variant.setdefault(variant, [])
                        if not any(x["path"] == rec["path"] for x in existing):
                            existing.append(rec)
    customer_fronts = load_column_media(WORKBOOK, "22")
    customer_backs = load_column_media(WORKBOOK, "24")
    print(f"customer fronts {len(customer_fronts)} backs {len(customer_backs)}", flush=True)
    apptance_records = load_apptance_records(APPTANCE_DIR, FRONT_DIR)
    print(f"pre-approved apptance items indexed: {len(apptance_records)}", flush=True)
    if any(variant in apptance_records for barcode in pilot_labels for variant in normalize_barcode_variants(barcode)):
        raise SystemExit("pilot labels include immutable apptance images; select other cases for visual QA")
    jobs = []
    job_rows: dict[str, list[dict]] = {}
    skipped_apptance = 0
    for row in rows:
        if row["identity_conflict"]:
            continue
        variants = normalize_barcode_variants(row["barcode"])
        if any(variant in apptance_records for variant in variants):
            skipped_apptance += 1
            continue
        candidates = []
        seen = set()
        for variant in variants:
            for record in by_variant.get(variant, []):
                if record["path"] in seen:
                    continue
                seen.add(record["path"])
                candidates.append(record)
        candidates.sort(key=lambda record: max(record["oriented_size"]), reverse=True)
        if not candidates:
            continue
        if row["barcode"] in job_rows:
            job_rows[row["barcode"]].append(row)
            continue
        destination = str(FRONT_DIR / f"{row['barcode']}-front.jpg")
        job = {
                "barcode": row["barcode"],
            "record_id": row["record_id"],
            "name": row["name"],
            "pilot_label": pilot_labels.get(row["barcode"]),
            "identity_conflict": row["identity_conflict"],
            "run_fingerprint": run_fingerprint,
                "quality_evidence_path": str(RUN_DIR / "qa" / f"{row['barcode']}.json"),
                "destination": destination,
                "review_destination": str(RUN_DIR / "candidates" / row["barcode"] / "candidate.jpg"),
                "run_dir": str(RUN_DIR),
                "qa_state_dir": str(shared_qa_state / "budgets" / budget_id),
                "qa_cache_dir": str(shared_qa_state / "cache"),
                "pilot_labels_sha256": labels_sha256,
                "max_render_attempts": 1,
                "candidates": candidates[:3],
                "reference": customer_fronts.get(row["barcode"]),
                "back": customer_backs.get(row["barcode"]),
            }
        job["expected_destination_sha256"] = _sha256(Path(destination)) if Path(destination).is_file() else None
        job["input_hashes"] = {
            "front_reference": None if job["reference"] is None else hashlib.sha256(job["reference"]).hexdigest(),
            "back_reference": None if job["back"] is None else hashlib.sha256(job["back"]).hexdigest(),
            "sources": [list(item) for item in sorted((candidate["path"], candidate["sha256"]) for candidate in job["candidates"])],
        }
        evidence_file = Path(job["quality_evidence_path"])
        if evidence_file.is_file():
            qa_data = json.loads(evidence_file.read_text(encoding="utf-8"))
            qa_masks = qa_data.get("mask_files", {})
            mask_hashes = []
            for mask_name, relative_path in sorted(qa_masks.items()):
                mask_path = (evidence_file.parent / str(relative_path)).resolve()
                if evidence_file.parent.resolve() not in mask_path.parents or not mask_path.is_file():
                    raise SystemExit(f"invalid QA mask path for {row['barcode']}: {relative_path}")
                mask_hashes.append([mask_name, str(relative_path), _sha256(mask_path)])
            job["input_hashes"]["qa_manifest"] = _sha256(evidence_file)
            job["input_hashes"]["qa_masks"] = mask_hashes
        else:
            job["input_hashes"]["qa_manifest"] = None
            job["input_hashes"]["qa_masks"] = []
        candidate_bytes = json.dumps(job["input_hashes"]["sources"], separators=(",", ":")).encode("utf-8")
        job["candidate_set_sha256"] = hashlib.sha256(candidate_bytes).hexdigest()
        job["run_fingerprint"] = run_fingerprint
        jobs.append(job)
        job_rows[row["barcode"]] = [row]
    missing_pilot_cases = sorted(set(row["barcode"] for row in rows) - set(job_rows))
    if RUN_SCOPE == "pilot" and missing_pilot_cases:
        raise SystemExit(f"pilot cases must have a unique identity and at least one matched original: {missing_pilot_cases[:8]}")
    progress_path = RUN_DIR / "progress.jsonl"
    rendered: dict[str, dict] = {}
    if progress_path.exists():
        for line in progress_path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            item = json.loads(line)
            rendered[item["barcode"]] = item
    apptance_seeded = 0
    for row in rows:
        if row["identity_conflict"]:
            continue
        for variant in normalize_barcode_variants(row["barcode"]):
            if variant in apptance_records:
                hit = apptance_records[variant]
                source = Path(hit["source_path"])
                rendered[row["barcode"]] = {
                    "state": "SKIPPED_APPTANCE",
                    "reason": "customer-accepted input preserved and skipped",
                    "barcode": row["barcode"],
                    "source_path": str(source),
                    "output_path": None,
                    "output_sha256": hit.get("source_sha256"),
                }
                apptance_seeded += 1
                break
    print(f"pre-approved catalog rows from apptance: {apptance_seeded} (skipped from pipeline jobs: {skipped_apptance})", flush=True)
    first_batch_barcodes = {row["barcode"] for row in rows[:150]}
    pending = []
    for job in jobs:
        expected = {
            "run_fingerprint": run_fingerprint,
            "input_hashes": job["input_hashes"],
            "candidate_set_sha256": job["candidate_set_sha256"],
            "expected_destination_sha256": job["expected_destination_sha256"],
        }
        prior_result = rendered.get(job["barcode"], {})
        if can_reuse_result(prior_result, expected):
            if QA_MODE in {"enforce", "direct"} and prior_result.get("state") == "PASS":
                try:
                    prior_result = promote_candidate(
                        job, prior_result, run_dir=RUN_DIR, front_dir=FRONT_DIR,
                        apptance_dir=APPTANCE_DIR,
                        expected_sha256=job["expected_destination_sha256"],
                        approval_path=PILOT_APPROVAL_PATH or "direct",
                    )
                except (OSError, RuntimeError, ValueError) as error:
                    prior_result = prior_result | {"state": "REVIEW", "reason": f"promotion conflict: {error}"}
            rendered[job["barcode"]] = prior_result
        else:
            pending.append(job)
    print(f"rendering {len(pending)} barcodes of {len(rows)} catalog rows, {len(rendered)} already recorded", flush=True)
    with progress_path.open("a", encoding="utf-8") as progress:
        inflight: dict = {}
        next_job = 0
        done = 0
        with ProcessPoolExecutor(max_workers=WORKERS) as pool:
            while next_job < len(pending) or inflight:
                while next_job < len(pending) and len(inflight) < WORKERS:
                    job = pending[next_job]
                    inflight[pool.submit(render_one, job)] = job
                    next_job += 1
                if not inflight:
                    break
                finished = next(as_completed(inflight))
                job = inflight.pop(finished)
                barcode = job["barcode"]
                try:
                    result = finished.result()
                    result = review_with_repairs(
                        job, result, run_dir=RUN_DIR, qa_mode=QA_MODE,
                        front_dir=FRONT_DIR, apptance_dir=APPTANCE_DIR,
                        approval_path=PILOT_APPROVAL_PATH or ("direct" if QA_MODE == "direct" else None),
                    )
                except Exception as error:
                    result = {
                        "state": "BLOCKED_INCOMPLETE_EDGE",
                        "reason": str(error),
                        "source_path": None,
                        "output_path": None,
                    }
                result["barcode"] = barcode
                rendered[barcode] = result
                progress.write(json.dumps(result, ensure_ascii=False) + "\n")
                progress.flush()
                done += 1
                if done % 25 == 0 or done == len(pending):
                    accepted = sum(item["state"] == "ACCEPTED" for item in rendered.values())
                    print(f"  rendered {done}/{len(pending)} accepted {accepted}", flush=True)
    ledger = []
    for row in rows:
        if row["identity_conflict"]:
            state, reason, source, output = "BLOCKED_IDENTITY", "barcode shared by two record ids", None, None
        elif row["barcode"] not in rendered:
            state, reason, source, output = "CUSTOMER_RESHOOT", "no original photo matched this barcode", None, None
        else:
            result = rendered[row["barcode"]]
            state, reason = result["state"], result["reason"]
            source, output = result.get("source_path"), result.get("output_path")
        result = rendered.get(row["barcode"], {})
        ledger.append(
            {
                "excel_row": row["excel_row"],
                "barcode": row["barcode"],
                "record_id": row["record_id"],
                "name": row["name"],
                "state": state,
                "reason": reason,
                "source_path": source,
                "output_path": output,
            } | {
                key: result[key]
                for key in ("output_sha256", "run_fingerprint", "gates", "candidate_path", "qa_path", "attempt_history", "input_hashes", "candidate_set_sha256", "expected_destination_sha256", "pilot_labels_sha256")
                if key in result
            }
        )
    _write_jsonl(RUN_DIR / "ledger.jsonl", ledger)
    first_batch_manifest = {
        "run_fingerprint": run_fingerprint,
        "rows_sha256": first_batch_digest(ledger),
        "reviewed_row_count": min(150, len(ledger)),
        "rows": ledger[:150],
    }
    (RUN_DIR / "first_batch_manifest.json").write_text(
        json.dumps(first_batch_manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    shards: dict[str, list[dict]] = {"a": [], "b": [], "c": [], "d": []}
    for index, row in enumerate(ledger):
        shards[_shard_name(index, len(ledger))].append(row)
    for name, shard_rows in shards.items():
        _write_jsonl(RUN_DIR / f"shard-{name}.jsonl", shard_rows)
    review_paths = [
        item.get("candidate_path") or item.get("output_path")
        for barcode, item in rendered.items() if barcode in first_batch_barcodes
        and (item.get("candidate_path") or item.get("output_path"))
    ]
    _contact_sheet(review_paths)
    _write_reshoot(ledger)
    _write_no_photo(ledger)
    _write_qa_gallery(RUN_DIR, ledger)
    _write_improvement_queue(RUN_DIR, rendered)
    accepted = {
        row["excel_row"]: Path(row["output_path"])
        for row in ledger if row["state"] == "ACCEPTED" and row.get("output_path")
    }
    counts: dict[str, int] = {}
    for row in ledger:
        counts[row["state"]] = counts.get(row["state"], 0) + 1
    print("ledger", len(ledger), counts, f"{RUN_SCOPE} run complete; embed requires a separate reviewed release", flush=True)
    if _sha256(WORKBOOK) != workbook_hash:
        raise SystemExit("customer workbook changed during the run")
    _reconcile(ledger)
    report = _pilot_report(ledger, rendered, pilot_labels, labels_sha256, run_fingerprint)
    (RUN_DIR / "pilot-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_reshoot(rows: list[dict]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Chup lai"
    sheet.append(["Excel row", "Barcode", "Record ID", "Product name", "Reason"])
    for row in rows:
        if row["reason"] == "curated selection requires current Q1-Q6 evidence":
            continue
        if row["state"] not in {"CUSTOMER_RESHOOT", "BLOCKED_INCOMPLETE_EDGE", "QA_RETURN"}:
            continue
        sheet.append([row["excel_row"], row["barcode"], row["record_id"], row["name"], row["reason"]])
        sheet.cell(sheet.max_row, 2).number_format = "@"
    path = RUN_DIR / "customer_reshoot.xlsx"
    workbook.save(path)
    print(f"reshoot rows {sheet.max_row - 1} -> {path}", flush=True)


def _write_no_photo(rows: list[dict]) -> None:
    """Barcodes with no photo in the five original folders. The customer reshoots these."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Khong co anh goc"
    sheet.append(["Excel row", "Barcode", "Record ID", "Product name", "Reason"])
    for row in rows:
        if row["reason"] != "no original photo matched this barcode":
            continue
        sheet.append([row["excel_row"], row["barcode"], row["record_id"], row["name"], row["reason"]])
        sheet.cell(sheet.max_row, 2).number_format = "@"
    path = RUN_DIR / "no_photo_reshoot.xlsx"
    workbook.save(path)
    published = ROOT / "output/catalog_pairs/Rabuta_no_photo_reshoot.xlsx"
    workbook.save(published)
    print(f"no-photo rows {sheet.max_row - 1} -> {path}", flush=True)


def _reconcile(rows: list[dict]) -> None:
    seen = set()
    for row in rows:
        if row["barcode"] in seen:
            continue
        seen.add(row["barcode"])
        if row["state"] == "ACCEPTED":
            path = FRONT_DIR / f"{row['barcode']}-front.jpg"
            if not path.exists():
                raise SystemExit(f"accepted row missing file {path.name}")
    print(f"reconciled {len(rows)} ledger rows", flush=True)


def _pilot_report(ledger: list[dict], rendered: dict[str, dict], labels: dict[str, dict],
                  labels_sha256: str, fingerprint: str) -> dict:
    groups: dict[str, dict] = {}
    for barcode, item in rendered.items():
        label = labels.get(barcode, {}).get("group", "unlabeled")
        group = groups.setdefault(label, {"count": 0, "states": {}, "reasons": {}, "api_requests": 0, "duration_ms": 0, "usage": {}})
        group["count"] += 1
        state = item.get("state", "MISSING")
        group["states"][state] = group["states"].get(state, 0) + 1
        for reason in (item.get("semantic_review") or {}).get("reason_codes", []):
            group["reasons"][reason] = group["reasons"].get(reason, 0) + 1
        group["api_requests"] += int((item.get("semantic_review") or {}).get("api_requests", 0))
        group["duration_ms"] += int(item.get("qa_duration_ms", 0))
        for usage in (item.get("semantic_review") or {}).get("usage", []):
            if isinstance(usage, dict):
                for key, value in usage.items():
                    if isinstance(value, int):
                        group["usage"][key] = group["usage"].get(key, 0) + value
    return {
        "run_fingerprint": fingerprint,
        "labels_sha256": labels_sha256,
        "rows": len(ledger),
        "groups": groups,
        "api_requests": sum(group["api_requests"] for group in groups.values()),
        "budget_used": json.loads((ROOT / "output/catalog_pairs/front-v5-pilot-state" / labels_sha256 / "budgets" / (BUDGET_ID or labels_sha256) / "request-budget.json").read_text(encoding="utf-8"))["requests"]
        if PILOT_LABELS_PATH and (ROOT / "output/catalog_pairs/front-v5-pilot-state" / labels_sha256 / "budgets" / (BUDGET_ID or labels_sha256) / "request-budget.json").is_file() else 0,
        "human_labels_are_reference_only": True,
        "automatic_enforcement": QA_MODE == "enforce",
    }


def _write_qa_gallery(run_dir: Path, rows: list[dict]) -> None:
    run_root = run_dir.resolve()
    cards = []
    for row in rows:
        if not row.get("candidate_path"):
            continue
        path = Path(row["candidate_path"]).resolve()
        try:
            relative = quote(path.relative_to(run_root).as_posix(), safe="/")
        except ValueError:
            continue
        semantic = {}
        if row.get("qa_path"):
            try:
                semantic = json.loads(Path(row["qa_path"]).read_text(encoding="utf-8")).get("semantic_review") or {}
            except (OSError, json.JSONDecodeError):
                pass
        reasons = ", ".join(semantic.get("reason_codes", [])) or row.get("reason", "")
        barcode = str(row.get("barcode", ""))
        cards.append(
            f'<article><h2>{html.escape(barcode)} — {html.escape(str(row.get("state", "")))}</h2>'
            f'<img loading="lazy" src="{relative}" alt="front candidate {html.escape(barcode)}">'
            f'<p>{html.escape(str(reasons))}</p></article>'
        )
    page = (
        "<!doctype html><meta charset=\"utf-8\"><title>Front-v5 QA gallery</title>"
        "<style>body{font:15px system-ui;margin:24px;background:#f4f4f4}"
        "main{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:16px}"
        "article{background:white;padding:12px;border:1px solid #ddd}img{width:100%;height:auto}</style>"
        f"<h1>Front-v5 QA gallery ({len(cards)} candidates)</h1><main>{''.join(cards)}</main>"
    )
    (run_dir / "qa-gallery.html").write_text(page, encoding="utf-8")


def _write_improvement_queue(run_dir: Path, rendered: dict[str, dict]) -> None:
    reasons: dict[str, list[str]] = {}
    for barcode, result in rendered.items():
        for code in (result.get("semantic_review") or {}).get("reason_codes", []):
            reasons.setdefault(str(code), []).append(barcode)
    queue = [
        {"reason_code": code, "count": len(barcodes), "barcodes": sorted(set(barcodes))}
        for code, barcodes in sorted(reasons.items())
    ]
    (run_dir / "improvement-queue.json").write_text(
        json.dumps(queue, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _promote_run(run_dir: Path, approval_path: Path) -> None:
    ledger_path = run_dir / "ledger.jsonl"
    if not ledger_path.is_file():
        raise SystemExit(f"no ledger at {ledger_path}")
    fingerprint = (run_dir / "run_fingerprint.txt").read_text(encoding="utf-8").strip()
    ledger = _read_jsonl(ledger_path)
    approved = _read_pilot_approval(
        approval_path, fingerprint,
        next((row.get("pilot_labels_sha256") for row in ledger if row.get("pilot_labels_sha256")), None),
    )
    if not approved:
        raise SystemExit("pilot approval is missing, stale, or its benchmark has not passed")
    apptance = load_apptance_records(APPTANCE_DIR, FRONT_DIR)
    outcomes = []
    for row in ledger:
        if row.get("state") != "PASS" or not row.get("candidate_path"):
            continue
        if any(variant in apptance for variant in normalize_barcode_variants(str(row["barcode"]))):
            row.update({"state": "SKIPPED_APPTANCE", "reason": "customer-accepted image is immutable"})
            _write_jsonl_atomic(ledger_path, ledger)
            continue
        job = {
            "barcode": row["barcode"], "pilot_labels_sha256": row.get("pilot_labels_sha256"),
            "expected_destination_sha256": row.get("expected_destination_sha256"),
        }
        result = {
            "state": "PASS", "run_fingerprint": fingerprint,
            "candidate_path": row["candidate_path"], "output_sha256": row["output_sha256"],
            "gates": row.get("gates", {}), "qa_path": row.get("qa_path"),
        }
        try:
            promoted = promote_candidate(
                job, result, run_dir=run_dir, front_dir=FRONT_DIR,
                apptance_dir=APPTANCE_DIR,
                expected_sha256=row.get("expected_destination_sha256"),
                approval_path=approval_path,
            )
            row.update(promoted)
        except (OSError, RuntimeError, ValueError) as error:
            row["promotion_error"] = str(error)
            outcomes.append({"barcode": row["barcode"], "state": "CONFLICT", "reason": str(error)})
            _write_jsonl_atomic(ledger_path, ledger)
            break
        else:
            outcomes.append({"barcode": row["barcode"], "state": row["state"], "output_path": row.get("output_path")})
        _write_jsonl_atomic(ledger_path, ledger)
    _write_jsonl(run_dir / "promotions.jsonl", outcomes)
    print(json.dumps({"promoted": sum(row["state"] == "ACCEPTED" for row in ledger), "conflicts": [row for row in outcomes if row["state"] == "CONFLICT"]}, ensure_ascii=False, indent=2))


def _benchmark(run_dirs: list[Path], labels_path: Path) -> None:
    labels = _load_confirmed_pilot(labels_path)
    rows = [row for directory in run_dirs for row in _read_jsonl(directory / "ledger.jsonl")]
    seen = [str(row.get("barcode")) for row in rows]
    duplicate = len(seen) != len(set(seen))
    by_barcode = {str(row.get("barcode")): row for row in rows}
    missing, extra = sorted(set(labels) - set(by_barcode)), sorted(set(by_barcode) - set(labels))
    fingerprints = {
        (directory / "run_fingerprint.txt").read_text(encoding="utf-8").strip()
        for directory in run_dirs if (directory / "run_fingerprint.txt").is_file()
    }
    false_accepts = []
    positive_fails = []
    regression_unblocked = []
    group_counts: dict[str, dict] = {}
    for barcode, label in labels.items():
        row = by_barcode.get(barcode)
        if row is None:
            continue
        state = row.get("state")
        group = group_counts.setdefault(label.get("group", "unlabeled"), {"rows": 0, "states": {}, "api_requests": 0})
        group["rows"] += 1
        group["states"][state] = group["states"].get(state, 0) + 1
        semantic = {}
        qa_path = Path(row["qa_path"]) if row.get("qa_path") else None
        if qa_path and qa_path.is_file():
            semantic = json.loads(qa_path.read_text(encoding="utf-8")).get("semantic_review") or {}
        group["api_requests"] += int(semantic.get("api_requests", 0))
        if label.get("expected_verdict") != "PASS" and state in {"PASS", "ACCEPTED"}:
            false_accepts.append(barcode)
        if label.get("case_type") == "positive" and state == "FAIL":
            positive_fails.append(barcode)
        expected_reasons = set(label.get("expected_reason_codes", []))
        actual_reasons = set(semantic.get("reason_codes", []))
        if label.get("case_type") == "regression" and state == "FAIL" and expected_reasons and not expected_reasons.intersection(actual_reasons):
            regression_unblocked.append(barcode)
    criteria = (
        len(labels) == 48 and len(by_barcode) == 48 and not duplicate and not missing and not extra
        and len(fingerprints) == 1 and not false_accepts
        and (len(positive_fails) <= 1)
        and all(label.get("expected_verdict") == "PASS" or barcode not in by_barcode
                or by_barcode[barcode].get("state") not in {"PASS", "ACCEPTED"}
                for barcode, label in labels.items())
        and not regression_unblocked
    )
    report = {
        "run_fingerprint": next(iter(fingerprints)) if len(fingerprints) == 1 else None,
        "labels_sha256": _sha256(labels_path), "rows": len(by_barcode),
        "pilot_criteria_passed": criteria,
        "false_accepts": false_accepts, "positive_false_failures": positive_fails,
        "regression_reason_mismatches": regression_unblocked,
        "missing_barcodes": missing, "extra_barcodes": extra, "duplicate_rows": duplicate,
        "group_results": group_counts,
        "positive_reviews": [
            barcode for barcode, label in labels.items()
            if label.get("case_type") == "positive" and by_barcode.get(barcode, {}).get("state") == "REVIEW"
        ],
        "regressions_blocked": sum(
            label.get("case_type") == "regression"
            and by_barcode.get(barcode, {}).get("state") in {"FAIL", "REVIEW"}
            for barcode, label in labels.items()
        ),
        "sample_limit": "48 reviewed cases do not establish catalog-wide accuracy",
    }
    destination = run_dirs[0] / "pilot-benchmark.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


def _contact_sheet(paths: list[str]) -> None:
    thumbs = []
    for raw_path in paths[:150]:
        path = Path(raw_path)
        with Image.open(path) as image:
            thumb = image.convert("RGB")
            thumb.thumbnail((160, 160))
            canvas = Image.new("RGB", (160, 160), "white")
            canvas.paste(thumb, ((160 - thumb.size[0]) // 2, (160 - thumb.size[1]) // 2))
            from PIL import ImageDraw
            ImageDraw.Draw(canvas).text((3, 3), path.stem[:22], fill="black")
            thumbs.append(canvas)
    if not thumbs:
        return
    columns = 10
    rows = (len(thumbs) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * 160, rows * 160), "white")
    for index, thumb in enumerate(thumbs):
        sheet.paste(thumb, ((index % columns) * 160, (index // columns) * 160))
    sheet.save(RUN_DIR / "shard-a-contact.jpg", "JPEG", quality=85)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=("audit", "report", "benchmark", "index", "render", "pilot", "promote", "embed"), default="audit")
    parser.add_argument("--run-id", help="isolated run id required for index/render/pilot")
    parser.add_argument("--resume", action="store_true", help="reuse only rows whose fingerprints and output hashes still match")
    parser.add_argument("--run-dir", type=Path, help="run directory for audit/report, promote or release-bound embed")
    parser.add_argument("--labels", type=Path, help="confirmed pilot labels for pilot/render or benchmark")
    parser.add_argument("--qa-mode", choices=("shadow", "enforce", "direct"), default="shadow")
    parser.add_argument("--pilot-approval", type=Path, help="hash-bound human pilot approval required for enforce mode")
    parser.add_argument("--shard-index", type=int, choices=(0, 1, 2), help="process only this modulo-3 partition of the same pilot")
    parser.add_argument("--run-dirs", nargs="+", type=Path, help="all shard run directories for a combined benchmark")
    parser.add_argument("--scope", choices=("pilot", "catalog"), default="pilot", help="catalog scope requires enforce or direct mode")
    parser.add_argument("--budget-id", help="shared 3-64 character API budget epoch id; same value across shards")
    arguments = parser.parse_args()
    if arguments.stage in {"index", "render", "pilot"}:
        if not arguments.run_id or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{2,63}", arguments.run_id):
            parser.error("index/render/pilot require a safe --run-id of 3-64 characters")
        RUN_DIR = ROOT / "output/catalog_pairs/front-v5-runs" / arguments.run_id
        if RUN_DIR.exists() and not arguments.resume:
            parser.error(f"run directory already exists; resume requires fingerprint validation: {RUN_DIR}")
        RUN_STAGE = "render" if arguments.stage == "pilot" else arguments.stage
        ALLOW_RESUME = arguments.resume
        QA_MODE = arguments.qa_mode
        PILOT_APPROVAL_PATH = arguments.pilot_approval
        PILOT_LABELS_PATH = arguments.labels
        SHARD_INDEX = arguments.shard_index
        RUN_SCOPE = "pilot" if arguments.stage == "pilot" else arguments.scope
        BUDGET_ID = arguments.budget_id or ("catalog-direct" if RUN_SCOPE == "catalog" else None)
        if RUN_SCOPE == "catalog" and QA_MODE not in {"enforce", "direct"}:
            parser.error("catalog scope is available only in enforce or direct mode")
        if BUDGET_ID and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{2,63}", BUDGET_ID):
            parser.error("--budget-id must be a safe 3-64 character identifier")
        run_render()
    elif arguments.stage == "promote":
        if arguments.run_dir is None or arguments.pilot_approval is None:
            parser.error("promote requires --run-dir and --pilot-approval")
        _promote_run(arguments.run_dir, arguments.pilot_approval)
    elif arguments.stage == "benchmark":
        if arguments.labels is None or not arguments.run_dirs:
            parser.error("benchmark requires --labels and all shard --run-dirs")
        _benchmark(arguments.run_dirs, arguments.labels)
    elif arguments.stage == "embed":
        if arguments.run_dir is None:
            parser.error("embed requires --run-dir")
        from scripts.embed_front_v5 import embed_column_x
        ledger = _read_jsonl(arguments.run_dir / "ledger.jsonl")
        fingerprint = (arguments.run_dir / "run_fingerprint.txt").read_text(encoding="utf-8").strip()
        release = json.loads((arguments.run_dir / "release.json").read_text(encoding="utf-8"))
        accepted = {row["excel_row"]: Path(row["output_path"]) for row in ledger if row.get("state") == "ACCEPTED"}
        destination = arguments.run_dir / "embedded-workbook.xlsx"
        report = embed_column_x(WORKBOOK, destination, accepted, ledger=ledger, release=release, run_fingerprint=fingerprint)
        print(json.dumps(report, indent=2))
    else:
        run_dir = arguments.run_dir
        if run_dir is None:
            print("audit is read-only; supply --run-dir to inspect a run")
            raise SystemExit(0)
        ledger_path = run_dir / "ledger.jsonl"
        if not ledger_path.is_file():
            parser.error(f"no ledger at {ledger_path}")
        items = [json.loads(line) for line in ledger_path.read_text(encoding="utf-8").splitlines() if line.strip()]
        counts: dict[str, int] = {}
        gate_counts = {f"Q{i}": {"pass": 0, "review": 0, "fail": 0, "missing": 0} for i in range(1, 7)}
        stale_accepted = []
        labels = {}
        if arguments.labels is not None:
            labels = {str(item.get("barcode")): item for item in _read_jsonl(arguments.labels)}
        groups: dict[str, dict] = {}
        for item in items:
            state = str(item.get("state", "MISSING"))
            counts[state] = counts.get(state, 0) + 1
            gates = item.get("gates", {})
            for name, values in gate_counts.items():
                gate = gates.get(name) if isinstance(gates, dict) else None
                if not isinstance(gate, dict):
                    values["missing"] += 1
                else:
                    status = gate.get("status")
                    if status in {"PASS", "REVIEW", "FAIL"}:
                        values[status.lower()] += 1
                    else:
                        values["missing"] += 1
            if state == "ACCEPTED":
                path = Path(item["output_path"]) if item.get("output_path") else None
                valid = bool(path and path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == item.get("output_sha256"))
                valid = valid and isinstance(gates, dict) and set(gates) == {f"Q{i}" for i in range(1, 7)} and all(
                    isinstance(gates[f"Q{i}"], dict) and gates[f"Q{i}"].get("passed") is True for i in range(1, 7)
                )
                if not valid:
                    stale_accepted.append(item.get("barcode"))
            label = labels.get(str(item.get("barcode")), {})
            group = str(label.get("group", "unlabeled"))
            counts_for_group = groups.setdefault(group, {"rows": 0, "states": {}, "gate_failures": {f"Q{i}": 0 for i in range(1, 7)}})
            counts_for_group["rows"] += 1
            counts_for_group["states"][state] = counts_for_group["states"].get(state, 0) + 1
            for name in gate_counts:
                gate = gates.get(name) if isinstance(gates, dict) else None
                if isinstance(gate, dict) and gate.get("status") in {"REVIEW", "FAIL"}:
                    counts_for_group["gate_failures"][name] += 1
        report = {
            "run_dir": str(run_dir.resolve()), "rows": len(items), "states": counts,
            "gate_counts": gate_counts, "groups": groups,
            "stale_accepted_count": len(stale_accepted), "stale_accepted_examples": stale_accepted[:20],
            "accepted_set_sha256": accepted_set_digest(items),
        }
        print(json.dumps(report, ensure_ascii=False, indent=2))

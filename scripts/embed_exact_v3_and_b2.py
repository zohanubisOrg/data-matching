#!/usr/bin/env python3
"""
scripts/embed_exact_v3_and_b2.py
---------------------------------
Nhúng CHÍNH XÁC ảnh từ:
- Cột 23 (W): Ảnh gốc của khách hàng trong Excel
- Cột 24 (X): Ảnh webshop mặt trước từ ĐÚNG kho front-v3 (output/catalog_pairs/tagged/front-v3)
- Cột 25 (Y): Ảnh mặt sau từ ĐÚNG kho back-v2 (output/catalog_pairs/tagged/back-v2)

Không dùng bất kỳ kho nào khác.
Chuẩn hóa toàn bộ drawing IDs để Excel mở mượt mà 0 warnings.
"""

from __future__ import annotations

import os
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.original_backup.xlsx"
FRONT_V3_DIR = ROOT / "output/catalog_pairs/tagged/front-v3"
BACK_DIR = ROOT / "output/catalog_pairs/tagged/back-v2"

OUTPUT_EXCEL = ROOT / "output/catalog_pairs/Rabuta_master_SKU_categories_ingredients_final_audit_v3.xlsx"
THUMBS_DIR = ROOT / "output/catalog_pairs/thumbnails_cache_exact_v3"
THUMBS_DIR.mkdir(parents=True, exist_ok=True)


def find_file(folder: Path, barcode: str, suffix: str):
    for sep in ("-", "_"):
        cand = folder / f"{barcode}{sep}{suffix}.jpg"
        if cand.exists():
            return cand
    cand_p = folder / f"{barcode}.jpg"
    if cand_p.exists():
        return cand_p
    if barcode.startswith("0"):
        stripped = barcode.lstrip("0")
        cand_p = folder / f"{stripped}-{suffix}.jpg"
        if cand_p.exists():
            return cand_p
        cand_p2 = folder / f"{stripped}_{suffix}.jpg"
        if cand_p2.exists():
            return cand_p2
    if len(barcode) == 12:
        with_zero = "0" + barcode
        cand_p = folder / f"{with_zero}-{suffix}.jpg"
        if cand_p.exists():
            return cand_p
    return None


def get_thumb(img_path: Path, max_size=115):
    thumb_path = THUMBS_DIR / f"thumb_{img_path.name}"
    if thumb_path.exists() and thumb_path.stat().st_size > 0:
        if thumb_path.stat().st_mtime >= img_path.stat().st_mtime:
            return thumb_path
    try:
        with PILImage.open(img_path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_size, max_size), PILImage.Resampling.LANCZOS)
            im.save(thumb_path, "JPEG", quality=90, optimize=True)
            return thumb_path
    except Exception:
        return img_path


def main():
    print("=" * 80)
    print("  NHÚNG ĐÚNG CHUẨN: [ẢNH GỐC KHÁCH] | [FRONT-V3] | [BACK-V2]")
    print("=" * 80)

    print(f"--> Đang đọc file gốc: {EXCEL_SOURCE}...")
    wb = openpyxl.load_workbook(EXCEL_SOURCE)
    ws = wb.active

    ws.insert_cols(24, 2)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c24 = ws.cell(3, 24)
    c24.value = "Ảnh webshop mặt trước"
    c24.font = header_font
    c24.fill = header_fill
    c24.alignment = align_center

    c25 = ws.cell(3, 25)
    c25.value = "Ảnh mặt sau"
    c25.font = header_font
    c25.fill = header_fill
    c25.alignment = align_center

    ws.column_dimensions["W"].width = 18
    ws.column_dimensions["X"].width = 18
    ws.column_dimensions["Y"].width = 18

    front_count = 0
    back_count = 0
    total = 0

    print("--> Đang nhúng trực tiếp từ front-v3 và back-v2...")
    for r in range(4, ws.max_row + 1):
        bc = ws.cell(r, 1).value
        if not bc:
            continue
        barcode = str(bc).strip()
        total += 1

        if ws.row_dimensions[r].height is None or ws.row_dimensions[r].height < 95:
            ws.row_dimensions[r].height = 95

        # 1. Nhúng ĐÚNG từ front-v3
        f3_p = find_file(FRONT_V3_DIR, barcode, "front")
        if f3_p:
            t = get_thumb(f3_p)
            ws.add_image(OpenpyxlImage(str(t)), f"X{r}")
            front_count += 1

        # 2. Nhúng ĐÚNG từ back-v2
        b2_p = find_file(BACK_DIR, barcode, "back")
        if b2_p:
            t = get_thumb(b2_p)
            ws.add_image(OpenpyxlImage(str(t)), f"Y{r}")
            back_count += 1

        if total % 300 == 0 or r == ws.max_row:
            print(f"    [Tiến độ] {total}/{ws.max_row - 3} SKU | Front-v3={front_count}, Back-v2={back_count}")

    print(f"\n--> Đang lưu ra file: {OUTPUT_EXCEL}...")
    ws.title = "Audit front-v3 vs back-v2"
    wb.save(OUTPUT_EXCEL)

    print("--> Đang chuẩn hóa toàn bộ Drawing IDs trong XML để mở không bị cảnh báo...")
    with zipfile.ZipFile(OUTPUT_EXCEL, "r") as zin:
        files = {n: zin.read(n) for n in zin.namelist()}

    dr_xml = files["xl/drawings/drawing1.xml"].decode("utf-8")
    dr_root = ET.fromstring(dr_xml)
    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
    doc_prs = dr_root.findall(".//xdr:cNvPr", ns)

    for idx, d in enumerate(doc_prs, start=1):
        d.set("id", str(idx))
        d.set("name", f"Picture {idx}")

    files["xl/drawings/drawing1.xml"] = ET.tostring(dr_root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(OUTPUT_EXCEL, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)

    print("\n" + "=" * 80)
    print("  HOÀN THÀNH FILE EXCEL ĐỐI CHIẾU CHUẨN XÁC THEO YÊU CẦU CỦA BẠN")
    print("=" * 80)
    print(f"  * Cột 23 (W): 'Ảnh mặt trước' (Ảnh gốc khách hàng - giữ nguyên 100%)")
    print(f"  * Cột 24 (X): 'Ảnh webshop mặt trước' (Nhúng ĐÚNG từ front-v3: {front_count} ảnh)")
    print(f"  * Cột 25 (Y): 'Ảnh mặt sau' (Nhúng ĐÚNG từ back-v2: {back_count} ảnh)")
    print(f"  * File đích : {OUTPUT_EXCEL}")
    print("=" * 80 + "\n")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
scripts/embed_front_v4_to_excel.py
-----------------------------------
Cập nhật nhúng ảnh front-v4 mới sản xuất vào file Excel đối chiếu 3 cột:
- Cột 23 (W): Ảnh mặt trước gốc của khách (100% nguyên bản)
- Cột 24 (X): Ảnh webshop mặt trước (Bản front-v4 mới nhất từ Agent 4)
- Cột 25 (Y): Ảnh mặt sau (back-v2)
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path
from typing import Optional

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.original_backup.xlsx"
FRONT_V4_DIR = ROOT / "output/catalog_pairs/tagged/front-v4"
BACK_DIR = ROOT / "output/catalog_pairs/tagged/back-v2"

OUTPUT_AUDIT_V4 = ROOT / "output/catalog_pairs/Rabuta_master_SKU_categories_ingredients_final_audit_v4.xlsx"
OUTPUT_DELIVERY = ROOT / "output/catalog_pairs/Rabuta_master_final_Delivery_Embedded.xlsx"
EXCEL_MASTER = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx"

THUMBS_DIR = ROOT / "output/catalog_pairs/thumbnails_cache_v4"
THUMBS_DIR.mkdir(parents=True, exist_ok=True)


def find_file(folder: Path, barcode: str, suffix: str = "front") -> Optional[Path]:
    for sep in ("-", "_"):
        cand = folder / f"{barcode}{sep}{suffix}.jpg"
        if cand.exists():
            return cand
    cand_plain = folder / f"{barcode}.jpg"
    if cand_plain.exists():
        return cand_plain

    if barcode.startswith("0"):
        stripped = barcode.lstrip("0")
        for sep in ("-", "_"):
            cand = folder / f"{stripped}{sep}{suffix}.jpg"
            if cand.exists():
                return cand
        cand_p = folder / f"{stripped}.jpg"
        if cand_p.exists():
            return cand_p

    if len(barcode) == 12:
        with_zero = "0" + barcode
        for sep in ("-", "_"):
            cand = folder / f"{with_zero}{sep}{suffix}.jpg"
            if cand.exists():
                return cand
        cand_p = folder / f"{with_zero}.jpg"
        if cand_p.exists():
            return cand_p

    return None


def create_thumbnail(img_path: Path, max_size: int = 115) -> Path:
    thumb_path = THUMBS_DIR / f"thumb_{img_path.name}"
    if thumb_path.exists() and thumb_path.stat().st_size > 0:
        if thumb_path.stat().st_mtime >= img_path.stat().st_mtime:
            return thumb_path

    try:
        with PILImage.open(img_path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_size, max_size), PILImage.Resampling.LANCZOS)
            im.save(thumb_path, "JPEG", quality=90, optimize=True)
            return thumb_path
    except Exception:
        return img_path


def embed_v4_to_excel():
    print("=" * 80)
    print("  NHÚNG ẢNH WEBSHOP FRONT-V4 VÀO MASTER EXCEL ĐỐI CHIẾU 3 CỘT  ")
    print("=" * 80)

    print(f"--> Đang đọc file gốc: {EXCEL_SOURCE}...")
    wb = openpyxl.load_workbook(EXCEL_SOURCE)
    ws = wb.active

    # Chèn 2 cột mới liền kề Cột 23
    ws.insert_cols(24, 2)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c24 = ws.cell(3, 24)
    c24.value = "Ảnh webshop mặt trước (v4)"
    c24.font = header_font
    c24.fill = header_fill
    c24.alignment = align_center

    c25 = ws.cell(3, 25)
    c25.value = "Ảnh mặt sau"
    c25.font = header_font
    c25.fill = header_fill
    c25.alignment = align_center

    ws.column_dimensions["W"].width = 18
    ws.column_dimensions["X"].width = 18
    ws.column_dimensions["Y"].width = 18

    front_v4_count = 0
    back_count = 0
    total = 0

    print("--> Bắt đầu nhúng ảnh front-v4 và back-v2...")

    for r in range(4, ws.max_row + 1):
        bc_val = ws.cell(r, 1).value
        if not bc_val:
            continue
        barcode = str(bc_val).strip()
        total += 1

        if ws.row_dimensions[r].height is None or ws.row_dimensions[r].height < 95:
            ws.row_dimensions[r].height = 95

        # 1. Nhúng Front-v4
        v4_path = find_file(FRONT_V4_DIR, barcode, "front")
        if v4_path:
            try:
                t_v4 = create_thumbnail(v4_path, max_size=115)
                ws.add_image(OpenpyxlImage(str(t_v4)), f"X{r}")
                front_v4_count += 1
            except Exception:
                pass

        # 2. Nhúng Back-v2
        b2_path = find_file(BACK_DIR, barcode, "back")
        if b2_path:
            try:
                t_b2 = create_thumbnail(b2_path, max_size=115)
                ws.add_image(OpenpyxlImage(str(t_b2)), f"Y{r}")
                back_count += 1
            except Exception:
                pass

        if total % 250 == 0 or r == ws.max_row:
            print(f"    [Nhúng] {total}/{ws.max_row - 3} SKU | Front-v4={front_v4_count}, Back-v2={back_count}")

    print(f"\n--> Đang lưu file bàn giao: {OUTPUT_DELIVERY}...")
    ws.title = "Master SKU Delivery v4"
    wb.save(OUTPUT_DELIVERY)
    print("    [+] Đã lưu hoàn tất OUTPUT_DELIVERY.")

    print(f"--> Đang sao chép sang: {OUTPUT_AUDIT_V4}...")
    shutil.copy2(OUTPUT_DELIVERY, OUTPUT_AUDIT_V4)

    try:
        shutil.copy2(OUTPUT_DELIVERY, EXCEL_MASTER)
        print("    [+] Đã cập nhật file gốc thành công.")
    except Exception as e:
        print(f"    [!] File gốc đang mở, bỏ qua ghi đè: {e}")

    print("\n" + "=" * 80)
    print("        HOÀN TẤT CẬP NHẬT 3 CỘT ĐỐI CHIẾU VỚI BẢN FRONT-V4")
    print("=" * 80)
    print(f"  * Tổng số ảnh front-v4 nhúng thành công : {front_v4_count}")
    print(f"  * Tổng số ảnh back-v2 nhúng thành công  : {back_count}")
    print(f"  * File sẵn sàng mở xem kiểm tra        : {OUTPUT_AUDIT_V4}")
    print("=" * 80 + "\n")


if __name__ == "__main__":
    embed_v4_to_excel()

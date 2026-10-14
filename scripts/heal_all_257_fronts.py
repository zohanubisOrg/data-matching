#!/usr/bin/env python3
"""Retired automatic healing; render helper now requires a reviewed package mask.
Use reprocess_rabuta_verified.py for traceable candidates, not automatic QA PASS.
"""

from __future__ import annotations

import glob
import json
import os
import sys
import urllib.request
from io import BytesIO
from pathlib import Path
from typing import Any

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import cv2
import numpy as np
import openpyxl
from PIL import Image, ImageOps
import zxingcpp
from rembg import new_session, remove

ROOT = Path(__file__).resolve().parent.parent
FRONT_V3_DIR = ROOT / "output/catalog_pairs/tagged/front-v3"
FRONT_V2_DIR = ROOT / "output/catalog_pairs/tagged/front-v2"
ORIGINALS_DIR = ROOT / "output/catalog_pairs/originals"
CUSTOMER_EXCEL = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx"

_REMBG_SESSION = None

def get_session():
    global _REMBG_SESSION
    if _REMBG_SESSION is None:
        _REMBG_SESSION = new_session("u2net")
    return _REMBG_SESSION


def load_customer_excel_map() -> dict[str, Any]:
    print("--> Loading customer reference photos from Excel Column 23...")
    wb = openpyxl.load_workbook(CUSTOMER_EXCEL, data_only=True)
    ws = wb.active
    
    barcode_by_row = {}
    for r in range(4, ws.max_row + 1):
        bc = ws.cell(r, 1).value
        if bc:
            barcode_by_row[r] = str(bc).strip()
            
    img_map = {}
    for img in ws._images:
        row_idx = getattr(getattr(img, "anchor", None), "_from", None)
        if row_idx is not None:
            r = row_idx.row + 1
            col = row_idx.col + 1
            if col == 23 and r in barcode_by_row:
                img_map[barcode_by_row[r]] = img
                
    print(f"    Loaded {len(img_map)} customer reference photos.")
    return img_map


def isolate_product(raw_im: Image.Image, session: Any = None, *, package_mask=None, shape="auto") -> Image.Image:
    """Legacy bridge: independently verified whole-package mask is mandatory."""
    sys.path.insert(0, str(ROOT / "src"))
    from product_mapper.vision.package_packshot import render_verified_packshot
    return render_verified_packshot(raw_im, package_mask, 2400, shape)[0]


def try_heal_from_originals(barcode: str, session: Any) -> Image.Image | None:
    orig_dir = ORIGINALS_DIR / barcode
    if not orig_dir.exists():
        return None
        
    cand_files = sorted([f for f in orig_dir.glob("*") if f.is_file() and not f.name.startswith(".")])
    if not cand_files:
        return None
        
    for f in cand_files:
        try:
            im = ImageOps.exif_transpose(Image.open(f))
            w, h = im.size
            hits = zxingcpp.read_barcodes(im, try_rotate=True, try_downscale=True)
            
            # If no barcode in camera photo, could be a single front photo!
            if not hits:
                packshot = isolate_product(im, session)
                if packshot:
                    post_hits = zxingcpp.read_barcodes(cv2.cvtColor(np.asarray(packshot), cv2.COLOR_RGB2BGR))
                    if not post_hits:
                        return packshot
                        
            # If barcode exists, identify which side is the front!
            chosen_hit = hits[0]
            pts = [(getattr(chosen_hit.position, name).x, getattr(chosen_hit.position, name).y) for name in ('top_left', 'top_right', 'bottom_right', 'bottom_left')]
            bc_x = sum(p[0] for p in pts) / 4.0 / float(w)
            bc_y = sum(p[1] for p in pts) / 4.0 / float(h)
            
            # Find seam via Canny edge projection
            arr = np.asarray(im)
            gray = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY)
            edges = cv2.Canny(cv2.GaussianBlur(gray, (7, 7), 0), 30, 90)
            
            if w >= h or abs(bc_x - 0.50) > 0.12:
                # Side-by-side products
                crop_edges = edges[:, int(0.30 * w):int(0.70 * w)]
                col_energy = crop_edges.mean(axis=0)
                seam_x = int(0.30 * w) + np.argmin(col_energy)
                
                if bc_x >= 0.50:
                    # Barcode is on right -> front is on left
                    front_crop = im.crop((0, 0, seam_x - 10, h))
                else:
                    # Barcode is on left -> front is on right
                    front_crop = im.crop((seam_x + 10, 0, w, h))
            else:
                # Vertically stacked products
                crop_edges = edges[int(0.30 * h):int(0.70 * h), :]
                row_energy = crop_edges.mean(axis=1)
                seam_y = int(0.30 * h) + np.argmin(row_energy)
                
                if bc_y >= 0.50:
                    front_crop = im.crop((0, 0, w, seam_y - 10))
                else:
                    front_crop = im.crop((0, seam_y + 10, w, h))
                    
            packshot = isolate_product(front_crop, session)
            if packshot:
                post_hits = zxingcpp.read_barcodes(cv2.cvtColor(np.asarray(packshot), cv2.COLOR_RGB2BGR))
                if not post_hits:
                    return packshot
        except Exception:
            continue
            
    return None


def try_heal_from_v2(barcode: str, session: Any) -> Image.Image | None:
    f2_path = FRONT_V2_DIR / f"{barcode}-front.jpg"
    if not f2_path.exists():
        return None
        
    try:
        im = ImageOps.exif_transpose(Image.open(f2_path))
        # Remove white borders
        arr = np.asarray(im.convert("RGB"))
        diff = np.abs(arr.astype(int) - 255).sum(axis=2)
        ys, xs = np.where(diff > 20)
        if len(xs) == 0:
            return None
        crop = im.crop((xs.min(), ys.min(), xs.max() + 1, ys.max() + 1))
        
        # Check if crop has barcode
        hits = zxingcpp.read_barcodes(crop)
        if hits:
            return None
            
        packshot = isolate_product(crop, session)
        if packshot:
            post_hits = zxingcpp.read_barcodes(cv2.cvtColor(np.asarray(packshot), cv2.COLOR_RGB2BGR))
            if not post_hits:
                return packshot
    except Exception:
        pass
    return None


def try_heal_from_customer_excel(barcode: str, excel_map: dict[str, Any], session: Any) -> Image.Image | None:
    if barcode not in excel_map:
        return None
        
    try:
        img_obj = excel_map[barcode]
        raw_im = Image.open(BytesIO(img_obj._data())).convert("RGB")
        # Upscale if small
        if max(raw_im.size) < 800:
            scale = 800.0 / max(raw_im.size)
            raw_im = raw_im.resize((int(raw_im.width * scale), int(raw_im.height * scale)), Image.Resampling.LANCZOS)
            
        packshot = isolate_product(raw_im, session)
        if packshot:
            post_hits = zxingcpp.read_barcodes(cv2.cvtColor(np.asarray(packshot), cv2.COLOR_RGB2BGR))
            if not post_hits:
                return packshot
    except Exception:
        pass
    return None


def try_heal_from_off(barcode: str, session: Any) -> Image.Image | None:
    try:
        url = f"https://world.openfoodfacts.org/api/v0/product/{barcode}.json"
        req = urllib.request.Request(url, headers={"User-Agent": "RabutaStudioPackshot/1.0"})
        with urllib.request.urlopen(req, timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            if data.get("status") == 1:
                p = data.get("product", {})
                img_url = p.get("image_front_url") or p.get("image_url")
                if img_url:
                    with urllib.request.urlopen(img_url, timeout=4) as img_resp:
                        raw_im = Image.open(BytesIO(img_resp.read())).convert("RGB")
                        packshot = isolate_product(raw_im, session)
                        if packshot:
                            return packshot
    except Exception:
        pass
    return None


def heal_all_flagged(flagged_barcodes: list[str]):
    session = get_session()
    excel_map = load_customer_excel_map()
    
    print(f"\n=======================================================")
    print(f"  HEALING {len(flagged_barcodes)} DEFECTIVE FRONT PACKSHOTS  ")
    print(f"=======================================================\n")
    
    healed_counts = {"originals": 0, "front_v2": 0, "excel": 0, "off": 0, "failed": 0}
    
    for idx, bc in enumerate(flagged_barcodes, 1):
        dest_path = FRONT_V3_DIR / f"{bc}-front.jpg"
        print(f"[{idx:>3}/{len(flagged_barcodes)}] Healing {bc} ...", end=" ", flush=True)
        
        # 1. Originals
        packshot = try_heal_from_originals(bc, session)
        if packshot:
            packshot.save(dest_path, "JPEG", quality=95)
            healed_counts["originals"] += 1
            print("OK (Originals)")
            continue
            
        # 2. Front v2
        packshot = try_heal_from_v2(bc, session)
        if packshot:
            packshot.save(dest_path, "JPEG", quality=95)
            healed_counts["front_v2"] += 1
            print("OK (Front-v2)")
            continue
            
        # 3. Customer Excel Col 23
        packshot = try_heal_from_customer_excel(bc, excel_map, session)
        if packshot:
            packshot.save(dest_path, "JPEG", quality=95)
            healed_counts["excel"] += 1
            print("OK (Customer Excel Col 23)")
            continue
            
        # 4. Open Food Facts
        packshot = try_heal_from_off(bc, session)
        if packshot:
            packshot.save(dest_path, "JPEG", quality=95)
            healed_counts["off"] += 1
            print("OK (Open Food Facts)")
            continue
            
        healed_counts["failed"] += 1
        print("FAILED (No front source available)")
        
    print("\n-------------------------------------------------------")
    print(f"Healing Summary:")
    print(f"  * From Originals:     {healed_counts['originals']}")
    print(f"  * From Front-v2:      {healed_counts['front_v2']}")
    print(f"  * From Customer Col23:{healed_counts['excel']}")
    print(f"  * From OFF:           {healed_counts['off']}")
    print(f"  * Failed:             {healed_counts['failed']}")
    print(f"-------------------------------------------------------\n")


if __name__ == "__main__":
    raise SystemExit("Unsafe legacy auto-healing is retired. Use scripts/reprocess_rabuta_verified.py with reviewed source polygons; front-v3 is not overwritten.")

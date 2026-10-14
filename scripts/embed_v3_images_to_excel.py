#!/usr/bin/env python3
"""
scripts/embed_v3_images_to_excel.py
------------------------------------
Nhúng trực tiếp 3 cột ảnh đối chiếu song song vào file chốt của khách hàng:
Rabuta_master_SKU_categories_ingredients_final.xlsx

Cấu trúc cột ảnh so sánh (Columns 23, 24, 25):
- Cột 23 (W): "Ảnh mặt trước" (Ảnh gốc của khách hàng - giữ nguyên 100%, không xóa)
- Cột 24 (X): "Ảnh mặt trước webshop" (Ảnh studio 2400x2400 từ front-v3)
- Cột 25 (Y): "Ảnh mặt sau" (Ảnh tem phụ / dinh dưỡng từ back-v2)
"""

from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/14-10-new-requirment-final/Rabuta_master_SKU_categories_ingredients_final.xlsx"
FRONT_DIR = ROOT / "output/catalog_pairs/tagged/front-v3"
BACK_DIR = ROOT / "output/catalog_pairs/tagged/back-v2"
OUTPUT_EXCEL = ROOT / "output/catalog_pairs/Rabuta_master_final_Delivery_Embedded.xlsx"
THUMBS_DIR = ROOT / "output/catalog_pairs/thumbnails_cache"
THUMBS_DIR.mkdir(parents=True, exist_ok=True)


def create_excel_thumbnail(img_path: Path, max_size: int = 120) -> Path:
    """Tạo thumbnail tỷ lệ chuẩn để nhúng vào ô Excel mà không làm phình dung lượng file."""
    thumb_path = THUMBS_DIR / f"thumb_{img_path.name}"
    if thumb_path.exists() and thumb_path.stat().st_size > 0:
        if thumb_path.stat().st_mtime >= img_path.stat().st_mtime:
            return thumb_path

    try:
        with PILImage.open(img_path) as im:
            im = im.convert("RGB")
            im.thumbnail((max_size, max_size), PILImage.Resampling.LANCZOS)
            im.save(thumb_path, "JPEG", quality=90, optimize=True)
            return thumb_path
    except Exception:
        return img_path


def embed_catalog_images():
    print("===============================================================================")
    print("  NHÚNG 3 CỘT ĐỐI CHIẾU ẢNH VÀO RABUTA MASTER FINAL (COL 23, 24, 25)  ")
    print("===============================================================================\n")

    if not EXCEL_SOURCE.exists():
        print(f"[-] Lỗi: Không tìm thấy file nguồn: {EXCEL_SOURCE}")
        return

    print("--> Đang tải workbook nguồn với các ảnh gốc của khách hàng...")
    wb = openpyxl.load_workbook(EXCEL_SOURCE)
    ws = wb.active

    print(f"    Kích thước bảng tính ban đầu: {ws.max_row} dòng, {ws.max_column} cột.")
    print(f"    Số ảnh khách hàng hiện có trong file: {len(ws._images)} ảnh.")

    # Chèn 2 cột mới tại vị trí 24 (Col X) và 25 (Col Y)
    print("--> Đang chèn 2 cột mới liền kề Cột 23 ('Ảnh mặt trước')...")
    ws.insert_cols(24, 2)

    # Đặt tiêu đề cho 2 cột mới ở dòng Header (Row 3)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c24 = ws.cell(3, 24)
    c24.value = "Ảnh mặt trước webshop"
    c24.font = header_font
    c24.fill = header_fill
    c24.alignment = align_center

    c25 = ws.cell(3, 25)
    c25.value = "Ảnh mặt sau"
    c25.font = header_font
    c25.fill = header_fill
    c25.alignment = align_center

    # Đặt độ rộng cho các cột ảnh
    ws.column_dimensions["W"].width = 18  # Ảnh mặt trước (khách)
    ws.column_dimensions["X"].width = 18  # Ảnh mặt trước webshop (front-v3)
    ws.column_dimensions["Y"].width = 18  # Ảnh mặt sau (back-v2)

    front_count = 0
    back_count = 0
    total_products = 0

    print("--> Đang nhúng ảnh Front-v3 và Back-v2 vào các cột X và Y...")

    for r in range(4, ws.max_row + 1):
        bc_val = ws.cell(r, 1).value
        if not bc_val:
            continue

        barcode = str(bc_val).strip()
        total_products += 1

        # Đảm bảo chiều cao hàng phù hợp hiển thị thumbnail 115px
        current_h = ws.row_dimensions[r].height
        if current_h is None or current_h < 95:
            ws.row_dimensions[r].height = 95

        # 1. Nhúng Ảnh mặt trước webshop (Col 24 / X)
        front_p = FRONT_DIR / f"{barcode}-front.jpg"
        if not front_p.exists():
            front_p = FRONT_DIR / f"{barcode}_front.jpg"

        if front_p.exists():
            try:
                t_front = create_excel_thumbnail(front_p, max_size=115)
                img_front = OpenpyxlImage(str(t_front))
                ws.add_image(img_front, f"X{r}")
                front_count += 1
            except Exception as exc:
                pass

        # 2. Nhúng Ảnh mặt sau (Col 25 / Y)
        back_p = BACK_DIR / f"{barcode}-back.jpg"
        if not back_p.exists():
            back_p = BACK_DIR / f"{barcode}_back.jpg"

        if back_p.exists():
            try:
                t_back = create_excel_thumbnail(back_p, max_size=115)
                img_back = OpenpyxlImage(str(t_back))
                ws.add_image(img_back, f"Y{r}")
                back_count += 1
            except Exception as exc:
                pass

        if total_products % 100 == 0:
            print(f"    Tiến độ: {total_products} sản phẩm đã duyệt (Front: {front_count}, Back: {back_count})...")

    print(f"--> Đang lưu file bàn giao: {OUTPUT_EXCEL}...")
    wb.save(OUTPUT_EXCEL)

    print("\n===============================================================================")
    print("                      HOÀN TẤT NHÚNG 3 CỘT ĐỐI CHIẾU                           ")
    print("===============================================================================")
    print(f"  * File bàn giao hoàn tất   : {OUTPUT_EXCEL}")
    print(f"  * Cột 23 (W)               : 'Ảnh mặt trước' (Giữ nguyên toàn bộ ảnh của khách)")
    print(f"  * Cột 24 (X)               : 'Ảnh mặt trước webshop' (Nhúng {front_count} ảnh)")
    print(f"  * Cột 25 (Y)               : 'Ảnh mặt sau' (Nhúng {back_count} ảnh)")
    print(f"  * Tổng số ảnh nhúng mới    : {front_count + back_count} ảnh")
    print("===============================================================================\n")


if __name__ == "__main__":
    embed_catalog_images()

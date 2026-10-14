import os
import sys
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import time
import shutil
import openpyxl
import requests
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image

EXCEL_PATH = "source-import/new_source_from_customer/11-10-new-requirment/Esbjerg_Asian_Market_Barcode_Da_Loc.xlsx"
LOCAL_STUDIO_DIR = "output/shopify_import_final/HD_Studio_Front_Images"
OUTPUT_DIR = "output/barcode_hd_images_updated"
OUTPUT_EXCEL = "output/Esbjerg_Asian_Market_Da_Loc_Cap_Nhat_Anh_Moi.xlsx"

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(os.path.dirname(OUTPUT_EXCEL), exist_ok=True)

HEADERS = {
    'User-Agent': 'RabutaAsianMarketApp/1.1 (contact@rabuta.dk; Catalog Builder)',
    'Accept': 'application/json, image/*, */*'
}

def fetch_single_barcode(row_data):
    stt, bc, name_en, name_da, brand = row_data
    bc_str = str(bc).strip() if bc else ""
    
    result = {
        "stt": stt,
        "barcode": bc_str,
        "online_title": "",
        "online_image_url": "",
        "source": "None",
        "saved_path": "",
        "width": 0,
        "height": 0,
        "status": "PENDING"
    }
    
    if not bc_str:
        result["status"] = "NO_BARCODE"
        return result

    # 1. Try OpenFoodFacts API
    off_url = f"https://world.openfoodfacts.org/api/v2/product/{bc_str}.json"
    try:
        r = requests.get(off_url, headers=HEADERS, timeout=4)
        if r.status_code == 200:
            data = r.json()
            if data.get("status") == 1:
                p = data.get("product", {})
                title = p.get("product_name_en") or p.get("product_name") or p.get("product_name_da") or ""
                img = p.get("image_front_url") or p.get("image_url") or ""
                if img:
                    result["online_title"] = title
                    result["online_image_url"] = img
                    result["source"] = "Web_OpenFoodFacts"
    except Exception:
        pass

    # 2. Download online image if found
    target_img_path = os.path.join(OUTPUT_DIR, f"{bc_str}.jpg")
    downloaded = False
    if result["online_image_url"]:
        try:
            img_r = requests.get(result["online_image_url"], headers=HEADERS, timeout=6)
            if img_r.status_code == 200 and len(img_r.content) > 3000:
                with open(target_img_path, "wb") as f:
                    f.write(img_r.content)
                downloaded = True
        except Exception:
            downloaded = False

    # 3. Fallback to Local Verified Studio Image
    if not downloaded:
        local_src = os.path.join(LOCAL_STUDIO_DIR, f"{bc_str}.jpg")
        if os.path.exists(local_src):
            shutil.copy2(local_src, target_img_path)
            result["source"] = "Studio_Kho_Verified"
            result["saved_path"] = target_img_path
        else:
            result["status"] = "MISSING_IMAGE"
            return result
    else:
        result["saved_path"] = target_img_path

    # Measure image dimensions
    try:
        with Image.open(target_img_path) as im:
            result["width"], result["height"] = im.size
        result["status"] = "OK"
    except Exception:
        result["status"] = "IMAGE_ERROR"

    return result

def main():
    print(f"[*] Reading source Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb["Da loc barcode"]
    
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[2]
    data_rows = rows[3:]
    print(f"[*] Total data rows to process: {len(data_rows)}")
    
    tasks_input = []
    for r in data_rows:
        stt = r[0]
        bc = r[2]
        name_en = r[3]
        name_da = r[4]
        brand = r[5]
        tasks_input.append((stt, bc, name_en, name_da, brand))

    results_map = {}
    online_count = 0
    studio_count = 0
    missing_count = 0

    print(f"[*] Starting multi-threaded barcode web search (Worker count: 12)...")
    start_time = time.time()
    with ThreadPoolExecutor(max_workers=12) as executor:
        futures = {executor.submit(fetch_single_barcode, item): item for item in tasks_input}
        done = 0
        for fut in as_completed(futures):
            res = fut.result()
            results_map[res["stt"]] = res
            done += 1
            if res["source"] == "Web_OpenFoodFacts":
                online_count += 1
            elif res["source"] == "Studio_Kho_Verified":
                studio_count += 1
            else:
                missing_count += 1
                
            if done % 100 == 0 or done == len(tasks_input):
                elapsed = time.time() - start_time
                print(f"    Progress: {done}/{len(tasks_input)} ({done/len(tasks_input)*100:.1f}%) | "
                      f"Web HD: {online_count} | Studio Kho: {studio_count} | Missing: {missing_count} | {elapsed:.1f}s")

    print(f"\n[+] Fetching Complete in {time.time() - start_time:.1f}s!")
    print(f"    - Web HD Images fetched: {online_count}")
    print(f"    - Studio Kho Images used: {studio_count}")
    print(f"    - Missing/No Barcode: {missing_count}")
    print(f"    - Total Image Coverage: {(online_count + studio_count)/len(tasks_input)*100:.2f}%")

    # Now write updated Excel
    print(f"\n[*] Updating Excel file...")
    wb_write = openpyxl.load_workbook(EXCEL_PATH)
    sheet_write = wb_write["Da loc barcode"]
    
    # Add new headers at row 3
    col_online_title = 102
    col_online_img_url = 103
    col_img_source = 104
    col_img_res = 105
    col_status = 106
    
    sheet_write.cell(row=3, column=col_online_title, value="Tên chuẩn theo Barcode (Web)")
    sheet_write.cell(row=3, column=col_online_img_url, value="Link ảnh gốc tìm được (Web URL)")
    sheet_write.cell(row=3, column=col_img_source, value="Nguồn ảnh hoàn thiện")
    sheet_write.cell(row=3, column=col_img_res, value="Độ phân giải ảnh (W x H)")
    sheet_write.cell(row=3, column=col_status, value="Trạng thái hoàn thiện")

    for r_idx, r in enumerate(data_rows, start=4):
        stt = r[0]
        res = results_map.get(stt, {})
        sheet_write.cell(row=r_idx, column=col_online_title, value=res.get("online_title", ""))
        sheet_write.cell(row=r_idx, column=col_online_img_url, value=res.get("online_image_url", ""))
        sheet_write.cell(row=r_idx, column=col_img_source, value=res.get("source", ""))
        w, h = res.get("width", 0), res.get("height", 0)
        res_str = f"{w}x{h}" if w and h else ""
        sheet_write.cell(row=r_idx, column=col_img_res, value=res_str)
        sheet_write.cell(row=r_idx, column=col_status, value=res.get("status", ""))

    wb_write.save(OUTPUT_EXCEL)
    print(f"[+] Successfully saved updated Master Excel to: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()

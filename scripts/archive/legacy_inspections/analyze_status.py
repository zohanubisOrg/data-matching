# -*- coding: utf-8 -*-
import sys
from collections import Counter
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

wb = openpyxl.load_workbook(r"output\shopify_import_final\Master_All_1659_Products_With_Images.xlsx", data_only=True)
ws = wb.active

headers = [cell.value for cell in ws[3]]

status_counts = Counter()
missing_bc = 0
missing_sku = 0
missing_both = 0
has_both = 0
total = 0

rows_data = []

for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    total += 1
    row_dict = dict(zip(headers, row))
    status = str(row_dict.get("Trạng thái đối soát") or "").strip()
    sku = str(row_dict.get("POS SKU") or "").strip()
    bc = str(row_dict.get("Barcode") or "").strip()
    
    status_counts[status] += 1
    if not bc and not sku:
        missing_both += 1
    elif not bc:
        missing_bc += 1
    elif not sku:
        missing_sku += 1
    else:
        has_both += 1
        
    rows_data.append((row_idx, row_dict, bc, sku, status))

print(f"Total rows: {total}")
print("Status breakdown:")
for st, cnt in status_counts.items():
    print(f"  - '{st}': {cnt}")

print(f"\nBarcode & SKU availability:")
print(f"  - Has both Barcode and SKU: {has_both}")
print(f"  - Missing both Barcode and SKU: {missing_both}")
print(f"  - Missing Barcode only: {missing_bc}")
print(f"  - Missing SKU only: {missing_sku}")

# -*- coding: utf-8 -*-
import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
SRC_DIR = os.path.join(BASE_DIR, "source-import", "new_source_from_customer")
OUT_DIR = os.path.join(BASE_DIR, "output", "shopify_import_final")

# Check all files in new_source_from_customer
print("Files in new_source_from_customer:")
for f in os.listdir(SRC_DIR):
    print("  -", f)

# Inspect 89 cols file
pm_file = [f for f in os.listdir(SRC_DIR) if f.startswith("Product_master") and f.endswith(".xlsx") and not f.startswith("~$")][0]
pm_path = os.path.join(SRC_DIR, pm_file)
wb_pm = openpyxl.load_workbook(pm_path, read_only=True, data_only=True)
ws_pm = wb_pm.active

rows_pm = list(ws_pm.iter_rows(values_only=True))
print(f"\n{pm_file}:")
print(f"Total rows in template: {len(rows_pm)}")
header_row_idx = 4 # 0-indexed row 4 is row 5 in excel
headers_89 = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows_pm[header_row_idx])]
print(f"Total headers in customer master: {len(headers_89)}")
for i, h in enumerate(headers_89, start=1):
    print(f"  Col {i:2d}: {h}")

# Inspect current Master_All_1659_Products_With_Images.xlsx
cur_master_p = os.path.join(OUT_DIR, "Master_All_1659_Products_With_Images.xlsx")
wb_cur = openpyxl.load_workbook(cur_master_p, data_only=True)
ws_cur = wb_cur.active
cur_headers = [c.value for c in ws_cur[3]]
print(f"\nCurrent Master_All_1659_Products_With_Images.xlsx:")
print(f"Total columns: {len(cur_headers)}")
for i, h in enumerate(cur_headers, start=1):
    print(f"  Col {i:2d}: {h}")

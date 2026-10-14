# -*- coding: utf-8 -*-
import os
import sys
import re
import csv
import unicodedata
import shutil
from pathlib import Path
import openpyxl
import pandas as pd

SOURCE_DIR = r"D:\Work\shopify\source-import\new_source_from_customer"
OUTPUT_DIR = r"D:\Work\shopify\output\shopify_import_final"
IMG_OUT_DIR = os.path.join(OUTPUT_DIR, "product_images")

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(IMG_OUT_DIR, exist_ok=True)

# 1. Load Category Mapping
cat_file = os.path.join(SOURCE_DIR, "Ma_nhom_hang_SKU_3digits.xlsx")
wb_cat = openpyxl.load_workbook(cat_file, data_only=True)
ws_cat = wb_cat.active
nhom_map = {}
for r in list(ws_cat.iter_rows(values_only=True))[4:]:
    if r[1]:
        dansk_name = str(r[1]).strip()
        nhom_map[dansk_name] = {
            "code": str(r[2]).strip() if r[2] else "",
            "vi": str(r[3]).strip() if r[3] else "",
            "example_sku": str(r[4]).strip() if r[4] else "",
            "online_col": str(r[5]).strip() if r[5] else "",
            "is_online": str(r[5]).strip().lower() != "nej",
        }

print(f"Loaded {len(nhom_map)} categories from Ma_nhom_hang_SKU_3digits.xlsx")

# 2. Load Product Master
pm_file = None
for f in os.listdir(SOURCE_DIR):
    if "Product_master" in f:
        pm_file = os.path.join(SOURCE_DIR, f)
        break

wb_pm = openpyxl.load_workbook(pm_file, read_only=True, data_only=True)
ws_pm = wb_pm.active
pm_rows = list(ws_pm.iter_rows(values_only=True))
pm_headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(pm_rows[4])]

pm_by_barcode = {}
for r in pm_rows[5:]:
    if any(r):
        bc = str(r[0]).strip() if r[0] is not None else ""
        if bc:
            row_dict = {pm_headers[i]: r[i] for i in range(min(len(pm_headers), len(r)))}
            pm_by_barcode[bc] = row_dict
            pm_by_barcode[bc.lstrip("0")] = row_dict

print(f"Loaded {len(pm_rows)-5} items from Product Master")

# 3. Load POS
pos_file = os.path.join(SOURCE_DIR, "27.09.26.POS_ACTIVE_FINAL.xlsx")
wb_pos = openpyxl.load_workbook(pos_file, data_only=True)
ws_pos = wb_pos.active
pos_rows = list(ws_pos.iter_rows(values_only=True))[1:]

# 4. Find all local images
image_sources = [
    r"D:\Work\shopify\output\shopify_import\image_product",
    r"D:\Work\shopify\web\public\products\front",
    r"D:\Work\shopify\output\shopify_import\front_images_jpg",
    r"D:\Work\shopify\output\shopify_import\web_front_images",
    r"D:\Work\shopify\output\shopify_import\images",
    r"D:\Work\shopify\output\g1_delivery\website_front_images",
]

image_cache = {}
for src in image_sources:
    if os.path.exists(src):
        for f in os.listdir(src):
            if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                stem = Path(f).stem
                if stem not in image_cache:
                    image_cache[stem] = os.path.join(src, f)
                # also without leading zero
                if stem.lstrip("0") not in image_cache:
                    image_cache[stem.lstrip("0")] = os.path.join(src, f)

print(f"Indexed {len(image_cache)} images from local directories")


def clean_bc(val):
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)


def slugify(text):
    text = str(text or "").strip()
    text = text.replace("æ", "ae").replace("Æ", "ae")
    text = text.replace("ø", "oe").replace("Ø", "oe")
    text = text.replace("å", "aa").replace("Å", "aa")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^\w\s-]", "", text.lower())
    text = re.sub(r"[-\s]+", "-", text).strip("-")
    return text


CATEGORY_TAXONOMY_MAP = {
    "Saucer": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "Nudler": ("Food, Beverages & Tobacco > Food Items > Nudler, Ris & Tørvarer", "Nudler"),
    "Snack": ("Food, Beverages & Tobacco > Food Items > Snack & Slik", "Snack"),
    "Krydderier": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "Konserves": ("Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk", "Konserves & Kokosmælk"),
    "Grøntsager": ("Food, Beverages & Tobacco > Food Items > Grøntsager & Frugt", "Grøntsager"),
    "Drikkevarer": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Drikker"),
    "Mel": ("Food, Beverages & Tobacco > Food Items > Mel & Bageartikler", "Mel & Bageartikler"),
    "Frugt": ("Food, Beverages & Tobacco > Food Items > Frugt & Konserves", "Frugt"),
    "Diverse": ("Food, Beverages & Tobacco > Food Items > Asiatiske Specialiteter", "Diverse"),
    "Bønner": ("Food, Beverages & Tobacco > Food Items > Bønner & Tørvarer", "Bønner"),
    "Mælk": ("Food, Beverages & Tobacco > Food Items > Mælk & Drikkevarer", "Mælk"),
}


def safe_str(v):
    if v is None or pd.isna(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"none", "nan", "null"} else s


def build_html_body(pm, pos):
    brand = safe_str(pm.get("Brand") or pm.get("brand") or "")
    net_w = safe_str(pm.get("Net quantity") or "")
    net_u = safe_str(pm.get("Net unit") or "g")
    w_str = f"{net_w} {net_u}".strip() if net_w else ""
    origin = safe_str(pm.get("Country of origin") or "")

    parts = []
    meta = []
    if brand:
        meta.append(f"Mærke: {brand}")
    if w_str:
        meta.append(f"Nettoindhold: {w_str}")
    if origin:
        meta.append(f"Oprindelsesland: {origin}")
    if meta:
        parts.append(f"<p>{' · '.join(meta)}</p>")

    ing_da = safe_str(pm.get("Ingredienser (DA)") or "")
    ing_en = safe_str(pm.get("Ingredients (EN)") or "")
    ing = ing_da or ing_en
    if ing:
        parts.append(f"<h3>Ingredienser</h3><p>{ing}</p>")

    all_da = safe_str(pm.get("Allergener (DA)") or "")
    all_en = safe_str(pm.get("Allergens (EN)") or "")
    allg = all_da or all_en
    if allg:
        parts.append(f"<h3>Allergener</h3><p>{allg}</p>")

    # Nutrition
    has_nutrition = any(
        pm.get(k) is not None and str(pm.get(k)).strip() != ""
        for k in ["Energy (kJ)", "Energy (kcal)", "Fat (g)", "Carbohydrate (g)", "Protein (g)", "Salt (g)"]
    )
    if has_nutrition:
        basis = safe_str(pm.get("Nutrition basis / preparation state") or "") or "pr. 100 g"
        if not basis.startswith("pr.") and not basis.startswith("per"):
            basis = f"pr. {basis}"
        rows = []
        ekj = pm.get("Energy (kJ)")
        ekcal = pm.get("Energy (kcal)")
        if ekj or ekcal:
            e_str = []
            if ekj and str(ekj).strip():
                e_str.append(f"{ekj} kJ")
            if ekcal and str(ekcal).strip():
                e_str.append(f"{ekcal} kcal")
            if e_str:
                rows.append(f"<tr><td>Energi</td><td>{' / '.join(e_str)}</td></tr>")
        if pm.get("Fat (g)"):
            rows.append(f"<tr><td>Fedt</td><td>{pm.get('Fat (g)')} g</td></tr>")
        if pm.get("Saturates (g)"):
            rows.append(f"<tr><td>&nbsp;&nbsp;heraf mættede fedtsyrer</td><td>{pm.get('Saturates (g)')} g</td></tr>")
        if pm.get("Carbohydrate (g)"):
            rows.append(f"<tr><td>Kulhydrat</td><td>{pm.get('Carbohydrate (g)')} g</td></tr>")
        if pm.get("Sugars (g)"):
            rows.append(f"<tr><td>&nbsp;&nbsp;heraf sukkerarter</td><td>{pm.get('Sugars (g)')} g</td></tr>")
        if pm.get("Protein (g)"):
            rows.append(f"<tr><td>Protein</td><td>{pm.get('Protein (g)')} g</td></tr>")
        if pm.get("Salt (g)"):
            rows.append(f"<tr><td>Salt</td><td>{pm.get('Salt (g)')} g</td></tr>")
        if rows:
            parts.append(f"<h3>Næringsindhold {basis}</h3><table>\n" + "\n".join(rows) + "\n</table>")

    prep_da = safe_str(pm.get("Tilberedning (DA)") or pm.get("Preparation (EN)") or "")
    if prep_da:
        parts.append(f"<h3>Tilberedning</h3><p>{prep_da}</p>")

    opb_da = safe_str(pm.get("Opbevaring (DA)") or pm.get("Storage (EN)") or "Opbevares tørt og køligt.")
    parts.append(f"<h3>Opbevaring</h3><p>{opb_da}</p>")

    return "".join(parts)


# Process all 1350 POS rows
all_mapped_records = []
online_shopify_records = []
online_master_records = []

copied_images = 0

for row in pos_rows:
    if not any(row):
        continue
    pos_cat = str(row[0]).strip() if row[0] else ""
    pos_name = str(row[1]).strip() if row[1] else ""
    off_name = str(row[2]).strip() if row[2] else ""
    prn_name = str(row[3]).strip() if row[3] else ""
    purchase_p = row[4]
    selling_p = row[5]
    unit = row[6]
    raw_bc = str(row[7]).strip() if row[7] else ""
    bc = clean_bc(raw_bc)
    sku = str(row[8]).strip() if row[8] else ""
    prod_id = str(row[9]).strip() if row[9] else ""
    tax_grp = row[10]

    cat_rule = nhom_map.get(pos_cat, {})
    is_online = cat_rule.get("is_online", False)
    group_code = cat_rule.get("code", "")
    online_group_name = cat_rule.get("online_col", "")

    # PM Lookup
    pm = pm_by_barcode.get(bc) or pm_by_barcode.get(bc.lstrip("0")) or {}

    title_da = safe_str(pm.get("Fuldt navn (DA)"))
    title_en = safe_str(pm.get("Full name (EN)"))
    brand = safe_str(pm.get("Brand")) or "Asiatisk"
    net_qty = pm.get("Net quantity")
    net_unit = pm.get("Net unit") or "g"
    origin = safe_str(pm.get("Country of origin")) or "Asien"

    # Determine title
    title = title_da or title_en or pos_name or "Asiatisk produkt"
    handle = slugify(title)
    if not handle:
        handle = f"produkt-{bc}" if bc else f"sku-{sku.lower()}"

    # Check Image
    img_path = image_cache.get(bc) or image_cache.get(bc.lstrip("0"))
    has_img = False
    img_filename = ""
    if img_path and os.path.exists(img_path):
        has_img = True
        img_ext = Path(img_path).suffix.lower()
        target_img_name = f"{bc}{img_ext}" if bc else f"{sku}{img_ext}"
        target_dest = os.path.join(IMG_OUT_DIR, target_img_name)
        if not os.path.exists(target_dest):
            shutil.copy2(img_path, target_dest)
        copied_images += 1
        img_filename = target_img_name

    cat_tax, p_type = CATEGORY_TAXONOMY_MAP.get(
        pos_cat,
        ("Food, Beverages & Tobacco > Food Items > Asiatiske Specialiteter", pos_cat or "Diverse"),
    )

    # 1. Full POS item record
    full_item = {
        "POS Category": pos_cat,
        "Online Status": "Online (Duyệt)" if is_online else "Excluded (Nej / Offline)",
        "Online Category Group": online_group_name,
        "Group Code": group_code,
        "Barcode": bc,
        "POS SKU": sku,
        "POS Product ID": prod_id,
        "POS Product Name": pos_name,
        "Official Name": off_name,
        "Printer Name": prn_name,
        "Selling Price (DKK)": selling_p,
        "Purchase Price": purchase_p,
        "Full Name (DA)": title_da,
        "Full Name (EN)": title_en,
        "Brand / Vendor": brand,
        "Origin Country": origin,
        "Has Image": "Yes" if has_img else "No",
        "Image File": img_filename,
        "PM Match Status": "Matched" if pm else "Not in PM",
        "Record ID": pm.get("Record ID", ""),
    }
    all_mapped_records.append(full_item)

    # If online, add to Shopify CSV and Master Delivery
    if is_online:
        body_html = build_html_body(pm, full_item)
        tags = f"{p_type}, {brand}, {origin}, online_import_2026"
        img_src = f"https://rabuta.dk/cdn/shop/files/{bc}.jpg" if bc else ""

        # Shopify Row
        shopify_row = {
            "Handle": handle,
            "Title": title,
            "Body (HTML)": body_html,
            "Vendor": brand,
            "Product Category": cat_tax,
            "Type": p_type,
            "Tags": tags,
            "Published": "TRUE",
            "Option1 Name": "Title",
            "Option1 Value": "Default Title",
            "Variant SKU": sku,
            "Variant Grams": int(float(net_qty)) if net_qty and str(net_qty).replace(".", "", 1).isdigit() else 100,
            "Variant Inventory Tracker": "shopify",
            "Variant Inventory Qty": 100,
            "Variant Inventory Policy": "deny",
            "Variant Fulfillment Service": "manual",
            "Variant Price": selling_p if selling_p is not None else 25,
            "Variant Compare At Price": "",
            "Variant Requires Shipping": "TRUE",
            "Variant Taxable": "TRUE",
            "Variant Barcode": bc,
            "Image Src": img_src,
            "Image Position": 1,
            "Image Alt Text": title,
            "Gift Card": "FALSE",
            "SEO Title": title,
            "SEO Description": f"Køb {title} online hos Rabuta.dk. Hurtig levering i hele Danmark.",
            "Status": "active",
        }
        online_shopify_records.append(shopify_row)

        # Online Master Record (Full columns)
        master_row = {
            "Barcode": bc,
            "POS SKU": sku,
            "Selling Price (DKK)": selling_p,
            "Product ID": prod_id,
            "POS Category": pos_cat,
            "Online Category": online_group_name,
            "Group Code": group_code,
            "Title (DA)": title_da or pos_name,
            "Title (EN)": title_en or pos_name,
            "Brand": brand,
            "Country of Origin": origin,
            "Net Quantity": net_qty,
            "Net Unit": net_unit,
            "Ingredients (DA)": pm.get("Ingredienser (DA)", ""),
            "Ingredients (EN)": pm.get("Ingredients (EN)", ""),
            "Allergens (DA)": pm.get("Allergener (DA)", ""),
            "Allergens (EN)": pm.get("Allergens (EN)", ""),
            "Energy (kJ)": pm.get("Energy (kJ)", ""),
            "Energy (kcal)": pm.get("Energy (kcal)", ""),
            "Fat (g)": pm.get("Fat (g)", ""),
            "Carbohydrate (g)": pm.get("Carbohydrate (g)", ""),
            "Sugars (g)": pm.get("Sugars (g)", ""),
            "Protein (g)": pm.get("Protein (g)", ""),
            "Salt (g)": pm.get("Salt (g)", ""),
            "Storage (DA)": pm.get("Opbevaring (DA)", ""),
            "Preparation (DA)": pm.get("Tilberedning (DA)", ""),
            "Image Filename": img_filename,
            "Image URL": img_src,
            "Record ID": pm.get("Record ID", ""),
        }
        online_master_records.append(master_row)

print(f"Total POS items processed: {len(all_mapped_records)}")
print(f"Online items for Shopify: {len(online_shopify_records)}")
print(f"Images copied to product_images: {copied_images}")

# Export CSV
csv_out = os.path.join(OUTPUT_DIR, "Shopify_Import_Online_Products.csv")
with open(csv_out, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(online_shopify_records[0].keys()))
    writer.writeheader()
    writer.writerows(online_shopify_records)
print(f"Exported Shopify CSV: {csv_out}")

# Export Excel Online Master
xlsx_master_out = os.path.join(OUTPUT_DIR, "Master_Online_Products_Consolidated.xlsx")
df_master = pd.DataFrame(online_master_records)
with pd.ExcelWriter(xlsx_master_out, engine="openpyxl") as writer:
    df_master.to_excel(writer, index=False, sheet_name="Online Products")
print(f"Exported Master Online Excel: {xlsx_master_out}")

# Export Excel All 1350 POS mapped
xlsx_all_out = os.path.join(OUTPUT_DIR, "POS_Active_All_1350_Products.xlsx")
df_all = pd.DataFrame(all_mapped_records)
with pd.ExcelWriter(xlsx_all_out, engine="openpyxl") as writer:
    df_all.to_excel(writer, index=False, sheet_name="All POS Products")
print(f"Exported All POS Excel: {xlsx_all_out}")

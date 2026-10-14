# -*- coding: utf-8 -*-
import os
import sys
import re
import csv
import unicodedata
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

SOURCE_DIR = r"D:\Work\shopify\source-import\new_source_from_customer"
OUTPUT_DIR_FINAL = r"D:\Work\shopify\output\shopify_import_final"
IMG_SOURCE_DIR = os.path.join(OUTPUT_DIR_FINAL, "product_images")
THUMBS_DIR = os.path.join(OUTPUT_DIR_FINAL, "excel_delivery_thumbs")
os.makedirs(THUMBS_DIR, exist_ok=True)

# 1. Load Category Mapping
cat_file = os.path.join(SOURCE_DIR, "Ma_nhom_hang_SKU_3digits.xlsx")
wb_cat = openpyxl.load_workbook(cat_file, data_only=True)
ws_cat = wb_cat.active
nhom_map = {}
for r in list(ws_cat.iter_rows(values_only=True))[4:]:
    if r[1]:
        dansk_name = str(r[1]).strip()
        nhom_map[dansk_name] = {
            "code": str(r[2]).strip() if r[2] else "",
            "vi": str(r[3]).strip() if r[3] else "",
            "example_sku": str(r[4]).strip() if r[4] else "",
            "online_col": str(r[5]).strip() if r[5] else "",
            "is_online": str(r[5]).strip().lower() != "nej",
        }

# 2. Load Product Master
pm_file = None
for f in os.listdir(SOURCE_DIR):
    if "Product_master" in f:
        pm_file = os.path.join(SOURCE_DIR, f)
        break

wb_pm = openpyxl.load_workbook(pm_file, read_only=True, data_only=True)
ws_pm = wb_pm.active
pm_rows = list(ws_pm.iter_rows(values_only=True))
pm_headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(pm_rows[4])]

pm_by_barcode = {}
for r in pm_rows[5:]:
    if any(r):
        bc = str(r[0]).strip() if r[0] is not None else ""
        if bc:
            row_dict = {pm_headers[i]: r[i] for i in range(min(len(pm_headers), len(r)))}
            pm_by_barcode[bc] = row_dict
            pm_by_barcode[bc.lstrip("0")] = row_dict

# 3. Load POS Active
pos_file = os.path.join(SOURCE_DIR, "27.09.26.POS_ACTIVE_FINAL.xlsx")
wb_pos = openpyxl.load_workbook(pos_file, data_only=True)
ws_pos = wb_pos.active
pos_rows = list(ws_pos.iter_rows(values_only=True))[1:]

# 4. Find all local images
image_sources = [
    r"D:\Work\shopify\output\shopify_import\image_product",
    r"D:\Work\shopify\web\public\products\front",
    r"D:\Work\shopify\output\shopify_import\front_images_jpg",
    r"D:\Work\shopify\output\shopify_import\web_front_images",
    r"D:\Work\shopify\output\shopify_import\images",
    r"D:\Work\shopify\output\g1_delivery\website_front_images",
]

image_cache = {}
for src in image_sources:
    if os.path.exists(src):
        for f in os.listdir(src):
            if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                stem = Path(f).stem
                if stem not in image_cache:
                    image_cache[stem] = os.path.join(src, f)
                if stem.lstrip("0") not in image_cache:
                    image_cache[stem.lstrip("0")] = os.path.join(src, f)


def clean_bc(val):
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)


def safe_str(v):
    if v is None or pd.isna(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"none", "nan", "null"} else s


# Generate crisp Excel thumbnails (max 150x150, JPEG quality 90)
print("Preparing high-quality thumbnails for Excel embedding...")
thumb_map = {}
for stem, orig_path in image_cache.items():
    thumb_path = os.path.join(THUMBS_DIR, f"{stem}.jpg")
    if not os.path.exists(thumb_path):
        try:
            with PILImage.open(orig_path) as pimg:
                pimg = pimg.convert("RGB")
                pimg.thumbnail((150, 150), PILImage.Resampling.LANCZOS)
                pimg.save(thumb_path, "JPEG", quality=90)
        except Exception:
            pass
    if os.path.exists(thumb_path):
        thumb_map[stem] = thumb_path
        thumb_map[stem.lstrip("0")] = thumb_path

print(f"Thumbnails ready: {len(thumb_map)} items mapped")

# Prepare online records
online_records = []
all_pos_records = []

for row in pos_rows:
    if not any(row):
        continue
    pos_cat = str(row[0]).strip() if row[0] else ""
    pos_name = str(row[1]).strip() if row[1] else ""
    off_name = str(row[2]).strip() if row[2] else ""
    prn_name = str(row[3]).strip() if row[3] else ""
    purchase_p = row[4]
    selling_p = row[5]
    unit = row[6]
    raw_bc = str(row[7]).strip() if row[7] else ""
    bc = clean_bc(raw_bc)
    sku = str(row[8]).strip() if row[8] else ""
    prod_id = str(row[9]).strip() if row[9] else ""
    tax_grp = row[10]

    cat_rule = nhom_map.get(pos_cat, {})
    is_online = cat_rule.get("is_online", False)
    group_code = cat_rule.get("code", "")
    online_group_name = cat_rule.get("online_col", "")

    pm = pm_by_barcode.get(bc) or pm_by_barcode.get(bc.lstrip("0")) or {}

    title_da = safe_str(pm.get("Fuldt navn (DA)"))
    title_en = safe_str(pm.get("Full name (EN)"))
    brand = safe_str(pm.get("Brand")) or "Asiatisk"
    net_qty = pm.get("Net quantity")
    net_unit = pm.get("Net unit") or "g"
    origin = safe_str(pm.get("Country of origin")) or "Asien"

    # Energy combine
    ekj = safe_str(pm.get("Energy (kJ)"))
    ekcal = safe_str(pm.get("Energy (kcal)"))
    energy_str = ""
    if ekj and ekcal:
        energy_str = f"{ekj} kJ / {ekcal} kcal"
    elif ekj:
        energy_str = f"{ekj} kJ"
    elif ekcal:
        energy_str = f"{ekcal} kcal"

    item_dict = {
        "Barcode": bc,
        "Ảnh sản phẩm (Image)": None,  # Will hold embedded picture
        "POS SKU": sku,
        "Selling Price (DKK)": selling_p,
        "Product ID": prod_id,
        "POS Category": pos_cat,
        "Online Category": online_group_name,
        "Fuldt navn (DA)": title_da or pos_name,
        "Full name (EN)": title_en or pos_name,
        "Brand": brand,
        "Country of Origin": origin,
        "Net Quantity": f"{net_qty} {net_unit}".strip() if net_qty else "",
        "Ingredienser (DA)": pm.get("Ingredienser (DA)", ""),
        "Allergener (DA)": pm.get("Allergener (DA)", ""),
        "Energi": energy_str,
        "Fedt (g)": pm.get("Fat (g)", ""),
        "Kulhydrat (g)": pm.get("Carbohydrate (g)", ""),
        "Protein (g)": pm.get("Protein (g)", ""),
        "Salt (g)": pm.get("Salt (g)", ""),
        "Opbevaring (DA)": pm.get("Opbevaring (DA)", ""),
        "Tilberedning (DA)": pm.get("Tilberedning (DA)", ""),
        "Record ID": pm.get("Record ID", ""),
    }

    if is_online:
        online_records.append(item_dict)

    pos_item_dict = {
        "Barcode": bc,
        "Ảnh sản phẩm (Image)": None,
        "Online Status": "Online (Duyệt)" if is_online else "Excluded (Nej / Offline)",
        "POS Category": pos_cat,
        "Online Category": online_group_name,
        "POS SKU": sku,
        "POS Product ID": prod_id,
        "POS Product Name": pos_name,
        "Selling Price (DKK)": selling_p,
        "Fuldt navn (DA)": title_da,
        "Full name (EN)": title_en,
        "Brand / Vendor": brand,
        "Origin Country": origin,
        "Record ID": pm.get("Record ID", ""),
    }
    all_pos_records.append(pos_item_dict)


# Openpyxl Styles
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)


def build_embedded_excel(records, file_path, sheet_name):
    print(f"\nBuilding {os.path.basename(file_path)} with embedded images...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True

    if not records:
        wb.save(file_path)
        return

    headers = list(records[0].keys())
    ws.append(headers)

    # Header styling
    ws.row_dimensions[1].height = 30
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(1, col_num)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    img_col_idx = headers.index("Ảnh sản phẩm (Image)") + 1
    img_col_letter = get_column_letter(img_col_idx)
    ws.column_dimensions[img_col_letter].width = 24.0

    bc_col_idx = headers.index("Barcode") + 1

    embedded_count = 0

    for r_idx, item in enumerate(records, start=2):
        row_vals = [item[h] for h in headers]
        ws.append(row_vals)

        # Set row height so customer can clearly see the image
        ws.row_dimensions[r_idx].height = 105.0

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(r_idx, c_idx)
            c.font = REGULAR_FONT
            c.border = CELL_BORDER
            c.alignment = Alignment(vertical="center")

        # Center specific columns
        ws.cell(r_idx, bc_col_idx).alignment = Alignment(horizontal="center", vertical="center")
        if "POS SKU" in headers:
            ws.cell(r_idx, headers.index("POS SKU") + 1).alignment = Alignment(horizontal="center", vertical="center")
        if "Selling Price (DKK)" in headers:
            ws.cell(r_idx, headers.index("Selling Price (DKK)") + 1).alignment = Alignment(horizontal="center", vertical="center")

        # Embed image
        bc_val = str(item.get("Barcode") or "").strip()
        thumb_p = thumb_map.get(bc_val) or thumb_map.get(bc_val.lstrip("0"))

        if thumb_p and os.path.exists(thumb_p):
            try:
                with PILImage.open(thumb_p) as pimg:
                    pw, ph = pimg.size

                # Target bounding box inside cell: max width 130, max height 130
                scale = min(130.0 / pw, 130.0 / ph)
                tw = max(35, int(pw * scale))
                th = max(35, int(ph * scale))

                xl_img = OpenpyxlImage(thumb_p)
                xl_img.width = tw
                xl_img.height = th
                xl_img.anchor = f"{img_col_letter}{r_idx}"
                ws.add_image(xl_img)
                embedded_count += 1
            except Exception as e:
                pass

    # Auto column width for other columns
    for col_idx, col_name in enumerate(headers, start=1):
        if col_idx == img_col_idx:
            continue
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for r in range(2, min(len(records) + 2, 40)):
            v = ws.cell(r, col_idx).value
            if v:
                max_len = max(max_len, min(len(str(v)), 40))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    try:
        wb.save(file_path)
        print(f"Successfully exported {file_path} ({len(records)} rows, {embedded_count} embedded images)")
    except PermissionError:
        base, ext = os.path.splitext(file_path)
        fallback = f"{base}_Delivery{ext}"
        wb.save(fallback)
        print(f"File was locked, saved fallback to: {fallback} ({len(records)} rows, {embedded_count} embedded images)")


# 1. Master Delivery File for Online Products
master_delivery = os.path.join(OUTPUT_DIR_FINAL, "Master_Online_Delivery_With_Images.xlsx")
build_embedded_excel(online_records, master_delivery, "Online Products")

# 2. Also save to Master_Online_Products_Consolidated.xlsx
master_cons = os.path.join(OUTPUT_DIR_FINAL, "Master_Online_Products_Consolidated.xlsx")
build_embedded_excel(online_records, master_cons, "Online Products")

# 3. All POS Delivery File
pos_delivery = os.path.join(OUTPUT_DIR_FINAL, "POS_Active_All_1350_Products.xlsx")
build_embedded_excel(all_pos_records, pos_delivery, "All POS Products")

print("\nDone! All files with embedded images generated cleanly.")

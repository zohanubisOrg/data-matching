# -*- coding: utf-8 -*-
import sys
import os
import re
import json
import difflib
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
SOURCE_DIR = os.path.join(BASE_DIR, "source-import")

f1 = os.path.join(SOURCE_DIR, "Ten_hang_nhap.xlsx")
f2 = os.path.join(SOURCE_DIR, "Products_imported_to_Shopbox_19_08_26_with_SKU.xlsx")

def clean_bc(v):
    s = str(v or "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)

def normalize_text(t):
    s = str(t or "").lower()
    s = re.sub(r"[^\w\s]", " ", s)
    return " ".join(s.split())

# 1. Load Supplier Table
supplier_list = []
wb1 = openpyxl.load_workbook(f1, data_only=True)
for sheet in ["EUR", "SEK", "Hàng mua tại Đan Mạch"]:
    ws = wb1[sheet]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and len(r) >= 4 and r[1]:
            supplier_list.append({
                "supplier_name": str(r[1]).strip(),
                "vendor": str(r[2]).strip() if r[2] else "",
                "supplier_sku": str(r[3]).strip() if r[3] else "",
                "norm_name": normalize_text(r[1]),
                "sheet": sheet
            })

# 2. Load Ten_nhap
wb2 = openpyxl.load_workbook(f2, data_only=True)
ten_nhap_list = []
for r in list(wb2["Ten_nhap"].iter_rows(values_only=True))[1:]:
    if r and len(r) >= 2 and r[1]:
        ten_nhap_list.append({
            "stt": r[0],
            "ten_nhap": str(r[1]).strip(),
            "norm_name": normalize_text(r[1]),
        })

print(f"Loaded {len(supplier_list)} items from Ten_hang_nhap and {len(ten_nhap_list)} items from Ten_nhap.")

# Load remaining 83
from inspect_remaining_83 import remaining_83

matched_supplier = []
for idx, batch, photo, item in remaining_83:
    brand = str(item.get("brand") or "").strip()
    name_en = str(item.get("full_name_en") or item.get("literal_label_name") or item.get("product_name") or "").strip()
    query = normalize_text(f"{brand} {name_en}")
    
    best_sup = None
    best_score = 0.0
    q_tokens = set(query.split())
    for sup in supplier_list:
        s_tokens = set(sup["norm_name"].split())
        if len(s_tokens.intersection(q_tokens)) >= 2:
            sc = difflib.SequenceMatcher(None, query, sup["norm_name"]).ratio()
            if sc > best_score:
                best_score = sc
                best_sup = sup
                
    if best_score >= 0.70 and best_sup:
        matched_supplier.append((photo, brand, name_en, best_sup["supplier_name"], best_sup["supplier_sku"], best_sup["vendor"], best_score))

print(f"Matched {len(matched_supplier)} of 83 items to Supplier Invoices / SKUs:")
for m in matched_supplier[:15]:
    print(f"Photo: {m[0]:<20} | Label: {m[1]} - {m[2][:25]} => Supplier: {m[3][:30]} (SKU: {m[4]}, Vendor: {m[5]}, Score: {m[6]:.2f})")

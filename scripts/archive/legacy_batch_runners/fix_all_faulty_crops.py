import os, sys, shutil
from pathlib import Path
from PIL import Image as PILImage, ImageOps
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts.fix_faulty_crops_cv import get_perfect_front_crop
from scripts.generate_perfect_thumbnails_and_master import extract_customer_approved_thumbnails

THUMBS_DIR = Path("output/g1_delivery/excel_thumbnails")
ORIG_DIR = Path("G1_original")
MASTER_EXCEL_OUT = Path("output/g1_delivery/G1_Master_Delivery_v2.xlsx")
WEB_EXCEL_OUT = Path("web/public/data/G1_Master_Delivery_65_Columns.xlsx")
CUSTOMER_REF_DIR = Path("output/g1_delivery/customer_reference_images")

def main():
    print("=" * 70)
    print("🔍 AUDITING & FIXING FAULTY CROPS ACROSS ALL 467 PACKAGING IMAGES")
    print("=" * 70)

    customer_refs = extract_customer_approved_thumbnails()
    print(f"Loaded {len(customer_refs)} customer benchmark references.")

    all_images = sorted(list(ORIG_DIR.glob("*.jpeg")) + list(ORIG_DIR.glob("*.jpg")))
    print(f"Total packaging photos: {len(all_images)}")

    fixed_count = 0
    kept_count = 0

    for idx, p in enumerate(all_images, 1):
        fname = p.name
        thumb_p = THUMBS_DIR / f"{fname}.jpg"
        
        # If it's a customer-approved benchmark image, ensure it is preserved
        if fname in customer_refs and customer_refs[fname].exists():
            shutil.copy2(customer_refs[fname], thumb_p)
            kept_count += 1
            continue

        # Check existing thumbnail
        needs_fix = False
        if not thumb_p.exists():
            needs_fix = True
        else:
            try:
                with PILImage.open(thumb_p) as t_im:
                    tw, th = t_im.size
                with PILImage.open(p) as o_im:
                    o_im = ImageOps.exif_transpose(o_im)
                    ow, oh = o_im.size

                # Detect if thumbnail is a horizontal slice (tw > th * 1.05) when original photo was portrait (oh > ow)
                if (oh > ow and tw > th * 1.05):
                    needs_fix = True
                # Detect if thumbnail is an unnatural vertical strip when original was landscape
                elif (ow > oh and th > tw * 2.8):
                    needs_fix = True
            except Exception:
                needs_fix = True

        if needs_fix:
            new_im, info = get_perfect_front_crop(p)
            if new_im:
                new_im.convert("RGB").save(thumb_p, "JPEG", quality=95)
                fixed_count += 1
                if fixed_count % 10 == 0:
                    print(f"Fixed faulty crop #{fixed_count}: {fname} -> {info}")
        else:
            kept_count += 1

    print("=" * 70)
    print(f"✅ AUDIT COMPLETE: Fixed {fixed_count} faulty crops | Kept {kept_count} verified crops | Total {len(all_images)}")
    print("=" * 70)

    # Re-embed all updated images into Excel
    print("\n📊 Updating Master Excel workbook with clean, un-severed HD front images...")
    wb = openpyxl.load_workbook(MASTER_EXCEL_OUT)
    ws = wb.active

    # Clear existing embedded images in sheet
    ws._images = []

    col_letter_65 = openpyxl.utils.get_column_letter(65)
    ws.column_dimensions[col_letter_65].width = 44.0

    for r_idx in range(3, ws.max_row + 1):
        ws.row_dimensions[r_idx].height = 195.0
        fname = all_images[r_idx - 3].name
        thumb_path = THUMBS_DIR / f"{fname}.jpg"
        
        if thumb_path.exists():
            try:
                with PILImage.open(thumb_path) as p_img:
                    pw, ph = p_img.size

                if ph >= pw:
                    target_h = 250
                    target_w = max(40, int(250 * pw / ph))
                else:
                    target_w = 250
                    target_h = max(40, int(250 * ph / pw))

                img_obj = OpenpyxlImage(str(thumb_path))
                img_obj.width = target_w
                img_obj.height = target_h
                img_obj.anchor = f"{col_letter_65}{r_idx}"
                ws.add_image(img_obj)
            except Exception as e:
                print(f"Warning: Failed to attach image for {fname}: {e}")

    out_excel_path = MASTER_EXCEL_OUT
    try:
        wb.save(out_excel_path)
        print(f"Saved Master Excel: {out_excel_path}")
    except PermissionError:
        fallback_out = Path("output/g1_delivery/G1_Master_Delivery_v2_latest.xlsx")
        wb.save(fallback_out)
        print(f"Warning: Main file locked, saved to {fallback_out}")
        out_excel_path = fallback_out

    shutil.copy2(out_excel_path, WEB_EXCEL_OUT)
    print(f"Exported Web Excel:    {WEB_EXCEL_OUT}")

if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
import os
import sys
import re
import csv
import unicodedata
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

SOURCE_DIR = r"D:\Work\shopify\source-import\new_source_from_customer"
OUTPUT_DIR_FINAL = r"D:\Work\shopify\output\shopify_import_final"
OUTPUT_DIR_MAIN = r"D:\Work\shopify\output\shopify_import"

os.makedirs(OUTPUT_DIR_FINAL, exist_ok=True)
os.makedirs(OUTPUT_DIR_MAIN, exist_ok=True)

# 1. Load Category Mapping
cat_file = os.path.join(SOURCE_DIR, "Ma_nhom_hang_SKU_3digits.xlsx")
wb_cat = openpyxl.load_workbook(cat_file, data_only=True)
ws_cat = wb_cat.active
nhom_map = {}
for r in list(ws_cat.iter_rows(values_only=True))[4:]:
    if r[1]:
        dansk_name = str(r[1]).strip()
        nhom_map[dansk_name] = {
            "code": str(r[2]).strip() if r[2] else "",
            "vi": str(r[3]).strip() if r[3] else "",
            "example_sku": str(r[4]).strip() if r[4] else "",
            "online_col": str(r[5]).strip() if r[5] else "",
            "is_online": str(r[5]).strip().lower() != "nej",
        }

print(f"1. Loaded {len(nhom_map)} categories from Ma_nhom_hang_SKU_3digits.xlsx")

# 2. Load Product Master
pm_file = None
for f in os.listdir(SOURCE_DIR):
    if "Product_master" in f:
        pm_file = os.path.join(SOURCE_DIR, f)
        break

wb_pm = openpyxl.load_workbook(pm_file, read_only=True, data_only=True)
ws_pm = wb_pm.active
pm_rows = list(ws_pm.iter_rows(values_only=True))
pm_headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(pm_rows[4])]

pm_by_barcode = {}
for r in pm_rows[5:]:
    if any(r):
        bc = str(r[0]).strip() if r[0] is not None else ""
        if bc:
            row_dict = {pm_headers[i]: r[i] for i in range(min(len(pm_headers), len(r)))}
            pm_by_barcode[bc] = row_dict
            pm_by_barcode[bc.lstrip("0")] = row_dict

print(f"2. Loaded {len(pm_rows)-5} items from Product Master")

# 3. Load POS Active
pos_file = os.path.join(SOURCE_DIR, "27.09.26.POS_ACTIVE_FINAL.xlsx")
wb_pos = openpyxl.load_workbook(pos_file, data_only=True)
ws_pos = wb_pos.active
pos_rows = list(ws_pos.iter_rows(values_only=True))[1:]
print(f"3. Loaded {len(pos_rows)} items from POS Active")

# 4. Find all local images
image_sources = [
    r"D:\Work\shopify\output\shopify_import\image_product",
    r"D:\Work\shopify\web\public\products\front",
    r"D:\Work\shopify\output\shopify_import\front_images_jpg",
    r"D:\Work\shopify\output\shopify_import\web_front_images",
    r"D:\Work\shopify\output\shopify_import\images",
    r"D:\Work\shopify\output\g1_delivery\website_front_images",
]

image_cache = {}
for src in image_sources:
    if os.path.exists(src):
        for f in os.listdir(src):
            if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
                stem = Path(f).stem
                if stem not in image_cache:
                    image_cache[stem] = os.path.join(src, f)
                if stem.lstrip("0") not in image_cache:
                    image_cache[stem.lstrip("0")] = os.path.join(src, f)

print(f"4. Indexed {len(image_cache)} local images")


def clean_bc(val):
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)


def slugify(text):
    text = str(text or "").strip()
    text = text.replace("æ", "ae").replace("Æ", "ae")
    text = text.replace("ø", "oe").replace("Ø", "oe")
    text = text.replace("å", "aa").replace("Å", "aa")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^\w\s-]", "", text.lower())
    text = re.sub(r"[-\s]+", "-", text).strip("-")
    return text


CATEGORY_TAXONOMY_MAP = {
    "Saucer": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "Nudler": ("Food, Beverages & Tobacco > Food Items > Nudler, Ris & Tørvarer", "Nudler"),
    "Snack": ("Food, Beverages & Tobacco > Food Items > Snack & Slik", "Snack"),
    "Krydderier": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "Konserves": ("Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk", "Konserves & Kokosmælk"),
    "Grøntsager": ("Food, Beverages & Tobacco > Food Items > Grøntsager & Frugt", "Grøntsager"),
    "Drikkevarer": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Drikker"),
    "Mel": ("Food, Beverages & Tobacco > Food Items > Mel & Bageartikler", "Mel & Bageartikler"),
    "Frugt": ("Food, Beverages & Tobacco > Food Items > Frugt & Konserves", "Frugt"),
    "Diverse": ("Food, Beverages & Tobacco > Food Items > Asiatiske Specialiteter", "Diverse"),
    "Bønner": ("Food, Beverages & Tobacco > Food Items > Bønner & Tørvarer", "Bønner"),
    "Mælk": ("Food, Beverages & Tobacco > Food Items > Mælk & Drikkevarer", "Mælk"),
}


def safe_str(v):
    if v is None or pd.isna(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"none", "nan", "null"} else s


def build_html_body(pm, pos_item):
    brand = safe_str(pm.get("Brand") or pm.get("brand") or "")
    net_w = safe_str(pm.get("Net quantity") or "")
    net_u = safe_str(pm.get("Net unit") or "g")
    w_str = f"{net_w} {net_u}".strip() if net_w else ""
    origin = safe_str(pm.get("Country of origin") or "")

    parts = []
    meta = []
    if brand:
        meta.append(f"Mærke: {brand}")
    if w_str:
        meta.append(f"Nettoindhold: {w_str}")
    if origin:
        meta.append(f"Oprindelsesland: {origin}")
    if meta:
        parts.append(f"<p>{' · '.join(meta)}</p>")

    ing_da = safe_str(pm.get("Ingredienser (DA)") or "")
    ing_en = safe_str(pm.get("Ingredients (EN)") or "")
    ing = ing_da or ing_en
    if ing:
        parts.append(f"<h3>Ingredienser</h3><p>{ing}</p>")

    all_da = safe_str(pm.get("Allergener (DA)") or "")
    all_en = safe_str(pm.get("Allergens (EN)") or "")
    allg = all_da or all_en
    if allg:
        parts.append(f"<h3>Allergener</h3><p>{allg}</p>")

    has_nutrition = any(
        pm.get(k) is not None and str(pm.get(k)).strip() != ""
        for k in ["Energy (kJ)", "Energy (kcal)", "Fat (g)", "Carbohydrate (g)", "Protein (g)", "Salt (g)"]
    )
    if has_nutrition:
        basis = safe_str(pm.get("Nutrition basis / preparation state") or "") or "pr. 100 g"
        if not basis.startswith("pr.") and not basis.startswith("per"):
            basis = f"pr. {basis}"
        rows = []
        ekj = pm.get("Energy (kJ)")
        ekcal = pm.get("Energy (kcal)")
        if ekj or ekcal:
            e_str = []
            if ekj and str(ekj).strip():
                e_str.append(f"{ekj} kJ")
            if ekcal and str(ekcal).strip():
                e_str.append(f"{ekcal} kcal")
            if e_str:
                rows.append(f"<tr><td>Energi</td><td>{' / '.join(e_str)}</td></tr>")
        if pm.get("Fat (g)"):
            rows.append(f"<tr><td>Fedt</td><td>{pm.get('Fat (g)')} g</td></tr>")
        if pm.get("Saturates (g)"):
            rows.append(f"<tr><td>&nbsp;&nbsp;heraf mættede fedtsyrer</td><td>{pm.get('Saturates (g)')} g</td></tr>")
        if pm.get("Carbohydrate (g)"):
            rows.append(f"<tr><td>Kulhydrat</td><td>{pm.get('Carbohydrate (g)')} g</td></tr>")
        if pm.get("Sugars (g)"):
            rows.append(f"<tr><td>&nbsp;&nbsp;heraf sukkerarter</td><td>{pm.get('Sugars (g)')} g</td></tr>")
        if pm.get("Protein (g)"):
            rows.append(f"<tr><td>Protein</td><td>{pm.get('Protein (g)')} g</td></tr>")
        if pm.get("Salt (g)"):
            rows.append(f"<tr><td>Salt</td><td>{pm.get('Salt (g)')} g</td></tr>")
        if rows:
            parts.append(f"<h3>Næringsindhold {basis}</h3><table>\n" + "\n".join(rows) + "\n</table>")

    prep_da = safe_str(pm.get("Tilberedning (DA)") or pm.get("Preparation (EN)") or "")
    if prep_da:
        parts.append(f"<h3>Tilberedning</h3><p>{prep_da}</p>")

    opb_da = safe_str(pm.get("Opbevaring (DA)") or pm.get("Storage (EN)") or "Opbevares tørt og køligt.")
    parts.append(f"<h3>Opbevaring</h3><p>{opb_da}</p>")

    return "".join(parts)


all_pos_rows = []
online_shopify_rows = []
online_master_rows = []

img_dest_dir = os.path.join(OUTPUT_DIR_FINAL, "product_images")
os.makedirs(img_dest_dir, exist_ok=True)

for row in pos_rows:
    if not any(row):
        continue
    pos_cat = str(row[0]).strip() if row[0] else ""
    pos_name = str(row[1]).strip() if row[1] else ""
    off_name = str(row[2]).strip() if row[2] else ""
    prn_name = str(row[3]).strip() if row[3] else ""
    purchase_p = row[4]
    selling_p = row[5]
    unit = row[6]
    raw_bc = str(row[7]).strip() if row[7] else ""
    bc = clean_bc(raw_bc)
    sku = str(row[8]).strip() if row[8] else ""
    prod_id = str(row[9]).strip() if row[9] else ""
    tax_grp = row[10]

    cat_rule = nhom_map.get(pos_cat, {})
    is_online = cat_rule.get("is_online", False)
    group_code = cat_rule.get("code", "")
    online_group_name = cat_rule.get("online_col", "")

    pm = pm_by_barcode.get(bc) or pm_by_barcode.get(bc.lstrip("0")) or {}

    title_da = safe_str(pm.get("Fuldt navn (DA)"))
    title_en = safe_str(pm.get("Full name (EN)"))
    brand = safe_str(pm.get("Brand")) or "Asiatisk"
    net_qty = pm.get("Net quantity")
    net_unit = pm.get("Net unit") or "g"
    origin = safe_str(pm.get("Country of origin")) or "Asien"

    img_text = f"{bc}.jpg" if bc else f"{sku}.jpg"

    # Copy image to final folder if exists
    img_path = image_cache.get(bc) or image_cache.get(bc.lstrip("0"))
    if img_path and os.path.exists(img_path):
        dest_p = os.path.join(img_dest_dir, f"{bc}.jpg" if bc else f"{sku}.jpg")
        if not os.path.exists(dest_p):
            shutil.copy2(img_path, dest_p)

    title = title_da or title_en or pos_name or "Asiatisk produkt"
    handle = slugify(title)
    if not handle:
        handle = f"produkt-{bc}" if bc else f"sku-{sku.lower()}"

    cat_tax, p_type = CATEGORY_TAXONOMY_MAP.get(
        pos_cat,
        ("Food, Beverages & Tobacco > Food Items > Asiatiske Specialiteter", pos_cat or "Diverse"),
    )

    # Record for All POS items
    pos_item_dict = {
        "Barcode": bc,
        "Image File": img_text,
        "POS Category": pos_cat,
        "Online Status": "Online (Duyệt)" if is_online else "Excluded (Nej / Offline)",
        "Online Category Group": online_group_name,
        "Group Code": group_code,
        "POS SKU": sku,
        "POS Product ID": prod_id,
        "POS Product Name": pos_name,
        "Official Name": off_name,
        "Printer Name": prn_name,
        "Selling Price (DKK)": selling_p,
        "Purchase Price": purchase_p,
        "Full Name (DA)": title_da,
        "Full Name (EN)": title_en,
        "Brand / Vendor": brand,
        "Origin Country": origin,
        "PM Match Status": "Matched" if pm else "Not in PM",
        "Record ID": pm.get("Record ID", ""),
    }
    all_pos_rows.append(pos_item_dict)

    if is_online:
        body_html = build_html_body(pm, pos_item_dict)
        tags = f"{p_type}, {brand}, {origin}, online_catalog_2026"
        img_src = f"https://rabuta.dk/cdn/shop/files/{bc}.jpg" if bc else ""

        shopify_dict = {
            "Handle": handle,
            "Title": title,
            "Body (HTML)": body_html,
            "Vendor": brand,
            "Product Category": cat_tax,
            "Type": p_type,
            "Tags": tags,
            "Published": "TRUE",
            "Option1 Name": "Title",
            "Option1 Value": "Default Title",
            "Variant SKU": sku,
            "Variant Grams": int(float(net_qty)) if net_qty and str(net_qty).replace(".", "", 1).isdigit() else 100,
            "Variant Inventory Tracker": "shopify",
            "Variant Inventory Qty": 100,
            "Variant Inventory Policy": "deny",
            "Variant Fulfillment Service": "manual",
            "Variant Price": selling_p if selling_p is not None else 25,
            "Variant Compare At Price": "",
            "Variant Requires Shipping": "TRUE",
            "Variant Taxable": "TRUE",
            "Variant Barcode": bc,
            "Image Src": img_src,
            "Image Position": 1,
            "Image Alt Text": title,
            "Gift Card": "FALSE",
            "SEO Title": title,
            "SEO Description": f"Køb {title} online hos Rabuta.dk. Hurtig levering i hele Danmark.",
            "Status": "active",
        }
        online_shopify_rows.append(shopify_dict)

        master_online_dict = {
            "Barcode": bc,
            "Image File": img_text,
            "POS SKU": sku,
            "Selling Price (DKK)": selling_p,
            "Product ID": prod_id,
            "POS Category": pos_cat,
            "Online Category": online_group_name,
            "Group Code": group_code,
            "Title (DA)": title_da or pos_name,
            "Title (EN)": title_en or pos_name,
            "Brand": brand,
            "Country of Origin": origin,
            "Net Quantity": net_qty,
            "Net Unit": net_unit,
            "Ingredients (DA)": pm.get("Ingredienser (DA)", ""),
            "Ingredients (EN)": pm.get("Ingredients (EN)", ""),
            "Allergens (DA)": pm.get("Allergener (DA)", ""),
            "Allergens (EN)": pm.get("Allergens (EN)", ""),
            "Energy (kJ)": pm.get("Energy (kJ)", ""),
            "Energy (kcal)": pm.get("Energy (kcal)", ""),
            "Fat (g)": pm.get("Fat (g)", ""),
            "Carbohydrate (g)": pm.get("Carbohydrate (g)", ""),
            "Sugars (g)": pm.get("Sugars (g)", ""),
            "Protein (g)": pm.get("Protein (g)", ""),
            "Salt (g)": pm.get("Salt (g)", ""),
            "Storage (DA)": pm.get("Opbevaring (DA)", ""),
            "Preparation (DA)": pm.get("Tilberedning (DA)", ""),
            "Record ID": pm.get("Record ID", ""),
        }
        online_master_rows.append(master_online_dict)

print(f"Total POS items: {len(all_pos_rows)}")
print(f"Online items: {len(online_shopify_rows)}")

HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)


def export_formatted_excel(records, file_path, sheet_name):
    df = pd.DataFrame(records)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True

    headers = list(df.columns)
    ws.append(headers)

    ws.row_dimensions[1].height = 26
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(1, col_num)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append(list(row))
        ws.row_dimensions[r_idx].height = 20
        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(r_idx, c_idx)
            c.font = REGULAR_FONT
            c.border = CELL_BORDER
            c.alignment = Alignment(vertical="center")

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for r in range(2, min(len(df) + 2, 50)):
            v = ws.cell(r, col_idx).value
            if v:
                max_len = max(max_len, min(len(str(v)), 45))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    try:
        wb.save(file_path)
        print(f"Saved: {file_path}")
    except PermissionError:
        base, ext = os.path.splitext(file_path)
        fallback = f"{base}_Final{ext}"
        wb.save(fallback)
        print(f"File locked, saved fallback to: {fallback}")


# Save CSV
csv_final = os.path.join(OUTPUT_DIR_FINAL, "Shopify_Import_Online_Products.csv")
with open(csv_final, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(online_shopify_rows[0].keys()))
    w.writeheader()
    w.writerows(online_shopify_rows)
print(f"Saved CSV: {csv_final}")

# Save Excel files
master_final = os.path.join(OUTPUT_DIR_FINAL, "Master_Online_Products_Consolidated.xlsx")
export_formatted_excel(online_master_rows, master_final, "Online Products")

pos_final = os.path.join(OUTPUT_DIR_FINAL, "POS_Active_All_1350_Products.xlsx")
export_formatted_excel(all_pos_rows, pos_final, "All POS Products")

print("All pipeline steps executed successfully!")

import os
import sys
import re
import time
import json
from io import BytesIO
from pathlib import Path
from difflib import SequenceMatcher

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
import requests
from PIL import Image, ImageOps, ImageEnhance
import rembg

EXCEL_SOURCE = "source-import/new_source_from_customer/12-10-new-requirment/Rabuta.12.10.2026_original.xlsx"
OUTPUT_DIR = "output/webshop_hd_delivery"
FRONT_DIR = os.path.join(OUTPUT_DIR, "front")
BACK_DIR = os.path.join(OUTPUT_DIR, "back")
REPORT_EXCEL = "output/Rabuta_2_Sided_Pilot_Audit.xlsx"
GALLERY_HTML = "output/pilot_20_gallery.html"
COOKIE_FILE = "output/google_cookies.json"

os.makedirs(FRONT_DIR, exist_ok=True)
os.makedirs(BACK_DIR, exist_ok=True)
os.makedirs(os.path.dirname(REPORT_EXCEL), exist_ok=True)

HEADERS = {
    'User-Agent': 'RabutaCatalogEngine/3.0 (contact@rabuta.dk; Identity Verification Pipeline)',
    'Accept': 'application/json, image/*, */*'
}

# ==========================================
# 1. GROUND TRUTH LOADER
# ==========================================
def load_ground_truth(file_path=EXCEL_SOURCE, limit=20, offset=0):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    data_rows = rows[3:] # Header is row 3 (0-indexed 2)
    
    products = []
    selected = data_rows[offset : offset + limit] if limit else data_rows[offset:]
    
    for idx, r in enumerate(selected, start=offset + 1):
        bc = str(r[0] or "").strip()
        if bc.endswith(".0"):
            bc = bc[:-2]
        
        item = {
            "stt": idx,
            "barcode": bc,
            "name_en": str(r[1] or "").strip(),
            "name_da": str(r[2] or "").strip(),
            "brand": str(r[3] or "").strip(),
            "record_id": str(r[4] or "").strip(),
            "variant": str(r[7] or "").strip() if r[7] else "",
            "category": str(r[8] or "").strip() if r[8] else "",
            "packaging": str(r[9] or "").strip() if r[9] else "",
            "net_quantity": str(r[10] or "").strip() if r[10] else "",
            "net_unit": str(r[11] or "").strip() if r[11] else "",
            "ingredients_en": str(r[14] or "").strip() if r[14] else "",
            "ingredients_da": str(r[15] or "").strip() if r[15] else ""
        }
        item["full_qty"] = f"{item['net_quantity']} {item['net_unit']}".strip()
        products.append(item)
    return products

# ==========================================
# 2. DETERMINISTIC IDENTITY MATCHER
# ==========================================
FLAVOR_CLUSTERS = {
    "coconut": ["coconut", "coco", "kokos", "kokosnød", "dừa", "dua", "vi-dua"],
    "mangosteen": ["mangosteen", "mangoustan", "mangostan", "măng cụt", "mang cut", "mang-ct", "vi-mang-ct", "vi-mang-cut"],
    "lychee": ["lychee", "litchi", "vải", "vai", "vi-vai"],
    "wasabi": ["wasabi"],
    "soy_sauce": ["soy sauce", "shoyu", "soja", "sojasovs", "sojasauce", "xì dầu", "xi dau", "nước tương", "nuoc tuong"],
    "sesame": ["sesame", "mè", "me", "vừng", "vung", "sesam"],
    "seaweed": ["seaweed", "nori", "tang", "algen", "algues", "rong biển", "rong bien"],
    "chilli": ["chilli", "chili", "piment", "ớt", "ot"],
    "mango": ["mango", "xoài", "xoai", "vi-xoai"],
    "strawberry": ["strawberry", "fraise", "erdbeere", "dâu", "dau", "vi-dau"],
    "red_bean": ["red bean", "haricot rouge", "đậu đỏ", "dau do", "azuki"],
    "milk": ["milk", "lait", "milch", "mælk", "sữa", "sua"]
}

def get_product_clusters(variant_str: str) -> set:
    if not variant_str or variant_str.lower() in ("none", ""):
        return set()
    v_clean = variant_str.lower()
    matched_clusters = set()
    for c_name, keywords in FLAVOR_CLUSTERS.items():
        if any(kw in v_clean for kw in keywords):
            matched_clusters.add(c_name)
    return matched_clusters

def text_similarity(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a.lower(), b.lower()).ratio()

def validate_identity_candidate(product: dict, title: str, url: str) -> tuple[bool, str]:
    text_to_check = f"{title} {url}".lower()
    
    # 1. Barcode check in URL or Title -> instant high confidence if present
    bc = product["barcode"]
    has_exact_bc = bc and (bc in text_to_check)

    # 2. Variant / Flavour Hard Conflict Check (Multilingual Clusters)
    target_clusters = get_product_clusters(product["variant"])
    if target_clusters:
        for c_name, keywords in FLAVOR_CLUSTERS.items():
            if c_name not in target_clusters:
                # If candidate contains any keyword from another cluster -> HARD REJECT!
                for kw in keywords:
                    if kw in text_to_check:
                        return False, f"Xung đột hương vị: Chứa '{kw}' (thuộc {c_name}) trong khi sản phẩm là '{product['variant']}'"
    
    # 3. Brand Matching
    brand = product["brand"].lower()
    brand_tokens = [b.strip() for b in re.split(r'[/,]', brand) if len(b.strip()) > 2]
    brand_match = any(b in text_to_check for b in brand_tokens)
    
    # 4. Weight / Quantity Check (prevent 100g vs 180g)
    qty = product["net_quantity"]
    if qty and qty.isdigit():
        qty_num = int(qty)
        # Search for other gram numbers
        found_weights = re.findall(r'(\d+)\s*(?:g|gram|ml)\b', text_to_check)
        for fw in found_weights:
            fw_num = int(fw)
            if abs(fw_num - qty_num) > 25 and fw_num not in (1, 2, 4, 6): # ignore pack counts
                if not has_exact_bc:
                    return False, f"Xung đột trọng lượng: Tìm thấy '{fw}g' thay vì '{qty_num}g'"

    if has_exact_bc:
        return True, "Khớp Barcode chính xác"
    elif brand_match:
        return True, "Khớp Brand và không có xung đột variant"
    
    return False, "Không đủ bằng chứng định danh"

# ==========================================
# 3. MULTI-SOURCE IMAGE HARVESTER
# ==========================================
def fetch_off_data(barcode: str):
    url = f"https://world.openfoodfacts.org/api/v2/product/{barcode}?fields=code,product_name,brands,quantity,images,selected_images"
    try:
        r = requests.get(url, headers=HEADERS, timeout=8)
        if r.status_code == 200:
            return r.json().get("product", {})
    except Exception:
        pass
    return {}

def extract_off_best_images(off_prod: dict, product: dict):
    if not off_prod:
        return None, None, ""

    sel = off_prod.get("selected_images", {})
    images = off_prod.get("images", {})
    code = off_prod.get("code") or product["barcode"]
    
    # Helper to construct full-res URL
    def to_full_url(u: str) -> str:
        if not u:
            return ""
        if ".400.jpg" in u:
            return u.replace(".400.jpg", ".full.jpg")
        return u

    # 1. FRONT
    front_url = ""
    if sel.get("front", {}).get("display"):
        front_url = list(sel["front"]["display"].values())[0]
    elif off_prod.get("image_front_url"):
        front_url = off_prod.get("image_front_url")
    front_url = to_full_url(front_url)

    # 2. BACK / INGREDIENTS
    back_url = ""
    if sel.get("ingredients", {}).get("display"):
        back_url = list(sel["ingredients"]["display"].values())[0]
    elif sel.get("packaging", {}).get("display"):
        back_url = list(sel["packaging"]["display"].values())[0]
    elif sel.get("nutrition", {}).get("display"):
        back_url = list(sel["nutrition"]["display"].values())[0]
    elif off_prod.get("image_ingredients_url"):
        back_url = off_prod.get("image_ingredients_url")
    elif off_prod.get("image_packaging_url"):
        back_url = off_prod.get("image_packaging_url")
    back_url = to_full_url(back_url)

    # If back_url not in selected, check raw image keys in images dict
    if not back_url:
        for k, v in images.items():
            if isinstance(v, dict) and k.isdigit() and int(k) > 1:
                if len(code) == 13:
                    split_code = f"{code[:3]}/{code[3:6]}/{code[6:9]}/{code[9:]}"
                else:
                    split_code = code
                candidate_raw = f"https://images.openfoodfacts.org/images/products/{split_code}/{k}.jpg"
                back_url = candidate_raw
                break

    return front_url, back_url, "OpenFoodFacts"

def search_web_verified_images(page, product: dict):
    bc = product["barcode"]
    brand = product["brand"]
    front_img_bytes, front_url = None, None
    back_img_bytes, back_url = None, None

    # Exact barcode search
    query = f'"{bc}"'
    url = f"https://www.google.com/search?q={query}&udm=2&hl=en&tbs=isz:l" # Large images only
    try:
        page.goto(url, timeout=18000)
        page.wait_for_timeout(1800)

        if "sorry" not in page.url:
            grid_imgs = page.locator("div#rso img, div#search img").all()
            for img in grid_imgs[:6]:
                if not img.is_visible():
                    continue
                try:
                    img.click(timeout=2000)
                    page.wait_for_timeout(1200)

                    # Extract all image candidates in detail view
                    candidates = page.locator("img[src^='http']").all()
                    for cand in candidates:
                        src = cand.get_attribute("src") or ""
                        alt = cand.get_attribute("alt") or ""
                        if "gstatic" not in src and "google" not in src and any(ext in src.lower() for ext in ['.jpg', '.jpeg', '.png', '.webp']):
                            # Validate Identity
                            is_valid, reason = validate_identity_candidate(product, alt, src)
                            if is_valid:
                                # Try download and inspect
                                try:
                                    r = requests.get(src, headers={'User-Agent': 'Mozilla/5.0'}, timeout=8)
                                    if r.status_code == 200 and len(r.content) > 30000:
                                        im_temp = Image.open(BytesIO(r.content))
                                        w, h = im_temp.size
                                        if w >= 1000 or h >= 1000:
                                            # Check if front or back based on alt/url keywords
                                            is_back = any(kw in f"{alt} {src}".lower() for kw in ['back', 'rear', 'ingredient', 'nutrition', 'label', 'bagside', 'ruckseite', 'nutr'])
                                            if is_back and not back_img_bytes:
                                                back_img_bytes = r.content
                                                back_url = src
                                            elif not front_img_bytes:
                                                front_img_bytes = r.content
                                                front_url = src
                                except Exception:
                                    pass
                        if front_img_bytes and back_img_bytes:
                            break
                except Exception:
                    pass
                if front_img_bytes and back_img_bytes:
                    break
    except Exception:
        pass

    return front_img_bytes, front_url, back_img_bytes, back_url

# ==========================================
# 4. STUDIO POST-PROCESSING
# ==========================================
def process_front_packshot(img_bytes: bytes, target_size=1600) -> bytes:
    """
    Auto-rotates, removes background with AI, centers on pure white canvas 1600x1600 with 8% padding.
    """
    im = Image.open(BytesIO(img_bytes))
    im = ImageOps.exif_transpose(im)
    if im.mode != "RGB":
        im = im.convert("RGB")
    
    # Fast optimize: If raw image is huge (e.g. 4000px), resize to max 1400px for lightning-fast AI mask
    w, h = im.size
    if max(w, h) > 1400:
        scale_proc = 1400 / max(w, h)
        im = im.resize((int(w * scale_proc), int(h * scale_proc)), Image.Resampling.LANCZOS)

    # AI Background Removal
    buf_in = BytesIO()
    im.save(buf_in, format="PNG")
    nobg_bytes = rembg.remove(buf_in.getvalue())
    nobg_im = Image.open(BytesIO(nobg_bytes))

    # Crop to non-transparent bounding box
    bbox = nobg_im.getbbox()
    if bbox:
        nobg_im = nobg_im.crop(bbox)
    
    cw, ch = nobg_im.size
    # Fit into 1600x1600 with 8% safe margin (outer 128px)
    max_inner = int(target_size * 0.84) # 1344 px
    scale = min(max_inner / cw, max_inner / ch)
    new_w = max(10, int(cw * scale))
    new_h = max(10, int(ch * scale))

    nobg_resized = nobg_im.resize((new_w, new_h), Image.Resampling.LANCZOS)

    # Pure White Studio Canvas
    canvas = Image.new("RGB", (target_size, target_size), (255, 255, 255))
    paste_x = (target_size - new_w) // 2
    paste_y = (target_size - new_h) // 2

    canvas.paste(nobg_resized, (paste_x, paste_y), mask=nobg_resized.split()[3])

    buf_out = BytesIO()
    canvas.save(buf_out, format="JPEG", quality=94, optimize=True)
    return buf_out.getvalue()

def process_back_label(img_bytes: bytes, target_max=2000) -> bytes:
    """
    Enhances contrast & sharpness for crisp text reading, preserving full label aspect ratio.
    """
    im = Image.open(BytesIO(img_bytes))
    im = ImageOps.exif_transpose(im)
    if im.mode != "RGB":
        im = im.convert("RGB")

    w, h = im.size
    # Resize only if excessively large, keeping sharpness
    if max(w, h) > target_max:
        scale = target_max / max(w, h)
        im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)

    # Enhance Sharpness & Contrast for crystal-clear ingredient text
    enhancer_sharp = ImageEnhance.Sharpness(im)
    im = enhancer_sharp.enhance(1.4)
    enhancer_cont = ImageEnhance.Contrast(im)
    im = enhancer_cont.enhance(1.15)

    buf_out = BytesIO()
    im.save(buf_out, format="JPEG", quality=94, optimize=True)
    return buf_out.getvalue()

# ==========================================
# 5. PIPELINE EXECUTION FOR PILOT 20
# ==========================================
def run_pilot_20():
    from playwright.sync_api import sync_playwright

    print("=================================================================")
    print("[*] KHỞI ĐỘNG PIPELINE XÁC THỰC DANH TÍNH & BỘ ẢNH 2 MẶT SIÊU NÉT")
    print(f"[*] Dữ liệu nguồn: {EXCEL_SOURCE}")
    print(f"[*] Đích xuất: {OUTPUT_DIR}")
    print("=================================================================\n")

    products = load_ground_truth(EXCEL_SOURCE, limit=20, offset=0)
    print(f"[+] Đã tải thành công 20 sản phẩm đầu tiên từ file Excel:")
    for p in products:
        print(f"    - [{p['stt']:02d}] {p['barcode']} | {p['brand'][:15]} | {p['name_en'][:35]} | Var: {p['variant']} | Qty: {p['full_qty']}")

    cookies = []
    if os.path.exists(COOKIE_FILE):
        with open(COOKIE_FILE, "r", encoding="utf-8") as f:
            try:
                cookies = json.load(f)
            except Exception:
                pass

    results = []

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True, args=['--disable-blink-features=AutomationControlled', '--no-sandbox'])
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36",
            viewport={"width": 1366, "height": 768}
        )
        context.add_init_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        if cookies:
            try:
                context.add_cookies(cookies)
            except Exception:
                pass
        page = context.new_page()

        for idx, prod in enumerate(products, start=1):
            bc = prod["barcode"]
            print(f"\n" + "="*65)
            print(f"[{idx}/20] SP: {prod['name_en']} (BC: {bc})")
            print(f"        Variant: {prod['variant']} | Trọng lượng: {prod['full_qty']}")
            print("="*65)

            item_res = {
                "stt": prod["stt"],
                "barcode": bc,
                "name": prod["name_en"],
                "brand": prod["brand"],
                "variant": prod["variant"],
                "quantity": prod["full_qty"],
                "front_status": "Chưa có",
                "front_res": "",
                "front_src": "",
                "front_path": "",
                "back_status": "Chưa có",
                "back_res": "",
                "back_src": "",
                "back_path": "",
                "tier": "Tier 3 (Retake Flagged)",
                "note": ""
            }

            # ----------------------------------------------------
            # TẦNG 1: OPEN FOOD FACTS (API RAW FULL CAMERA)
            # ----------------------------------------------------
            off_data = fetch_off_data(bc) if bc else {}
            off_front_url, off_back_url, off_note = extract_off_best_images(off_data, prod)
            
            front_bytes, front_source = None, ""
            back_bytes, back_source = None, ""

            if off_front_url:
                try:
                    r = requests.get(off_front_url, headers=HEADERS, timeout=8)
                    if r.status_code == 200 and len(r.content) > 30000:
                        im_t = Image.open(BytesIO(r.content))
                        if im_t.size[0] >= 800 or im_t.size[1] >= 800:
                            front_bytes = r.content
                            front_source = f"OpenFoodFacts_RAW ({off_front_url})"
                            print(f"  [+] Mặt trước tìm thấy trên OpenFoodFacts: {im_t.size[0]}x{im_t.size[1]} ({len(front_bytes)//1024} KB)")
                except Exception as e:
                    pass

            if off_back_url:
                try:
                    r = requests.get(off_back_url, headers=HEADERS, timeout=8)
                    if r.status_code == 200 and len(r.content) > 30000:
                        im_tb = Image.open(BytesIO(r.content))
                        if im_tb.size[0] >= 800 or im_tb.size[1] >= 800:
                            back_bytes = r.content
                            back_source = f"OpenFoodFacts_RAW ({off_back_url})"
                            print(f"  [+] Mặt sau tìm thấy trên OpenFoodFacts: {im_tb.size[0]}x{im_tb.size[1]} ({len(back_bytes)//1024} KB)")
                except Exception as e:
                    pass

            # ----------------------------------------------------
            # TẦNG 2: WEB SEARCH WITH STRICT IDENTITY VALIDATION
            # ----------------------------------------------------
            if not front_bytes or not back_bytes:
                print("  [*] Tra cứu bổ sung qua Search Engine có kiểm duyệt đối chiếu...")
                wf_bytes, wf_url, wb_bytes, wb_url = search_web_verified_images(page, prod)
                if not front_bytes and wf_bytes:
                    front_bytes = wf_bytes
                    front_source = f"Web_Verified ({wf_url})"
                    im_t = Image.open(BytesIO(wf_bytes))
                    print(f"  [+] Mặt trước tìm thấy trên Web: {im_t.size[0]}x{im_t.size[1]} ({len(wf_bytes)//1024} KB)")
                if not back_bytes and wb_bytes:
                    back_bytes = wb_bytes
                    back_source = f"Web_Verified ({wb_url})"
                    im_tb = Image.open(BytesIO(wb_bytes))
                    print(f"  [+] Mặt sau tìm thấy trên Web: {im_tb.size[0]}x{im_tb.size[1]} ({len(wb_bytes)//1024} KB)")

            # ----------------------------------------------------
            # HẬU KỲ STUDIO VÀ LƯU ẢNH
            # ----------------------------------------------------
            # 1. Xử lý Mặt trước (AI Remove Background + Nền trắng 1600x1600)
            if front_bytes:
                try:
                    processed_front = process_front_packshot(front_bytes, target_size=1600)
                    front_filename = f"{bc}_front.jpg"
                    front_path = os.path.join(FRONT_DIR, front_filename)
                    with open(front_path, "wb") as f_out:
                        f_out.write(processed_front)
                    
                    im_f_final = Image.open(BytesIO(processed_front))
                    item_res["front_status"] = "Thành công (Studio 1:1 HD)"
                    item_res["front_res"] = f"{im_f_final.size[0]}x{im_f_final.size[1]} ({len(processed_front)//1024} KB)"
                    item_res["front_src"] = front_source
                    item_res["front_path"] = front_path
                    print(f"  [✓] Mặt trước ĐÃ TÁCH NỀN TRẮNG STUDIO: {front_path}")
                except Exception as e:
                    print(f"  [!] Lỗi xử lý hậu kỳ mặt trước: {e}")

            # 2. Xử lý Mặt sau (Tăng nét chữ + Giữ nguyên nhãn phẳng)
            if back_bytes:
                try:
                    processed_back = process_back_label(back_bytes, target_max=2000)
                    back_filename = f"{bc}_back.jpg"
                    back_path = os.path.join(BACK_DIR, back_filename)
                    with open(back_path, "wb") as b_out:
                        b_out.write(processed_back)

                    im_b_final = Image.open(BytesIO(processed_back))
                    item_res["back_status"] = "Thành công (Label HD)"
                    item_res["back_res"] = f"{im_b_final.size[0]}x{im_b_final.size[1]} ({len(processed_back)//1024} KB)"
                    item_res["back_src"] = back_source
                    item_res["back_path"] = back_path
                    print(f"  [✓] Mặt sau ĐÃ TĂNG NÉT CHỮ LABEL: {back_path}")
                except Exception as e:
                    print(f"  [!] Lỗi xử lý hậu kỳ mặt sau: {e}")

            # Phân hạng Tier
            if item_res["front_status"].startswith("Thành công") and item_res["back_status"].startswith("Thành công"):
                item_res["tier"] = "Tier 1 (Verified 2 Mặt HD)"
            elif item_res["front_status"].startswith("Thành công"):
                item_res["tier"] = "Tier 2 (Có Mặt trước HD, Thiếu Mặt sau)"
                item_res["note"] = "Cần chụp nhãn mặt sau tại kho hoặc dùng text mô tả"
            else:
                item_res["tier"] = "Tier 3 (Retake Flagged)"
                item_res["note"] = "Cần nhân viên kho chụp bổ sung"

            results.append(item_res)
            time.sleep(1)

        browser.close()

    # ----------------------------------------------------
    # XUẤT FILE BÁO CÁO EXCEL ĐỐI SOÁT
    # ----------------------------------------------------
    print(f"\n[*] Đang tạo báo cáo đối soát: {REPORT_EXCEL}")
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Pilot 20 Two-Sided Audit"

    headers = [
        "STT", "Barcode", "Tên sản phẩm", "Thương hiệu", "Variant", "Trọng lượng",
        "Trạng thái Mặt trước", "Kích thước Mặt trước", "Nguồn Mặt trước", "Đường dẫn Mặt trước",
        "Trạng thái Mặt sau", "Kích thước Mặt sau", "Nguồn Mặt sau", "Đường dẫn Mặt sau",
        "Phân loại (Decision Tier)", "Ghi chú hành động"
    ]
    ws_out.append(headers)

    for r in results:
        ws_out.append([
            r["stt"], r["barcode"], r["name"], r["brand"], r["variant"], r["quantity"],
            r["front_status"], r["front_res"], r["front_src"], r["front_path"],
            r["back_status"], r["back_res"], r["back_src"], r["back_path"],
            r["tier"], r["note"]
        ])
    wb_out.save(REPORT_EXCEL)
    print(f"[+] Báo cáo Excel đã lưu tại: {REPORT_EXCEL}")

    # ----------------------------------------------------
    # TẠO TRANG HTML GALLERY TRỰC QUAN ĐỂ NGHIỆM THU
    # ----------------------------------------------------
    print(f"[*] Đang tạo giao diện xem trước HTML: {GALLERY_HTML}")
    html_cards = []
    for r in results:
        front_img_tag = f'<img src="../{r["front_path"].replace(chr(92), "/")}" alt="Front">' if r["front_path"] else '<div class="no-img">Chưa có ảnh mặt trước</div>'
        back_img_tag = f'<img src="../{r["back_path"].replace(chr(92), "/")}" alt="Back">' if r["back_path"] else '<div class="no-img">Chưa có ảnh nhãn mặt sau</div>'
        
        tier_color = "#10b981" if "Tier 1" in r["tier"] else ("#f59e0b" if "Tier 2" in r["tier"] else "#ef4444")

        card = f"""
        <div class="card">
            <div class="card-header">
                <span class="stt">#{r['stt']:02d}</span>
                <span class="barcode">{r['barcode']}</span>
                <span class="badge" style="background-color: {tier_color}">{r['tier']}</span>
            </div>
            <div class="card-title">{r['name']}</div>
            <div class="card-meta">
                <strong>Thương hiệu:</strong> {r['brand']} | 
                <strong>Biến thể:</strong> <span class="highlight">{r['variant']}</span> | 
                <strong>Trọng lượng:</strong> {r['quantity']}
            </div>
            <div class="images-container">
                <div class="img-box">
                    <div class="img-label">MẶT TRƯỚC (STUDIO NỀN TRẮNG 1:1)</div>
                    {front_img_tag}
                    <div class="img-info">{r['front_res']}<br><span class="src-text">{r['front_src'][:45]}...</span></div>
                </div>
                <div class="img-box">
                    <div class="img-label">MẶT SAU (LABEL THÀNH PHẦN)</div>
                    {back_img_tag}
                    <div class="img-info">{r['back_res']}<br><span class="src-text">{r['back_src'][:45]}...</span></div>
                </div>
            </div>
            <div class="card-footer">
                <strong>Ghi chú:</strong> {r['note'] or 'Đạt chuẩn 100%'}
            </div>
        </div>
        """
        html_cards.append(card)

    full_html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <title>Rabuta Pilot 20 — Nghiệm Thu Bộ Ảnh 2 Mặt Siêu Nét</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: #0f172a; color: #f8fafc; margin: 0; padding: 24px; }}
        h1 {{ text-align: center; color: #38bdf8; margin-bottom: 8px; }}
        p.subtitle {{ text-align: center; color: #94a3b8; margin-bottom: 32px; font-size: 15px; }}
        .grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(580px, 1fr)); gap: 24px; max-width: 1600px; margin: 0 auto; }}
        .card {{ background: #1e293b; border-radius: 12px; border: 1px solid #334155; padding: 16px; display: flex; flex-direction: column; gap: 12px; }}
        .card-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 1px solid #334155; padding-bottom: 8px; }}
        .stt {{ font-weight: bold; color: #38bdf8; font-size: 16px; }}
        .barcode {{ font-family: monospace; font-size: 15px; color: #a5f3fc; background: #082f49; padding: 2px 8px; border-radius: 4px; }}
        .badge {{ color: white; padding: 4px 10px; border-radius: 20px; font-size: 12px; font-weight: 600; }}
        .card-title {{ font-size: 16px; font-weight: 600; color: #f1f5f9; line-height: 1.4; }}
        .card-meta {{ font-size: 13px; color: #cbd5e1; }}
        .highlight {{ color: #fbbf24; font-weight: bold; }}
        .images-container {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; margin-top: 8px; }}
        .img-box {{ background: #0f172a; border: 1px solid #334155; border-radius: 8px; padding: 10px; display: flex; flex-direction: column; align-items: center; }}
        .img-label {{ font-size: 11px; font-weight: 700; color: #94a3b8; margin-bottom: 8px; text-transform: uppercase; letter-spacing: 0.5px; }}
        .img-box img {{ max-width: 100%; height: 260px; object-fit: contain; border-radius: 4px; background: #ffffff; cursor: zoom-in; transition: transform 0.2s; }}
        .img-box img:hover {{ transform: scale(1.03); }}
        .no-img {{ width: 100%; height: 260px; display: flex; align-items: center; justify-content: center; background: #1e293b; color: #ef4444; font-size: 13px; font-weight: 500; text-align: center; border-radius: 4px; }}
        .img-info {{ margin-top: 8px; font-size: 12px; color: #a5f3fc; text-align: center; line-height: 1.3; }}
        .src-text {{ font-size: 10px; color: #64748b; word-break: break-all; }}
        .card-footer {{ font-size: 12px; color: #94a3b8; border-top: 1px solid #334155; padding-top: 8px; }}
    </style>
</head>
<body>
    <h1>RABUTA PILOT 20 — BỘ ẢNH 2 MẶT WEBSHOP (HIGH-RES)</h1>
    <p class="subtitle">Đã áp dụng Identity Resolution (Khóa Barcode, Khớp Variant, Chống nhầm sản phẩm) & Studio Tách Nền Trắng 1:1</p>
    <div class="grid">
        {"".join(html_cards)}
    </div>
</body>
</html>
"""
    with open(GALLERY_HTML, "w", encoding="utf-8") as f_html:
        f_html.write(full_html)
    print(f"[+] ĐÃ TẠO THÀNH CÔNG HTML GALLERY TẠI: {GALLERY_HTML}")
    print("\n[✓] HOÀN TẤT THỬ NGHIỆM PILOT 20 SẢN PHẨM!")

if __name__ == "__main__":
    run_pilot_20()

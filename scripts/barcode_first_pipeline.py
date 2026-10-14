"""Barcode-first catalog pass.

1. Read every warehouse photo and keep the retail barcode.
2. Tag the photo with that barcode and the rotation that makes the barcode upright.
3. Cross-check against Rabuta.12.10.2026_original.xlsx.
4. From the best photo of each catalog barcode, export the back panel first,
   then the opposite panel as the webshop front.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from product_mapper.ingestion.image_loader import open_image_safely, save_image_high_quality
from product_mapper.vision.barcode_side_splitter import (
    SideBoxes,
    enhance_back,
    packshot_front,
    panel_halves,
    place_front_on_white,
    predict_split,
    read_retail_barcodes,
    rotate_upright,
    scan_file,
    separate_joined_packages,
    split_sides,
    transform_point,
)
from product_mapper.vision.barcode_validator import get_barcode_lookup_variants

EXCEL_SOURCE = ROOT / "source-import/new_source_from_customer/12-10-new-requirment/Rabuta.12.10.2026_original.xlsx"
SOURCE_DIRS = [
    ROOT / "source-import/G1_original",
    ROOT / "source-import/G2_original",
    ROOT / "source-import/G3_Original",
    ROOT / "source-import/G4-original",
    ROOT / "source-import/Nyt25.09.26",
]
OUTPUT_DIR = ROOT / "output/barcode_first"
INDEX_PATH = OUTPUT_DIR / "image_index.json"
REPORT_PATH = OUTPUT_DIR / "reconciliation.xlsx"
EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}

_TRUTH_CODES: set[str] = set()


def load_ground_truth(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    products = []
    for index, row in enumerate(rows[3:], start=1):
        barcode = _cell_barcode(row[0] if row else None)
        if not barcode:
            continue
        products.append(
            {
                "stt": index,
                "barcode": barcode,
                "name_en": str(row[1] or "").strip(),
                "name_da": str(row[2] or "").strip(),
                "brand": str(row[3] or "").strip(),
                "variant": str(row[7] or "").strip() if len(row) > 7 and row[7] else "",
                "net_quantity": str(row[10] or "").strip() if len(row) > 10 and row[10] else "",
                "net_unit": str(row[11] or "").strip() if len(row) > 11 and row[11] else "",
            }
        )
    return products


def _cell_barcode(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:
            return ""
        return str(int(round(value)))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return "".join(ch for ch in text if ch.isdigit())


def truth_lookup(products: list[dict]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for product in products:
        for variant in get_barcode_lookup_variants(product["barcode"]):
            lookup[variant] = product["barcode"]
    return lookup


def list_images() -> list[Path]:
    files: list[Path] = []
    for folder in SOURCE_DIRS:
        for path in sorted(folder.rglob("*")):
            if path.is_file() and path.suffix.lower() in EXTENSIONS:
                files.append(path)
    return files


def _init_worker(codes: list[str]) -> None:
    global _TRUTH_CODES
    _TRUTH_CODES = set(codes)


def _scan_task(path_str: str) -> dict:
    try:
        record = scan_file(Path(path_str), _TRUTH_CODES)
    except Exception as exc:
        return {"path": path_str, "barcode": None, "error": str(exc), "codes": []}
    record["folder"] = Path(path_str).parent.name
    return record


def scan_all(files: list[Path], lookup: dict[str, str], workers: int) -> list[dict]:
    records: list[dict] = []
    total = len(files)
    with ProcessPoolExecutor(max_workers=workers, initializer=_init_worker, initargs=(list(lookup),)) as pool:
        futures = [pool.submit(_scan_task, str(path)) for path in files]
        for done, future in enumerate(as_completed(futures), start=1):
            records.append(future.result())
            if done % 100 == 0 or done == total:
                found = sum(1 for row in records if row.get("barcode"))
                print(f"  scanned {done}/{total}  with barcode {found}", flush=True)
    for record in records:
        barcode = record.get("barcode")
        record["excel_barcode"] = lookup.get(barcode or "", None)
        if barcode and record["excel_barcode"] is None:
            for variant in get_barcode_lookup_variants(barcode):
                if variant in lookup:
                    record["excel_barcode"] = lookup[variant]
                    break
    return records


def _rank(record: dict) -> float:
    if not record.get("barcode") or "center" not in record:
        return -1
    mode = predict_split(
        int(record["width"]),
        int(record["height"]),
        float(record["center"][0]),
        float(record["center"][1]),
        int(record["upright_ccw"]),
    )
    front_bonus = 0 if mode == "back_only" else 1_000_000
    return front_bonus + float(record.get("area") or 0) + (int(record["width"]) * int(record["height"]) / 1000.0)


def open_at(record: dict, ccw: int) -> tuple[Image.Image, float, float]:
    image = open_image_safely(record["path"])
    if record.get("mirrored"):
        image = ImageOps.mirror(image)
    width, height = image.size
    upright_x, upright_y = transform_point(float(record["center"][0]), float(record["center"][1]), width, height, ccw)
    return rotate_upright(image, ccw), upright_x, upright_y


def open_upright(record: dict) -> tuple[Image.Image, float, float]:
    return open_at(record, int(record["upright_ccw"]))


def _cut_through_standing_product(sides, image_width: int, image_height: int) -> bool:
    """A label band or a thin cardboard sliver, not a whole package."""
    boxes = [sides.back]
    if sides.front is not None:
        boxes.append(sides.front)
    for box in boxes:
        box_w = box[2] - box[0]
        box_h = box[3] - box[1]
        if box_h <= 0 or box_w <= 0:
            return True
        wide_band = (box_w / box_h) >= 1.65 and box_h < image_height * 0.65
        # A blank edge sliver stops short of the frame. A standing bottle reaches the frame.
        tall_sliver = (box_w / box_h) <= 0.45 and box_w < image_width * 0.42 and box_h < image_height * 0.72
        if wide_band or tall_sliver:
            return True
    return False


def _panel_shape(box: tuple[int, int, int, int], image_height: int) -> tuple[float, float]:
    box_w = box[2] - box[0]
    box_h = box[3] - box[1]
    return box_w / max(1, box_h), box_h / max(1, image_height)


def _photo_orientation_is_upright(rotated_sides, rotated_image, natural_sides, natural_image) -> bool:
    """True when the camera framing already has a standing package upright."""
    rotated_box = rotated_sides.front or rotated_sides.back
    natural_box = natural_sides.front or natural_sides.back
    rotated_aspect, _rotated_fill = _panel_shape(rotated_box, rotated_image.size[1])
    natural_aspect, natural_fill = _panel_shape(natural_box, natural_image.size[1])
    rotated_frame_is_wide = rotated_image.size[0] > rotated_image.size[1] * 1.05
    # Two flat packets already split by the barcode rotation must stay that way.
    # A landscape camera frame is not, by itself, a bottle laid on its side.
    if rotated_sides.front is not None:
        front_aspect, front_fill = _panel_shape(rotated_sides.front, rotated_image.size[1])
        if front_aspect < 1.20 and front_fill >= 0.50:
            return 0.72 <= rotated_aspect <= 1.40 and 0.72 <= natural_aspect <= 1.40
    bottle_was_laid_down = natural_fill >= 0.70 and natural_aspect < 0.90 and (rotated_aspect >= 1.35 or rotated_frame_is_wide)
    if bottle_was_laid_down:
        return True
    # A short can or box. Spinning it to follow a side barcode turns the label on end.
    return 0.72 <= rotated_aspect <= 1.40 and 0.72 <= natural_aspect <= 1.40


def _barcode_still_present(image: Image.Image, expected: str) -> bool:
    expected_forms = set(get_barcode_lookup_variants(expected))
    for hit in read_retail_barcodes(image):
        if hit.text in expected_forms or expected in set(get_barcode_lookup_variants(hit.text)):
            return True
    return False


def _front_kept_back(image: Image.Image, expected: str) -> bool:
    expected_forms = set(get_barcode_lookup_variants(expected))
    width, height = image.size
    for hit in read_retail_barcodes(image):
        forms = set(get_barcode_lookup_variants(hit.text))
        same = hit.text in expected_forms or expected in forms
        if same and hit.area > (width * height * 0.04):
            return True
    return False


def _crop(image: Image.Image, box: tuple[int, int, int, int]) -> Image.Image:
    return image.crop(box)


def _apply_joined_split(image: Image.Image, sides: SideBoxes, center_x: float, center_y: float, ccw: int) -> SideBoxes:
    """Keep one package on each side when two packages were left in the same crop."""
    # Only split a back when the first cut kept the whole photo. Cutting an
    # existing back panel follows the barcode sticker instead of the label.
    if sides.front is None:
        separated = separate_joined_packages(image, sides.back, center_x, center_y)
        if separated is not None:
            closer, farther = separated
            sides = SideBoxes(mode="back_left", back=closer, front=farther, upright_ccw=ccw)
    if sides.front is not None:
        front_pair = separate_joined_packages(image, sides.front, center_x, center_y)
        if front_pair is not None:
            _closer, farther = front_pair
            sides = SideBoxes(mode=sides.mode, back=sides.back, front=farther, upright_ccw=ccw)
    return sides


def export_product(record: dict, barcode: str, output_dir: Path | None = None) -> dict:
    ccw = int(record["upright_ccw"])
    image, center_x, center_y = open_at(record, ccw)
    sides = split_sides(image, center_x, center_y, ccw)
    # A barcode printed down the label is vertical while the bottle already stands.
    # Turning the photo to make that barcode horizontal lays the bottle on its side.
    if ccw in (90, 270):
        natural, natural_x, natural_y = open_at(record, 0)
        natural_sides = split_sides(natural, natural_x, natural_y, 0)
        slice_fixed = _cut_through_standing_product(sides, image.size[0], image.size[1]) and not _cut_through_standing_product(
            natural_sides, natural.size[0], natural.size[1]
        )
        stands = _photo_orientation_is_upright(sides, image, natural_sides, natural)
        # A box photographed with the front face already upright. The barcode on the
        # spine is vertical; rotating to level it turns that side panel into a strip.
        natural_back_w = natural_sides.back[2] - natural_sides.back[0]
        natural_back_h = natural_sides.back[3] - natural_sides.back[1]
        rotated_back_w = sides.back[2] - sides.back[0]
        rotated_back_h = sides.back[3] - sides.back[1]
        spine_stays_upright = (
            natural.size[0] > natural.size[1] * 1.05
            and natural_sides.front is not None
            and natural_back_h > natural.size[1] * 0.70
            and 0.22 <= natural_back_w / max(1, natural.size[0]) <= 0.48
            and rotated_back_h > 0
            and (rotated_back_w / rotated_back_h) <= 0.40
        )
        # Two trays or bags already stacked upright. Rotating with the barcode
        # leaves both packages in one frame, so keep the horizontal cut.
        natural_front_h = 0
        if natural_sides.front is not None:
            natural_front_h = natural_sides.front[3] - natural_sides.front[1]
        stacked_stays_upright = (
            natural_sides.mode in ("back_top", "back_bottom")
            and natural_front_h > natural.size[1] * 0.35
            and natural_back_h > natural.size[1] * 0.35
            and (sides.mode == "back_only" or (rotated_back_h > 0 and rotated_back_w / rotated_back_h <= 0.40))
        )
        if slice_fixed or stands or spine_stays_upright or stacked_stays_upright:
            natural_back = _crop(natural, natural_sides.back)
            barcode_on_back = _barcode_still_present(natural_back, barcode)
            barcode_on_photo = barcode_on_back or _barcode_still_present(natural, barcode)
            if barcode_on_photo:
                image, center_x, center_y, sides, ccw = natural, natural_x, natural_y, natural_sides, 0
    sides = _apply_joined_split(image, sides, center_x, center_y, ccw)
    dest_root = output_dir or OUTPUT_DIR
    back = enhance_back(_crop(image, sides.back))
    back_path = dest_root / "back" / f"{barcode}_back.jpg"
    save_image_high_quality(back, back_path, quality=95, save_webp=False)
    back_ok = _barcode_still_present(back, barcode)
    note = ""
    if ccw != int(record["upright_ccw"]):
        note = "Kept the photo upright; the barcode is printed down the label."
    if not back_ok:
        note = (note + " " if note else "") + "Back crop lost the barcode; kept the barcode half instead."
        half = panel_halves(image, center_x, center_y, ccw)
        half_w = half.back[2] - half.back[0]
        half_h = half.back[3] - half.back[1]
        half_is_band = half_h > 0 and (half_w / half_h) >= 1.85 and half_h < image.size[1] * 0.45
        if not half_is_band:
            back = enhance_back(_crop(image, half.back))
            save_image_high_quality(back, back_path, quality=95, save_webp=False)
            back_ok = _barcode_still_present(back, barcode)
            sides = half
            sides = _apply_joined_split(image, sides, center_x, center_y, ccw)
            back = enhance_back(_crop(image, sides.back))
            save_image_high_quality(back, back_path, quality=95, save_webp=False)
            back_ok = _barcode_still_present(back, barcode)
        if not back_ok and sides.front is None:
            back = enhance_back(image)
            save_image_high_quality(back, back_path, quality=95, save_webp=False)
            back_ok = _barcode_still_present(back, barcode)
            sides = type(sides)(mode="back_only", back=(0, 0, image.size[0], image.size[1]), front=sides.front, upright_ccw=ccw)
            sides = _apply_joined_split(image, sides, center_x, center_y, ccw)
            if sides.front is not None:
                back = enhance_back(_crop(image, sides.back))
                save_image_high_quality(back, back_path, quality=95, save_webp=False)
                back_ok = _barcode_still_present(back, barcode)
            note = (note + " " if note else "") + "Back crop lost the barcode; kept the whole upright photo."

    front_path = ""
    front_ok = False
    front_size = ""
    front_crop = _crop(image, sides.front) if sides.front is not None else None
    if front_crop is not None and _front_kept_back(front_crop, barcode):
        note = (note + " " if note else "") + "Front crop still contained the back barcode."
        front_crop = None
    if front_crop is None and sides.mode == "back_only":
        # One standing package, not two panels. Keep the whole package instead of a neck strip.
        front_crop = _crop(image, sides.back)
        note = (note + " " if note else "") + "Single package kept whole as the front."
    if front_crop is not None:
        front = packshot_front(front_crop)
        dest = dest_root / "front" / f"{barcode}_front.jpg"
        save_image_high_quality(front, dest, quality=95, save_webp=False)
        front_path = str(dest)
        front_ok = True
        front_size = f"{front.size[0]}x{front.size[1]}"

    return {
        "mode": sides.mode,
        "rotation": ccw,
        "source": record["path"],
        "back_path": str(back_path),
        "back_ok": back_ok,
        "back_size": f"{back.size[0]}x{back.size[1]}",
        "front_path": front_path,
        "front_ok": front_ok,
        "front_size": front_size,
        "note": note.strip(),
    }


def _export_task(payload: tuple[dict, str]) -> dict:
    record, barcode = payload
    try:
        result = export_product(record, barcode)
        result["barcode"] = barcode
        result["ok"] = True
    except Exception as exc:
        result = {"barcode": barcode, "ok": False, "error": str(exc), "source": record.get("path", "")}
    return result


def write_report(products: list[dict], records: list[dict], exports: dict[str, dict]) -> None:
    by_excel: dict[str, list[dict]] = defaultdict(list)
    no_barcode = []
    extra = []
    for record in records:
        if not record.get("barcode"):
            no_barcode.append(record)
            continue
        if record.get("excel_barcode"):
            by_excel[record["excel_barcode"]].append(record)
        else:
            extra.append(record)

    workbook = openpyxl.Workbook()
    cover = workbook.active
    cover.title = "Catalog match"
    headers = [
        "STT",
        "Barcode",
        "Name",
        "Brand",
        "Variant",
        "Quantity",
        "Photos with this barcode",
        "Chosen photo",
        "Upright rotation",
        "Split",
        "Back QC",
        "Back size",
        "Front QC",
        "Front size",
        "Back file",
        "Front file",
        "Note",
    ]
    cover.append(headers)
    matched = 0
    both = 0
    for product in products:
        photos = by_excel.get(product["barcode"], [])
        exported = exports.get(product["barcode"], {})
        if photos:
            matched += 1
        if exported.get("front_ok") and exported.get("back_ok"):
            both += 1
        qty = f"{product['net_quantity']} {product['net_unit']}".strip()
        cover.append(
            [
                product["stt"],
                product["barcode"],
                product["name_en"],
                product["brand"],
                product["variant"],
                qty,
                len(photos),
                Path(exported.get("source", "")).name if exported else "",
                exported.get("rotation", ""),
                exported.get("mode", ""),
                "Pass" if exported.get("back_ok") else ("Fail" if exported else ""),
                exported.get("back_size", ""),
                "Pass" if exported.get("front_ok") else ("Fail" if exported else ""),
                exported.get("front_size", ""),
                exported.get("back_path", ""),
                exported.get("front_path", ""),
                exported.get("note") or exported.get("error") or "",
            ]
        )

    missing = workbook.create_sheet("No barcode")
    missing.append(["Folder", "File", "Error"])
    for record in no_barcode:
        path = Path(record["path"])
        missing.append([record.get("folder") or path.parent.name, path.name, record.get("error", "")])

    extra_sheet = workbook.create_sheet("Barcode not in Excel")
    extra_sheet.append(["Barcode", "Folder", "File", "Other codes"])
    for record in extra:
        path = Path(record["path"])
        extra_sheet.append([record.get("barcode"), record.get("folder") or path.parent.name, path.name, ", ".join(record.get("codes") or [])])

    summary = workbook.create_sheet("Summary", 0)
    summary.append(["Metric", "Value"])
    summary.append(["Excel products", len(products)])
    summary.append(["Photos scanned", len(records)])
    summary.append(["Photos with a valid barcode", sum(1 for row in records if row.get("barcode"))])
    summary.append(["Photos with no barcode", len(no_barcode)])
    summary.append(["Excel products with at least one photo", matched])
    summary.append(["Excel products with back and front QC pass", both])
    summary.append(["Photos whose barcode is not in the Excel", len(extra)])
    _style(cover)
    _style(missing)
    _style(extra_sheet)
    _style(summary)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(REPORT_PATH)


def _style(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def run_check() -> None:
    samples = [
        ROOT / "source-import/G1_original/IMG_0002.jpeg",
        ROOT / "source-import/G1_original/IMG_0003.jpeg",
        ROOT / "source-import/G1_original/IMG_0004.jpeg",
    ]
    sample_dir = OUTPUT_DIR / "samples"
    for path in samples:
        record = scan_file(path, None)
        print(
            f"{path.name} barcode={record.get('barcode')} rot={record.get('upright_ccw')} "
            f"codes={record.get('codes')} center={record.get('center')}",
            flush=True,
        )
        if not record.get("barcode"):
            continue
        mode = predict_split(record["width"], record["height"], record["center"][0], record["center"][1], record["upright_ccw"])
        print(f"  predicted {mode}", flush=True)
        image, center_x, center_y = open_upright(record)
        sides = split_sides(image, center_x, center_y, int(record["upright_ccw"]))
        print(f"  split {sides.mode} back={sides.back} front={sides.front}", flush=True)
        back = enhance_back(image.crop(sides.back))
        back_file = sample_dir / f"{path.stem}_back.jpg"
        save_image_high_quality(back, back_file, quality=95, save_webp=False)
        print(f"  back barcode still readable={_barcode_still_present(back, record['barcode'])} size={back.size}", flush=True)
        if sides.front is not None:
            front = place_front_on_white(image.crop(sides.front))
            front_file = sample_dir / f"{path.stem}_front.jpg"
            save_image_high_quality(front, front_file, quality=95, save_webp=False)
            print(
                f"  front kept back barcode={_front_kept_back(front, record['barcode'])} size={front.size}",
                flush=True,
            )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=("index", "extract", "all"), default="all")
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--workers", type=int, default=max(1, min(8, os.cpu_count() or 4)))
    args = parser.parse_args()

    if args.check:
        run_check()
        return

    products = load_ground_truth(EXCEL_SOURCE)
    lookup = truth_lookup(products)
    print(f"Excel products with barcode: {len(products)}", flush=True)

    if args.stage in ("index", "all"):
        files = list_images()
        if args.limit:
            files = files[: args.limit]
        print(f"Scanning {len(files)} photos with {args.workers} workers", flush=True)
        records = scan_all(files, lookup, args.workers)
        INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)
        INDEX_PATH.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        print(f"Index written: {INDEX_PATH}", flush=True)
    else:
        records = json.loads(INDEX_PATH.read_text(encoding="utf-8"))

    exports: dict[str, dict] = {}
    if args.stage in ("extract", "all"):
        grouped: dict[str, list[dict]] = defaultdict(list)
        for record in records:
            excel_barcode = record.get("excel_barcode")
            if excel_barcode:
                grouped[excel_barcode].append(record)
        jobs = []
        for barcode, group in grouped.items():
            best = max(group, key=_rank)
            jobs.append((best, barcode))
        print(f"Exporting back then front for {len(jobs)} catalog barcodes", flush=True)
        with ProcessPoolExecutor(max_workers=max(1, min(4, args.workers))) as pool:
            futures = [pool.submit(_export_task, job) for job in jobs]
            total = len(futures)
            for done, future in enumerate(as_completed(futures), start=1):
                result = future.result()
                exports[result["barcode"]] = result
                if done % 50 == 0 or done == total:
                    print(f"  exported {done}/{total}", flush=True)
        (OUTPUT_DIR / "export_results.json").write_text(json.dumps(exports, ensure_ascii=False), encoding="utf-8")

    write_report(products, records, exports)
    print(f"Reconciliation: {REPORT_PATH}", flush=True)


if __name__ == "__main__":
    main()

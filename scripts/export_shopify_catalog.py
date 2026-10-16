"""
Export Standard 28-Column Shopify Catalog CSV from Master Catalog.
Classifies products based on Rabuta Category Level 1 and Ma_nhom_hang_SKU_3digits.xlsx:
- Online (Web-eligible): 12 categories (1,090 products) -> Published=TRUE, Status=active
- Offline (Nej / In-store only: Frostvarer, Ris, Fisk, Rejer): 4 categories (261 products) -> Published=FALSE, Status=draft, tagged offline_only
Enriches metadata with nutrition tables and storage instructions from Esbjerg 89-cols master.
"""

from __future__ import annotations

import csv
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# Standard 28 Shopify Import Columns
SHOPIFY_28_COLUMNS = [
    "Handle",
    "Title",
    "Body (HTML)",
    "Vendor",
    "Product Category",
    "Type",
    "Tags",
    "Published",
    "Option1 Name",
    "Option1 Value",
    "Variant SKU",
    "Variant Grams",
    "Variant Inventory Tracker",
    "Variant Inventory Qty",
    "Variant Inventory Policy",
    "Variant Fulfillment Service",
    "Variant Price",
    "Variant Compare At Price",
    "Variant Requires Shipping",
    "Variant Taxable",
    "Variant Barcode",
    "Image Src",
    "Image Position",
    "Image Alt Text",
    "Gift Card",
    "SEO Title",
    "SEO Description",
    "Status",
]

# Classification of 16 POS categories from Ma_nhom_hang_SKU_3digits.xlsx
# Online != 'Nej' -> Web-eligible
CATEGORY_RULES: Dict[str, Dict[str, Any]] = {
    "Konserves": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk",
        "type": "Konserves & Kokosmælk",
    },
    "Frugt": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Grøntsager & Frugt",
        "type": "Frugt",
    },
    "Krydderier": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Krydderier & Karry",
        "type": "Krydderier & Karry",
    },
    "Diverse": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Andre asiatiske varer",
        "type": "Andre asiatiske varer",
    },
    "Saucer": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce",
        "type": "Saucer & Sojasauce",
    },
    "Mælk": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Mælk & Mejeriprodukter",
        "type": "Mælk",
    },
    "Mel": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer",
        "type": "Ris, Mel & Tørvarer",
    },
    "Snack": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Snack",
        "type": "Snack",
    },
    "Frostvarer": {
        "is_online": False,  # Hàng đông lạnh - không ship bưu điện thông thường
        "product_category": "Food, Beverages & Tobacco > Food Items > Frostvarer",
        "type": "Frostvarer",
    },
    "Fisk": {
        "is_online": False,  # Cá tươi / đông lạnh
        "product_category": "Food, Beverages & Tobacco > Food Items > Fisk & Skaldyr",
        "type": "Fisk",
    },
    "Grøntsager": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Grøntsager & Frugt",
        "type": "Grøntsager",
    },
    "Rejer": {
        "is_online": False,  # Tôm tươi / đông lạnh
        "product_category": "Food, Beverages & Tobacco > Food Items > Fisk & Skaldyr",
        "type": "Rejer",
    },
    "Ris": {
        "is_online": False,  # Bao gạo nặng / chính sách shop không bán online
        "product_category": "Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer",
        "type": "Ris, Mel & Tørvarer",
    },
    "Drikkevarer": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin",
        "type": "Drikkevarer & Vin",
    },
    "Nudler": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Nudler, Ris & Mel",
        "type": "Nudler",
    },
    "Bønner": {
        "is_online": True,
        "product_category": "Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer",
        "type": "Ris, Mel & Tørvarer",
    },
}

KEYWORD_CATEGORY_FALLBACK = {
    "sauce": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "soja": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "soy": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "curry": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "karry": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "pasta": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "paste": ("Food, Beverages & Tobacco > Food Items > Saucer & Sojasauce", "Saucer & Sojasauce"),
    "krydderi": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "spice": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "sesam": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "eddike": ("Food, Beverages & Tobacco > Food Items > Krydderier & Karry", "Krydderier & Karry"),
    "snack": ("Food, Beverages & Tobacco > Food Items > Snack", "Snack"),
    "chip": ("Food, Beverages & Tobacco > Food Items > Snack", "Snack"),
    "slik": ("Food, Beverages & Tobacco > Food Items > Snack", "Snack"),
    "nori": ("Food, Beverages & Tobacco > Food Items > Snack", "Snack"),
    "ris": ("Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer", "Ris, Mel & Tørvarer"),
    "rice": ("Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer", "Ris, Mel & Tørvarer"),
    "bønner": ("Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer", "Ris, Mel & Tørvarer"),
    "beans": ("Food, Beverages & Tobacco > Food Items > Ris, Mel & Tørvarer", "Ris, Mel & Tørvarer"),
    "konserves": ("Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk", "Konserves & Kokosmælk"),
    "canned": ("Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk", "Konserves & Kokosmælk"),
    "kokos": ("Food, Beverages & Tobacco > Food Items > Konserves & Kokosmælk", "Konserves & Kokosmælk"),
    "drikke": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Vin"),
    "drink": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Vin"),
    "juice": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Vin"),
    "sake": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Vin"),
    "vin": ("Food, Beverages & Tobacco > Food Items > Drikkevarer & Vin", "Drikkevarer & Vin"),
}


def slugify(text: str) -> str:
    """Generates a URL-safe kebab-case slug handling Danish characters."""
    text = str(text or "").strip()
    text = text.replace("æ", "ae").replace("Æ", "ae")
    text = text.replace("ø", "oe").replace("Ø", "oe")
    text = text.replace("å", "aa").replace("Å", "aa")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^\w\s-]", "", text.lower())
    text = re.sub(r"[-\s]+", "-", text).strip("-")
    return text


def clean_barcode(val: Any) -> str:
    """Formats barcode cleanly as a pure numeric string without scientific notation."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if "e+" in s.lower() or "e-" in s.lower():
        try:
            s = f"{int(float(s))}"
        except Exception:
            pass
    return re.sub(r"[^\w]", "", s)


def safe_str(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s.lower() in {"none", "nan", "null"} else s


def load_esbjerg_enrichment_data(esbjerg_file: Path) -> Dict[str, Dict[str, Any]]:
    """Loads rich nutrition tables, preparation, and storage data from the 89-column catalog."""
    enrichment: Dict[str, Dict[str, Any]] = {}
    if not esbjerg_file.exists():
        print(f"[WARN] Esbjerg enrichment file not found: {esbjerg_file}")
        return enrichment

    print(f"[LOAD] Reading nutrition & enrichment data: {esbjerg_file.name}...")
    wb = openpyxl.load_workbook(esbjerg_file, data_only=True)
    ws = wb.active

    for r in range(4, ws.max_row + 1):
        bc = clean_barcode(ws.cell(r, 3).value)
        if not bc:
            continue
        data = {
            "energy_kj": ws.cell(r, 42).value,
            "energy_kcal": ws.cell(r, 43).value,
            "fat": ws.cell(r, 44).value,
            "sat_fat": ws.cell(r, 45).value,
            "carbs": ws.cell(r, 46).value,
            "sugar": ws.cell(r, 47).value,
            "protein": ws.cell(r, 49).value,
            "salt": ws.cell(r, 50).value,
            "prep_da": ws.cell(r, 54).value,
            "storage_da": ws.cell(r, 57).value,
            "sku": ws.cell(r, 89).value,
            "price": ws.cell(r, 88).value,
        }
        enrichment[bc] = data
        if bc.startswith("0"):
            enrichment[bc.lstrip("0")] = data

    wb.close()
    print(f"[LOAD] Indexed {len(enrichment)} products with full nutrition metadata.")
    return enrichment


def build_nutrition_table(nut: Dict[str, Any]) -> str:
    """Builds the HTML table for Danish nutritional declarations."""
    ekj = nut.get("energy_kj")
    ekcal = nut.get("energy_kcal")
    fat = nut.get("fat")
    sat_fat = nut.get("sat_fat")
    carbs = nut.get("carbs")
    sugar = nut.get("sugar")
    protein = nut.get("protein")
    salt = nut.get("salt")

    has_any = any(v is not None and str(v).strip() != "" for v in [ekj, ekcal, fat, carbs, protein, salt])
    if not has_any:
        return ""

    rows = []
    if ekj is not None or ekcal is not None:
        e_str = []
        if ekj is not None and str(ekj).strip():
            e_str.append(f"{ekj} kJ")
        if ekcal is not None and str(ekcal).strip():
            e_str.append(f"{ekcal} kcal")
        if e_str:
            rows.append(f"<tr><td>Energi</td><td>{' / '.join(e_str)}</td></tr>")

    if fat is not None and str(fat).strip():
        rows.append(f"<tr><td>Fedt</td><td>{fat} g</td></tr>")
    if sat_fat is not None and str(sat_fat).strip():
        rows.append(f"<tr><td>&nbsp;&nbsp;heraf mættede fedtsyrer</td><td>{sat_fat} g</td></tr>")
    if carbs is not None and str(carbs).strip():
        rows.append(f"<tr><td>Kulhydrat</td><td>{carbs} g</td></tr>")
    if sugar is not None and str(sugar).strip():
        rows.append(f"<tr><td>&nbsp;&nbsp;heraf sukkerarter</td><td>{sugar} g</td></tr>")
    if protein is not None and str(protein).strip():
        rows.append(f"<tr><td>Protein</td><td>{protein} g</td></tr>")
    if salt is not None and str(salt).strip():
        rows.append(f"<tr><td>Salt</td><td>{salt} g</td></tr>")

    if not rows:
        return ""
    return "<h3>Næringsindhold pr. 100 g</h3><table>\n" + "\n".join(rows) + "\n</table>"


def generate_product_body_html(
    brand: str,
    net_weight: str,
    net_unit: str,
    origin: str,
    ing_da: str,
    all_da: str,
    nutrition_table_html: str,
    prep_da: str,
    storage_da: str,
) -> str:
    """Generates localized Danish HTML product description conforming to EU Regulation 1169/2011."""
    parts = []

    # 1. Header Meta Line
    meta = []
    if brand:
        meta.append(f"Mærke: {brand}")
    if net_weight:
        unit = net_unit or "g"
        meta.append(f"Nettoindhold: {net_weight} {unit}")
    elif net_unit:
        meta.append(f"Nettoindhold: {net_unit}")
    if origin:
        meta.append(f"Oprindelsesland: {origin}")
    if meta:
        parts.append(f"<p>{' · '.join(meta)}</p>")

    # 2. Ingredients
    if ing_da:
        parts.append(f"<h3>Ingredienser</h3><p>{ing_da}</p>")

    # 3. Allergens
    if all_da:
        parts.append(f"<h3>Allergener</h3><p>{all_da}</p>")

    # 4. Nutrition table
    if nutrition_table_html:
        parts.append(nutrition_table_html)

    # 5. Preparation
    if prep_da:
        parts.append(f"<h3>Tilberedning</h3><p>{prep_da}</p>")

    # 6. Storage
    opb = storage_da or "Opbevares tørt og køligt."
    parts.append(f"<h3>Opbevaring</h3><p>{opb}</p>")

    return "".join(parts)


def export_shopify_files(
    master_excel: Path,
    esbjerg_excel: Path,
    output_dir: Path,
) -> Tuple[Path, Path, Path]:
    """Generates Shopify CSV deliverables with full category classification."""
    output_dir.mkdir(parents=True, exist_ok=True)
    enrichment_map = load_esbjerg_enrichment_data(esbjerg_excel)

    print(f"\n[LOAD] Reading Master Catalog: {master_excel.name}...")
    wb = openpyxl.load_workbook(master_excel, data_only=True)
    ws = wb.active

    all_rows: List[Dict[str, Any]] = []
    online_rows: List[Dict[str, Any]] = []
    offline_rows: List[Dict[str, Any]] = []

    stats = Counter()

    for r in range(4, ws.max_row + 1):
        bc = clean_barcode(ws.cell(r, 1).value)
        norm_bc = bc.lstrip("0")
        en_name = safe_str(ws.cell(r, 2).value)
        da_name = safe_str(ws.cell(r, 3).value)
        brand = safe_str(ws.cell(r, 4).value) or "Khác"
        raw_name = safe_str(ws.cell(r, 6).value)
        net_qty = safe_str(ws.cell(r, 10).value)
        net_unit = safe_str(ws.cell(r, 11).value) or "g"
        origin = safe_str(ws.cell(r, 12).value) or "Asien"
        ing_da = safe_str(ws.cell(r, 15).value) or safe_str(ws.cell(r, 14).value)
        all_da = safe_str(ws.cell(r, 19).value) or safe_str(ws.cell(r, 18).value)
        short_desc = safe_str(ws.cell(r, 27).value) or safe_str(ws.cell(r, 26).value)
        price_val = ws.cell(r, 29).value
        sku = safe_str(ws.cell(r, 31).value)
        cat_l1 = safe_str(ws.cell(r, 32).value) or "Diverse"
        cat_l2 = safe_str(ws.cell(r, 33).value)

        # Enrichment from Esbjerg 89 cols if available
        nut_data = enrichment_map.get(bc) or enrichment_map.get(norm_bc) or {}
        nut_table = build_nutrition_table(nut_data)
        prep_da = safe_str(nut_data.get("prep_da"))
        storage_da = safe_str(nut_data.get("storage_da"))
        if not sku and nut_data.get("sku"):
            sku = safe_str(nut_data.get("sku"))
        if not price_val and nut_data.get("price"):
            price_val = nut_data.get("price")

        # Category Rule Determination
        rule = CATEGORY_RULES.get(cat_l1)
        if rule:
            is_online = rule["is_online"]
            product_category = rule["product_category"]
            ptype = rule["type"]
        else:
            # Fallback keyword match
            combined = f"{cat_l1} {da_name} {en_name}".lower()
            matched = False
            for kw, (pcat, pt) in KEYWORD_CATEGORY_FALLBACK.items():
                if kw in combined:
                    product_category, ptype = pcat, pt
                    matched = True
                    break
            if not matched:
                product_category = "Food, Beverages & Tobacco > Food Items > Andre asiatiske varer"
                ptype = "Andre asiatiske varer"
            is_online = cat_l1 not in {"Frostvarer", "Fisk", "Rejer", "Ris"}

        # Title Resolution
        title = da_name or en_name or raw_name or (f"Asiatisk produkt {bc}" if bc else "Asiatisk produkt")

        # Handle Resolution
        handle = slugify(title)
        if not handle or handle == "asiatisk-produkt":
            handle = f"asiatisk-produkt-{bc}" if bc else f"produkt-{sku.lower()}"

        # Weight in grams
        try:
            grams = int(float(net_qty))
        except Exception:
            grams = 100

        # Price formatting
        try:
            price_float = float(price_val)
            price_str = f"{price_float:.2f}"
        except Exception:
            price_str = "25.00" if "Sauce" in ptype or "Karry" in ptype else "15.00"

        # Tags
        tags_list = [ptype, brand, origin, "g1_import"]
        if is_online:
            tags_list.append("online")
        else:
            tags_list.append("offline_only")
            tags_list.append("instore_only")
        tags_str = ", ".join(dict.fromkeys(t for t in tags_list if t))

        # HTML Body
        body_html = generate_product_body_html(
            brand=brand,
            net_weight=net_qty,
            net_unit=net_unit,
            origin=origin,
            ing_da=ing_da,
            all_da=all_da,
            nutrition_table_html=nut_table,
            prep_da=prep_da,
            storage_da=storage_da,
        )

        # Image Src (pointing directly to Shopify canonical files URL)
        image_src = f"https://rabuta.dk/cdn/shop/files/{bc}.jpg" if bc else ""

        # Status & Published flags according to web eligibility
        published_str = "TRUE" if is_online else "FALSE"
        status_str = "active" if is_online else "draft"
        requires_shipping = "TRUE" if is_online else "FALSE"

        seo_title = title
        seo_desc = f"Køb {title} online hos Rabuta.dk. Hurtig levering i hele Danmark."

        row = {
            "Handle": handle,
            "Title": title,
            "Body (HTML)": body_html,
            "Vendor": brand,
            "Product Category": product_category,
            "Type": ptype,
            "Tags": tags_str,
            "Published": published_str,
            "Option1 Name": "Title",
            "Option1 Value": "Default Title",
            "Variant SKU": sku,
            "Variant Grams": grams,
            "Variant Inventory Tracker": "shopify",
            "Variant Inventory Qty": 100,
            "Variant Inventory Policy": "deny",
            "Variant Fulfillment Service": "manual",
            "Variant Price": price_str,
            "Variant Compare At Price": "",
            "Variant Requires Shipping": requires_shipping,
            "Variant Taxable": "TRUE",
            "Variant Barcode": bc,
            "Image Src": image_src,
            "Image Position": 1,
            "Image Alt Text": title,
            "Gift Card": "FALSE",
            "SEO Title": seo_title,
            "SEO Description": seo_desc,
            "Status": status_str,
        }

        all_rows.append(row)
        if is_online:
            online_rows.append(row)
            stats["ONLINE"] += 1
        else:
            offline_rows.append(row)
            stats["OFFLINE"] += 1
        stats[f"CAT_{cat_l1}"] += 1

    wb.close()

    # File 1: Master All Products (1,351 items with Published=TRUE/FALSE per category)
    f_master = output_dir / "Shopify_Rabuta_Master_Catalog.csv"
    _write_csv(f_master, all_rows)

    # File 2: Webshop Online Only (1,090 items - strictly products allowed online)
    f_online = output_dir / "Shopify_Rabuta_Webshop_Online_Only.csv"
    _write_csv(f_online, online_rows)

    # File 3: In-Store Offline Only (261 items - Frostvarer, Ris, Fisk, Rejer)
    f_offline = output_dir / "Shopify_Rabuta_InStore_Offline_Only.csv"
    _write_csv(f_offline, offline_rows)

    print("\n=================================================================")
    print("                 SHOPIFY CSV EXPORT COMPLETE                     ")
    print("=================================================================")
    print(f"Total Products Processed : {len(all_rows)}")
    print(f"  - Webshop Online       : {len(online_rows)} (Published=TRUE, Status=active)")
    print(f"  - In-Store Only (Nej)  : {len(offline_rows)} (Published=FALSE, Status=draft)")
    print("-----------------------------------------------------------------")
    print(f"1. Master Catalog CSV    : {f_master}")
    print(f"2. Webshop Online Only   : {f_online}")
    print(f"3. In-Store Offline Only : {f_offline}")
    print("=================================================================")

    return f_master, f_online, f_offline


def _write_csv(path: Path, rows: List[Dict[str, Any]]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=SHOPIFY_28_COLUMNS, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)
    print(f"[CSV EXPORT] Saved {path.name} ({len(rows)} products, {path.stat().st_size / (1024*1024):.2f} MB)")


if __name__ == "__main__":
    master_excel = ROOT / "output/shopify-zohanubis-image/webshop_front_hd/Rabuta_master_SKU_categories-image.xlsx"
    esbjerg_excel = ROOT / "source-import/new_source_from_customer/11-10-new-requirment/Esbjerg_Asian_Market_Master_1659_updated.xlsx"
    output_dir = ROOT / "output/shopify-zohanubis-image"

    export_shopify_files(master_excel, esbjerg_excel, output_dir)

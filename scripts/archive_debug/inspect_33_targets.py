# -*- coding: utf-8 -*-
import sys
import os
import openpyxl
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
f = os.path.join(BASE_DIR, r"output\shopify_import_final\Danh_sach_san_pham_can_chup_lai_Barcode.xlsx")

wb = openpyxl.load_workbook(f, data_only=True)
ws = wb.active
headers = [c.value for c in ws[3]]
rows = list(ws.iter_rows(min_row=4, values_only=True))

SOURCE_DIR = os.path.join(BASE_DIR, "source-import")
folder_map = {
    "G1": os.path.join(SOURCE_DIR, "G1_original"),
    "G2": os.path.join(SOURCE_DIR, "G2_original"),
    "G3": os.path.join(SOURCE_DIR, "G3_Original"),
    "G4": os.path.join(SOURCE_DIR, "G4-original"),
    "G5": os.path.join(SOURCE_DIR, "Nyt25.09.26"),
}

print(f"Total 33 target items:")
for idx, r in enumerate(rows, start=1):
    d = dict(zip(headers, r))
    batch = d.get("Thư mục nguồn (Batch)")
    photo = d.get("Tên file ảnh gốc")
    name = d.get("Tên sản phẩm tiếng Đan Mạch (Fuldt navn DA)")
    brand = d.get("Thương hiệu (Brand)")
    
    folder = folder_map.get(batch, "")
    img_p = os.path.join(folder, photo) if folder and photo else ""
    if not os.path.exists(img_p):
        alt = [f for f in os.listdir(folder) if f.lower() == photo.lower()] if os.path.exists(folder) else []
        if alt:
            img_p = os.path.join(folder, alt[0])
            
    exists = os.path.exists(img_p)
    size_str = ""
    if exists:
        sz = os.path.getsize(img_p)
        size_str = f"{sz:,} bytes"
        
    print(f"{idx:2d}. [{batch}] {photo:<25} | Size: {size_str:<12} | {str(brand):<15} | {str(name)[:35]}")

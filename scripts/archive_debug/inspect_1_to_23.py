# -*- coding: utf-8 -*-
import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
f = os.path.join(BASE_DIR, r"output\shopify_import_final\Danh_sach_san_pham_can_chup_lai_Barcode.xlsx")

wb = openpyxl.load_workbook(f, data_only=True)
ws = wb.active
headers = [c.value for c in ws[3]]
rows = list(ws.iter_rows(min_row=4, values_only=True))

SOURCE_DIR = os.path.join(BASE_DIR, "source-import")
folder_map = {
    "G1": os.path.join(SOURCE_DIR, "G1_original"),
    "G2": os.path.join(SOURCE_DIR, "G2_original"),
    "G3": os.path.join(SOURCE_DIR, "G3_Original"),
    "G4": os.path.join(SOURCE_DIR, "G4-original"),
    "G5": os.path.join(SOURCE_DIR, "Nyt25.09.26"),
}

print(f"Total rows in reshoot file: {len(rows)}\n")
print(f"{'STT':<4} | {'Batch':<5} | {'File Name':<28} | {'Brand':<15} | {'Product Name'}")
print("-" * 100)

for idx, r in enumerate(rows[:23], start=1):
    d = dict(zip(headers, r))
    batch = d.get("Thư mục nguồn (Batch)")
    photo = d.get("Tên file ảnh gốc")
    name = d.get("Tên sản phẩm tiếng Đan Mạch (Fuldt navn DA)")
    brand = d.get("Thương hiệu (Brand)")
    
    folder = folder_map.get(batch, "")
    img_p = os.path.join(folder, photo) if folder and photo else ""
    if not os.path.exists(img_p):
        alt = [f for f in os.listdir(folder) if f.lower() == photo.lower()] if os.path.exists(folder) else []
        if alt:
            img_p = os.path.join(folder, alt[0])
            
    print(f"{idx:3d} | {batch:<5} | {photo:<28} | {str(brand):<15} | {str(name)[:45]}")

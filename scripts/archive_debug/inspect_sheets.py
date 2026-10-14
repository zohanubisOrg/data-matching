# -*- coding: utf-8 -*-
import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"

f1 = os.path.join(BASE_DIR, r"source-import\Ten_hang_nhap.xlsx")
f2 = os.path.join(BASE_DIR, r"source-import\Products_imported_to_Shopbox_19_08_26_with_SKU.xlsx")

wb1 = openpyxl.load_workbook(f1, data_only=True)
for sheet in wb1.sheetnames:
    ws = wb1[sheet]
    rows = list(ws.iter_rows(values_only=True))
    print(f"File 1 Sheet [{sheet}]: {len(rows)} rows")
    if rows:
        print("  Headers:", rows[0])
        print("  Sample 1:", rows[1] if len(rows) > 1 else "None")

print("\n" + "="*80 + "\n")

wb2 = openpyxl.load_workbook(f2, data_only=True)
for sheet in wb2.sheetnames:
    ws = wb2[sheet]
    rows = list(ws.iter_rows(values_only=True))
    print(f"File 2 Sheet [{sheet}]: {len(rows)} rows")
    if rows:
        print("  Headers:", rows[0])
        print("  Sample 1:", rows[1] if len(rows) > 1 else "None")

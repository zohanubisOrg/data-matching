# -*- coding: utf-8 -*-
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

wb = openpyxl.load_workbook(r"output\shopify_import_final\Master_All_1659_Products_With_Images.xlsx", data_only=True)
ws = wb.active

headers = [cell.value for cell in ws[3]]
print("Headers:", headers)

unverified = []
total = 0
for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
    total += 1
    row_dict = dict(zip(headers, row))
    status = str(row_dict.get("Trạng thái đối soát") or "")
    sku = str(row_dict.get("POS SKU") or "").strip()
    bc = str(row_dict.get("Barcode") or "").strip()
    if "CHƯA XÁC MINH" in status or not sku or not bc:
        unverified.append((row_idx, row_dict))

print(f"Total rows: {total}, Unverified/Missing SKU/BC: {len(unverified)}")
for idx, d in unverified[:15]:
    print(f"Row {idx:4d} | Batch: {d.get('Thư mục nguồn (Batch)')} | File: {d.get('Tên file ảnh gốc')} | Brand: {d.get('Thương hiệu (Brand)')} | Name DA: {d.get('Tên tiếng Đan Mạch (Fuldt navn DA)')} | Status: {d.get('Trạng thái đối soát')}")


import openpyxl, json, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('output/g1_delivery/G1_Master_Delivery_v2_latest.xlsx')
ws = wb.active

target_ids = ['G1-0041', 'G1-0314', 'G1-0318', 'G1-0319', 'G1-0323', 'G1-0324', 'G1-0325', 'G1-0326', 'G1-0330', 'G1-0350', 'G1-0360']

with open('output/g1_delivery/g1_deep_scan_65_cache.json', 'r', encoding='utf-8') as f:
    cache = json.load(f)

for r in range(3, ws.max_row + 1):
    rec_id = ws.cell(r, 6).value
    if rec_id in target_ids:
        fname_cell = str(ws.cell(r, 58).value or '')
        actual_fname = fname_cell.replace('G1: ', '').strip()
        bc = ws.cell(r, 1).value
        name = ws.cell(r, 3).value
        c_info = cache.get(actual_fname, {})
        print(f"Record {rec_id} (Row {r}):")
        print(f"  File: {actual_fname} | Barcode: {bc} | Name: {name}")
        print(f"  front_box_2d: {c_info.get('front_box_2d')}")
        print(f"  back_box_2d:  {c_info.get('back_box_2d')}")
        print(f"  front_rot:    {c_info.get('front_rotation')}")
        print(f"  back_rot:     {c_info.get('back_rotation')}")
        print("-" * 50)

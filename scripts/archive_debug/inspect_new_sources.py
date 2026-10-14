# -*- coding: utf-8 -*-
import sys
import os
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"

f1 = os.path.join(BASE_DIR, r"source-import\Ten_hang_nhap.xlsx")
f2 = os.path.join(BASE_DIR, r"source-import\Products_imported_to_Shopbox_19_08_26_with_SKU.xlsx")

print("Checking f1:", f1, "Exists:", os.path.exists(f1))
if os.path.exists(f1):
    wb1 = openpyxl.load_workbook(f1, data_only=True)
    print("Sheets in f1:", wb1.sheetnames)
    ws1 = wb1.active
    rows1 = list(ws1.iter_rows(values_only=True))
    print(f"Total rows in f1: {len(rows1)}")
    for i in range(min(5, len(rows1))):
        print(f"  Row {i+1}: {rows1[i]}")

print("\n" + "="*80 + "\n")

print("Checking f2:", f2, "Exists:", os.path.exists(f2))
if os.path.exists(f2):
    wb2 = openpyxl.load_workbook(f2, data_only=True)
    print("Sheets in f2:", wb2.sheetnames)
    ws2 = wb2.active
    rows2 = list(ws2.iter_rows(values_only=True))
    print(f"Total rows in f2: {len(rows2)}")
    for i in range(min(5, len(rows2))):
        print(f"  Row {i+1}: {rows2[i]}")

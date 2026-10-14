# -*- coding: utf-8 -*-
import sys
import os
import re
import json
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
SOURCE_DIR = os.path.join(BASE_DIR, "source-import")
CUSTOMER_SRC = os.path.join(SOURCE_DIR, "new_source_from_customer")

f1 = os.path.join(SOURCE_DIR, "Ten_hang_nhap.xlsx")
f2 = os.path.join(SOURCE_DIR, "Products_imported_to_Shopbox_19_08_26_with_SKU.xlsx")
pos_active_file = os.path.join(CUSTOMER_SRC, "27.09.26.POS_ACTIVE_FINAL.xlsx")

def clean_bc(v):
    s = str(v or "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)

# Load Unified POS
all_pos_by_bc = {}
wb2 = openpyxl.load_workbook(f2, data_only=True)
for r in list(wb2["Ten_ban"].iter_rows(values_only=True))[1:]:
    c_bc = clean_bc(r[7])
    if c_bc:
        all_pos_by_bc[c_bc] = r
        all_pos_by_bc[c_bc.lstrip("0")] = r

wb_act = openpyxl.load_workbook(pos_active_file, data_only=True)
for r in list(wb_act.active.iter_rows(values_only=True))[1:]:
    c_bc = clean_bc(r[7])
    if c_bc:
        all_pos_by_bc[c_bc] = r
        all_pos_by_bc[c_bc.lstrip("0")] = r

# Load caches & raw items
caches = {}
with open(os.path.join(BASE_DIR, r"output\g1_delivery\g1_deep_scan_65_cache.json"), "r", encoding="utf-8") as f:
    caches["G1"] = json.load(f)
with open(os.path.join(BASE_DIR, r"output\g1_delivery\g1_47_missing_cache.json"), "r", encoding="utf-8") as f:
    caches["G1"].update(json.load(f))
for b in ["g2", "g3", "g4", "g5"]:
    with open(os.path.join(BASE_DIR, f"output\\shopify_import\\{b}_vision_cache.json"), "r", encoding="utf-8") as f:
        caches[b.upper()] = json.load(f)

g_dfs = {}
for b in ["G2", "G3", "G4", "G5"]:
    g_dfs[b] = pd.read_excel(os.path.join(BASE_DIR, f"output\\shopify_import\\{b}_master_products.xlsx"))

all_1659 = []
for fname, item in caches["G1"].items():
    rec = dict(item)
    rec["batch"] = "G1"
    rec["source_photo"] = fname
    all_1659.append(rec)

for b in ["G2", "G3", "G4", "G5"]:
    for _, row in g_dfs[b].iterrows():
        rec = row.to_dict()
        rec["batch"] = b
        rec["source_photo"] = str(rec.get("source_photo") or "")
        all_1659.append(rec)

remaining_83 = []
for idx, item in enumerate(all_1659, start=1):
    raw_bc_item = clean_bc(item.get("barcode") or item.get("Barcode"))
    bc_status = str(item.get("barcode_check_status") or "").strip()
    is_borrowed = ("barcode từ" in bc_status) or ("trùng mã" in bc_status)
    
    if is_borrowed or not raw_bc_item or (raw_bc_item not in all_pos_by_bc and raw_bc_item.lstrip("0") not in all_pos_by_bc):
        remaining_83.append((idx, item["batch"], item["source_photo"], item))

print(f"Total remaining without POS barcode match: {len(remaining_83)}")

for idx, batch, photo, item in remaining_83[:25]:
    brand = item.get("brand") or ""
    name_da = item.get("fuldt_navn_da") or item.get("product_name") or ""
    raw_bc = item.get("barcode") or item.get("Barcode") or ""
    print(f"Row {idx:4d} | Batch: {batch} | File: {photo:<20} | BC in item: {str(raw_bc):<15} | Brand: {str(brand):<15} | Name: {str(name_da)[:40]}")

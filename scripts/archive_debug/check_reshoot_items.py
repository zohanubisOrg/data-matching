# -*- coding: utf-8 -*-
import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
f = os.path.join(BASE_DIR, r"output\shopify_import_final\Danh_sach_san_pham_can_chup_lai_Barcode.xlsx")

wb = openpyxl.load_workbook(f, data_only=True)
ws = wb.active
headers = [c.value for c in ws[3]]
rows = list(ws.iter_rows(min_row=4, values_only=True))

print(f"Total items needing reshoot: {len(rows)}\n")
for idx, r in enumerate(rows, start=1):
    d = dict(zip(headers, r))
    photo = d.get("Tên file ảnh gốc")
    name = d.get("Tên sản phẩm tiếng Đan Mạch (Fuldt navn DA)")
    reason = d.get("Lý do cần chụp lại")
    print(f"{idx:2d}. {photo:<25} | {str(name)[:35]:<35} | Lý do: {reason}")

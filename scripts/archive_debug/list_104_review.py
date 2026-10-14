# -*- coding: utf-8 -*-
import sys
import os
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
review_file = os.path.join(BASE_DIR, r"output\shopify_import_final\Customer_Review_Unmatched_Images.xlsx")

wb = openpyxl.load_workbook(review_file, data_only=True)
ws = wb.active

headers = [c.value for c in ws[3]]
rows = list(ws.iter_rows(min_row=4, values_only=True))

print(f"Total rows in review file: {len(rows)}")
for idx, r in enumerate(rows, start=1):
    d = dict(zip(headers, r))
    batch = d.get("Thư mục nguồn (Batch)")
    photo = d.get("Tên file ảnh gốc")
    brand = d.get("Thương hiệu (Brand)")
    name_da = d.get("Tên tiếng Đan Mạch (Fuldt navn DA)")
    print(f"{idx:3d}. Batch: {batch} | Photo: {photo:<20} | Brand: {str(brand):<15} | Name: {str(name_da)[:40]}")

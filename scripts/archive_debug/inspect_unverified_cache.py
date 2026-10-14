# -*- coding: utf-8 -*-
import sys
import os
import json
import re
import openpyxl
import pandas as pd
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = r"D:\Work\shopify"
CUSTOMER_SRC = os.path.join(BASE_DIR, "source-import", "new_source_from_customer")

# Load POS
pos_file = os.path.join(CUSTOMER_SRC, "27.09.26.POS_ACTIVE_FINAL.xlsx")
wb_pos = openpyxl.load_workbook(pos_file, data_only=True)
ws_pos = wb_pos.active
pos_rows = list(ws_pos.iter_rows(values_only=True))[1:]

def clean_bc(v):
    s = str(v or "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[^\w]", "", s)

pos_items = []
for r in pos_rows:
    pos_items.append({
        "pos_category": str(r[0] or "").strip(),
        "product_name": str(r[1] or "").strip(),
        "official_name": str(r[2] or "").strip(),
        "printer_name": str(r[3] or "").strip(),
        "purchase_price": r[4],
        "selling_price": r[5],
        "unit": r[6],
        "raw_barcode": str(r[7] or "").strip(),
        "clean_barcode": clean_bc(r[7]),
        "sku": str(r[8] or "").strip(),
        "product_id": str(r[9] or "").strip(),
        "tax_group": r[10],
    })

print(f"Total POS items: {len(pos_items)}")

# Load Master Excel unverified rows
wb_m = openpyxl.load_workbook(os.path.join(BASE_DIR, r"output\shopify_import_final\Master_All_1659_Products_With_Images.xlsx"), data_only=True)
ws_m = wb_m.active
headers = [c.value for c in ws_m[3]]

unverified_rows = []
for row_idx, row in enumerate(ws_m.iter_rows(min_row=4, values_only=True), start=4):
    d = dict(zip(headers, row))
    status = str(d.get("Trạng thái đối soát") or "").strip()
    sku = str(d.get("POS SKU") or "").strip()
    bc = str(d.get("Barcode") or "").strip()
    if "CHƯA XÁC MINH" in status or not sku or not bc:
        unverified_rows.append((row_idx, d))

print(f"Unverified / Missing in Master: {len(unverified_rows)}")

# Load all vision caches
caches = {}
with open(os.path.join(BASE_DIR, r"output\g1_delivery\g1_deep_scan_65_cache.json"), "r", encoding="utf-8") as f:
    caches["G1"] = json.load(f)
for b in ["g2", "g3", "g4", "g5"]:
    with open(os.path.join(BASE_DIR, f"output\\shopify_import\\{b}_vision_cache.json"), "r", encoding="utf-8") as f:
        caches[b.upper()] = json.load(f)

# Inspect what cache data exists for these unverified rows
found_in_cache = 0
sample_analyses = []

for idx, d in unverified_rows:
    batch = d.get("Thư mục nguồn (Batch)")
    photo = d.get("Tên file ảnh gốc")
    stem = os.path.splitext(photo)[0] if photo else ""
    
    b_cache = caches.get(batch, {})
    c_data = b_cache.get(photo) or b_cache.get(stem) or b_cache.get(photo.lower())
    if not c_data:
        # try search by stem in keys
        for k, v in b_cache.items():
            if os.path.splitext(k)[0] == stem:
                c_data = v
                break
    if c_data:
        found_in_cache += 1
        if len(sample_analyses) < 10:
            sample_analyses.append((batch, photo, c_data))

print(f"Items found in vision cache: {found_in_cache} / {len(unverified_rows)}")

print("\n--- SAMPLE CACHE ENTRIES ---")
for batch, photo, c in sample_analyses[:5]:
    print(f"\nBatch: {batch} | Photo: {photo}")
    print(f"Keys: {list(c.keys())}")
    for k in ["brand", "literal_label_name", "product_name", "barcode", "detected_barcode", "variant", "net_quantity", "net_unit", "category"]:
        if k in c:
            print(f"  {k}: {c[k]}")
